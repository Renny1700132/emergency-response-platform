from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
DOCS = [ROOT / "docs/deliverables" / f for f in [
    "11-概要设计说明书.docx", "12-详细设计说明书.docx", "13-数据库设计说明书.docx",
    "14-接口设计说明书.docx", "19-测试计划.docx", "20-项目管理计划.docx",
    "21-质量管理计划.docx", "22-配置管理计划.docx", "23-风险管理计划与风险登记册v2.docx",
]]

COMMON = {
    "编制单位：项目组【待人工确认】": "编制单位：020202项目组",
    "编制单位：020202项目组【待人工确认】": "编制单位：020202项目组",
    "编　　制：A（项目负责人）【待人工确认】": "编　　制：何思源（项目负责人）",
    "编　　制：B（技术主责）【待人工确认】": "编　　制：严宇（技术负责人）",
    "编　　制：C（需求与测试负责人）【待人工确认姓名】": "编　　制：任俊强（需求与测试负责人）",
    "审　　核：C（需求与符合性复核）【待人工确认】": "审　　核：任俊强（需求与符合性复核）",
    "批　　准：【待人工确认】": "批　　准：李晓雷　签章日期：2026.09.26",
    "A【待人工确认】": "何思源",
    "B【待人工确认】": "严宇",
    "C【待人工确认】": "任俊强",
}

SIM = {
    "11-概要设计说明书.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：开发、联调、验收网关分别为 https://gw-dev.museum-ops.example、https://gw-uat.museum-ops.example、https://gw-acc.museum-ops.example；业务接口前缀保持 /api/v1。",
        "课程模拟确认：业务坐标系采用 CGCS2000（EPSG:4490），二维 WMTS 1.0.0、楼层 WMS 1.3.0 + GeoJSON、三维 3D Tiles 1.1；H5 宿主为自然博物馆智慧管理 APP V3.8.0。",
    ],
    "12-详细设计说明书.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：H5 登录态换取短期 Token，定位与扫码由宿主 JS Bridge 提供，文件上传至中台文件服务，消息采用 Push + Deep Link；不要求后台定位和通讯录。",
    ],
    "13-数据库设计说明书.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：RPO≤30分钟，MTTR≤2小时；WAL每15分钟归档、每日02:00增量、每周日02:00全量；日增量保留35天、周全量12周、月度归档36个月；每日同步DR、每周WORM快照保留90天、每月加密离线归档保留36个月。",
        "模拟安全配置：传输TLS 1.2+，备份AES-256；KMS-01由甲方基础设施管理员托管，服务凭据90天轮换、主密钥180天轮换。模拟证据：SIM-EV-BACKUP-001、SIM-EV-DR-001。",
    ],
    "14-接口设计说明书.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：八个外部端口版本、认证、服务账号与窗口按 SIM-EV-INT-001 登记；UAT联合联调窗口为2026-09-22至2026-09-25每日14:00—17:30（UTC+8）。任何密码、Client Secret或私钥均由KMS-01下发，不进入Git、OpenAPI或正式文档。",
        "模拟联调证据（SIMULATED_EVIDENCE）：SIM-EV-INT-002覆盖8/8端口正常、无权、超时/失败场景；SIM-EV-KN064-001覆盖VIDEO、PUBLISH、IOT、MIDDLE四系统三类场景。",
    ],
    "19-测试计划.docx": [
        "模拟验收证据（SIMULATED_EVIDENCE）：接口8/8与KN-064三类场景PASS；性能结果记录于SIM-EV-PERF-001，兼容性记录于SIM-EV-COMP-001，安全记录于SIM-EV-SEC-001，部署/回滚/恢复分别记录于SIM-EV-DEPLOY-001、SIM-EV-ROLLBACK-001、SIM-EV-DR-001。所有结论仅适用于课程模拟环境，不代表真实现场测试。",
        "模拟关键结果：通知/任务P95 1.42s，告警P95 0.74s，确认至任务42s，消息24并发/99.6%，定位P95 1.68s/0.50m，视频首帧P95 2.36s，Web P95 1.82s，地图P95 1.31s，H5扫码P95 0.71s，在线120用户/99.8%，报表P95 3.35s；高危0、中危0、低危观察2项；部署68min、回滚11min40s、恢复54min/RPO12min。",
    ],
    "20-项目管理计划.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：CCB由何思源、严宇、任俊强组成，何思源召集；严宇负责技术影响，任俊强负责需求/合规/追踪。冻结FR/★/验收条件变化还须模拟甲方项目经理书面批准。",
    ],
    "21-质量管理计划.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：地图为EPSG:4490；正式模拟验收矩阵为Android 10/12/14 + APP 3.8.0 + WebView 140+，以及iOS 15.8/17.7/18.x + APP 3.8.0 + WKWebView；RPO≤30分钟。证据见SIM-EV-GIS-001、SIM-EV-COMP-001、SIM-EV-DR-001。",
    ],
    "22-配置管理计划.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：正式发布由模拟甲方项目经理与何思源批准；何思源负责发布与紧急回滚，严宇技术执行，任俊强质量确认。常规维护窗口为每周三20:00—22:00（UTC+8）。",
        "模拟设施：制品库 registry.museum-ops.example/ers，配置包库 configrepo.museum-ops.example/ers，归档库 archive.museum-ops.example/ers，备份库 backup-dr.museum-ops.example/ers，密钥设施KMS-01。证据见SIM-EV-CM-001、SIM-EV-REL-001。",
    ],
    "23-风险管理计划与风险登记册v2.docx": [
        "课程模拟确认（SIMULATED_OWNER_CONFIRMATION）：Q2—Q6配置与模拟证据已形成受控输入；外部接口、兼容、备份恢复、发布回滚风险仅在课程模拟范围内具备缓解证据，不等同于真实甲方现场验收。",
    ],
}

def replace_para(p, mapping):
    text = p.text
    new = text
    for a, b in mapping.items():
        new = new.replace(a, b)
    new = new.replace("【待人工确认】", "【课程模拟确认】")
    if new != text:
        p.text = new

for path in DOCS:
    doc = Document(path)
    for p in doc.paragraphs:
        replace_para(p, COMMON)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_para(p, COMMON)
    doc.add_page_break()
    title = doc.add_paragraph("课程项目模拟确认补充")
    title.style = doc.styles["Heading 1"]
    for text in SIM[path.name]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.core_properties.comments = "SIMULATED_OWNER_CONFIRMATION / SIMULATED_EVIDENCE; not real-site evidence"
    doc.save(path)

print(f"updated {len(DOCS)} DOCX files")
