from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document


EXPECTED_FR = {f"G2-FR-{i:03d}" for i in range(1, 40)}
EXPECTED_AC = {f"AC-G2-FR-{i:03d}-{j:02d}" for i in range(1, 40) for j in range(1, 4)}
EXPECTED_DESIGN = {f"DBD-TR-{i:03d}" for i in range(1, 40)}
NON_STAR = {"G2-FR-007", "G2-FR-030", "G2-FR-032", "G2-FR-033", "G2-FR-038"}


def tokens(text: str, pattern: str) -> set[str]:
    return set(re.findall(pattern, text))


def audit_text(text: str) -> dict:
    fr = tokens(text, r"(?<!AC-)G2-FR-\d{3}")
    ac = tokens(text, r"AC-G2-FR-\d{3}-\d{2}")
    design = tokens(text, r"DBD-TR-\d{3}")
    rows = {}
    for line in text.splitlines():
        if not line.startswith("| DBD-TR-"):
            continue
        m = re.search(r"(DBD-TR-\d{3}).*?(G2-FR-\d{3}).*?(非★|★)", line)
        if m:
            rows[m.group(2)] = {"design": m.group(1), "star": m.group(3)}
    star = {fr_id for fr_id, value in rows.items() if value["star"] == "★"}
    non_star = {fr_id for fr_id, value in rows.items() if value["star"] == "非★"}
    return {
        "fr_count": len(fr & EXPECTED_FR),
        "ac_count": len(ac & EXPECTED_AC),
        "design_count": len(design & EXPECTED_DESIGN),
        "matrix_row_count": len(rows),
        "star_count": len(star),
        "non_star_count": len(non_star),
        "missing_fr": sorted(EXPECTED_FR - fr),
        "missing_ac": sorted(EXPECTED_AC - ac),
        "missing_design": sorted(EXPECTED_DESIGN - design),
        "wrong_star": sorted((EXPECTED_FR - NON_STAR) ^ star),
        "wrong_non_star": sorted(NON_STAR ^ non_star),
    }


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--markdown", required=True)
    p.add_argument("--docx", required=True)
    p.add_argument("--out", required=True)
    args = p.parse_args()
    md = Path(args.markdown).read_text(encoding="utf-8")
    start = md.index("### 8.1 需求与验收数据设计追踪矩阵")
    end = md.index("### 8.2 未决事项和阻断关系")
    matrix = md[start:end]
    doc = Document(args.docx)
    doc_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(
        "| " + " | ".join(c.text.replace("\n", "<br>") for c in row.cells) + " |"
        for t in doc.tables for row in t.rows
    )
    result = {"markdown_matrix": audit_text(matrix), "formal_docx": audit_text(doc_text)}
    result["pass"] = all(
        v["fr_count"] == 39 and v["ac_count"] == 117 and v["design_count"] == 39
        and v["star_count"] == 34 and v["non_star_count"] == 5
        and not v["missing_fr"] and not v["missing_ac"] and not v["missing_design"]
        and not v["wrong_star"] and not v["wrong_non_star"]
        for v in result.values() if isinstance(v, dict)
    )
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["pass"] else 1)


if __name__ == "__main__":
    main()
