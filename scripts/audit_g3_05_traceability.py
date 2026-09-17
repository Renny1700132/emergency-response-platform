from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from docx import Document

EXPECTED_FR={f"G2-FR-{i:03d}" for i in range(1,40)}
EXPECTED_AC={f"AC-G2-FR-{i:03d}-{j:02d}" for i in range(1,40) for j in range(1,4)}
EXPECTED_DESIGN={f"DLD-TR-{i:03d}" for i in range(1,40)}
NON_STAR={"G2-FR-007","G2-FR-030","G2-FR-032","G2-FR-033","G2-FR-038"}

def audit(text):
    fr=set(re.findall(r"(?<!AC-)G2-FR-\d{3}",text)); ac=set(re.findall(r"AC-G2-FR-\d{3}-\d{2}",text)); design=set(re.findall(r"DLD-TR-\d{3}",text)); rows={}
    for line in text.splitlines():
        if "DLD-TR-" not in line: continue
        m=re.search(r"(DLD-TR-\d{3}).*?(G2-FR-\d{3}).*?(非★|★)",line,re.S)
        if m: rows[m.group(2)]={"design":m.group(1),"star":m.group(3)}
    star={k for k,v in rows.items() if v["star"]=="★"}; non={k for k,v in rows.items() if v["star"]=="非★"}
    return {"fr_count":len(fr&EXPECTED_FR),"ac_count":len(ac&EXPECTED_AC),"design_count":len(design&EXPECTED_DESIGN),"matrix_row_count":len(rows),"star_count":len(star),"non_star_count":len(non),"missing_fr":sorted(EXPECTED_FR-fr),"missing_ac":sorted(EXPECTED_AC-ac),"missing_design":sorted(EXPECTED_DESIGN-design),"wrong_star":sorted((EXPECTED_FR-NON_STAR)^star),"wrong_non_star":sorted(NON_STAR^non)}

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--markdown",required=True); ap.add_argument("--docx",required=True); ap.add_argument("--out",required=True); args=ap.parse_args()
    md=Path(args.markdown).read_text(encoding="utf-8"); start=md.index("### 12.1 39 FR / 117 AC"); end=md.index("### 12.2 关键数字")
    d=Document(args.docx); doctext="\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join("| "+" | ".join(c.text.replace("\n","<br>") for c in row.cells)+" |" for t in d.tables for row in t.rows)
    result={"markdown_matrix":audit(md[start:end]),"formal_docx":audit(doctext)}
    result["pass"]=all(v["fr_count"]==39 and v["ac_count"]==117 and v["design_count"]==39 and v["matrix_row_count"]==39 and v["star_count"]==34 and v["non_star_count"]==5 and not v["missing_fr"] and not v["missing_ac"] and not v["missing_design"] and not v["wrong_star"] and not v["wrong_non_star"] for v in result.values() if isinstance(v,dict))
    Path(args.out).write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8"); print(json.dumps(result,ensure_ascii=False,indent=2)); raise SystemExit(0 if result["pass"] else 1)

if __name__=="__main__": main()
