from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import yaml
from docx import Document


EXPECTED_FR = {f"G2-FR-{i:03d}" for i in range(1, 40)}
EXPECTED_AC = {f"AC-G2-FR-{i:03d}-{j:02d}" for i in range(1, 40) for j in range(1, 4)}
EXPECTED_TRACE = {f"API-TR-{i:03d}" for i in range(1, 40)}
NON_STAR = {"G2-FR-007", "G2-FR-030", "G2-FR-032", "G2-FR-033", "G2-FR-038"}
SAMPLE_TERMS = ["澜图", "遥感影像", "samgeo", "QGIS", "segment-geospatial"]


def refs(obj):
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k == "$ref" and isinstance(v, str):
                yield v
            else:
                yield from refs(v)
    elif isinstance(obj, list):
        for item in obj:
            yield from refs(item)


def resolve_local(doc, ref):
    cur = doc
    for part in ref.removeprefix("#/").split("/"):
        cur = cur[part.replace("~1", "/").replace("~0", "~")]
    return cur


def text_audit(text):
    frs = set(re.findall(r"(?<!AC-)G2-FR-\d{3}", text))
    acs = set(re.findall(r"AC-G2-FR-\d{3}-\d{2}", text))
    traces = set(re.findall(r"API-TR-\d{3}", text))
    star_rows, non_rows = set(), set()
    for line in text.splitlines():
        m = re.search(r"API-TR-\d{3}.*?(G2-FR-\d{3}).*?(非★|★)", line)
        if m:
            (non_rows if m.group(2) == "非★" else star_rows).add(m.group(1))
    return {
        "fr_count": len(frs & EXPECTED_FR), "ac_count": len(acs & EXPECTED_AC),
        "trace_count": len(traces & EXPECTED_TRACE), "star_count": len(star_rows),
        "non_star_count": len(non_rows), "missing_fr": sorted(EXPECTED_FR - frs),
        "missing_ac": sorted(EXPECTED_AC - acs), "missing_trace": sorted(EXPECTED_TRACE - traces),
        "star_mismatch": sorted((EXPECTED_FR - NON_STAR) ^ star_rows),
        "non_star_mismatch": sorted(NON_STAR ^ non_rows),
        "sample_pollution": [x for x in SAMPLE_TERMS if x.lower() in text.lower()],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--markdown", required=True); ap.add_argument("--yaml", required=True)
    ap.add_argument("--docx", required=True); ap.add_argument("--out", required=True)
    args = ap.parse_args()
    md = Path(args.markdown).read_text(encoding="utf-8")
    raw = Path(args.yaml).read_text(encoding="utf-8"); api = yaml.safe_load(raw)
    doc = Document(args.docx)
    doc_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(" | ".join(c.text.replace("\n", "<br>") for c in row.cells) for t in doc.tables for row in t.rows)
    operations = []
    for path, item in api.get("paths", {}).items():
        for method, op in item.items():
            if method.upper() not in {"GET", "POST", "PUT", "PATCH", "DELETE"}: continue
            operations.append((method.upper(), path, op))
    op_ids = [x[2].get("operationId") for x in operations]
    yaml_frs = {x for _, _, op in operations for x in op.get("x-requirements", [])}
    yaml_acs = {x for _, _, op in operations for x in op.get("x-acceptance", [])}
    ref_errors = []
    for ref in refs(api):
        if ref.startswith("#/"):
            try: resolve_local(api, ref)
            except Exception: ref_errors.append(ref)
    write_missing_idempotency = []
    for method, path, op in operations:
        if method == "GET": continue
        params = [p.get("$ref") if isinstance(p, dict) else None for p in op.get("parameters", [])]
        if "#/components/parameters/IdempotencyKey" not in params:
            write_missing_idempotency.append(f"{method} {path}")
    result = {
        "task": "G3-06", "openapi_version": api.get("openapi"),
        "operation_count": len(operations), "unique_operation_ids": len(set(op_ids)),
        "duplicate_operation_ids": sorted({x for x in op_ids if op_ids.count(x) > 1}),
        "unresolved_local_refs": sorted(set(ref_errors)),
        "write_missing_idempotency": write_missing_idempotency,
        "security_defined": bool(api.get("security") and api.get("components", {}).get("securitySchemes")),
        "yaml_fr_count": len(yaml_frs & EXPECTED_FR), "yaml_ac_count": len(yaml_acs & EXPECTED_AC),
        "yaml_missing_fr": sorted(EXPECTED_FR - yaml_frs), "yaml_missing_ac": sorted(EXPECTED_AC - yaml_acs),
        "issue_boundary": {
            "documented": "ISSUE-G3-01-001" in raw and "ISSUE-G3-01-001" in md and "ISSUE-G3-01-001" in doc_text,
            "pending_external_schema": "PENDING_EXTERNAL_EVIDENCE" in raw,
            "access_no_auto_replay": "NO_AUTOMATIC_REPLAY" in raw,
        },
        "markdown": text_audit(md), "formal_docx": text_audit(doc_text),
        "yaml_sample_pollution": [x for x in SAMPLE_TERMS if x.lower() in raw.lower()],
    }
    ta = [result["markdown"], result["formal_docx"]]
    result["pass"] = all([
        result["openapi_version"] == "3.0.3", len(operations) >= 40,
        len(set(op_ids)) == len(op_ids), not ref_errors, not write_missing_idempotency,
        result["security_defined"], result["yaml_fr_count"] == 39, result["yaml_ac_count"] == 117,
        not result["yaml_missing_fr"], not result["yaml_missing_ac"], not result["yaml_sample_pollution"],
        all(result["issue_boundary"].values()),
        all(x["fr_count"] == 39 and x["ac_count"] == 117 and x["trace_count"] == 39 and x["star_count"] == 34 and x["non_star_count"] == 5 and not x["missing_fr"] and not x["missing_ac"] and not x["missing_trace"] and not x["star_mismatch"] and not x["non_star_mismatch"] and not x["sample_pollution"] for x in ta),
    ])
    Path(args.out).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["pass"] else 1)


if __name__ == "__main__":
    main()
