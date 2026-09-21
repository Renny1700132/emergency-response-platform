"""Audit the G3-10 M3 consistency candidate without claiming a freeze."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]


def read(path: str) -> str:
    return (ROOT / path).read_text(encoding="utf-8")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def docx_text(path: Path) -> str:
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    chunks += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(chunks)


def task_status(tasks: str, task_id: str) -> str:
    match = re.search(rf"^\| {re.escape(task_id)} \|.*?\| ([^|]+) \|$", tasks, re.MULTILINE)
    if not match:
        raise AssertionError(f"missing task row: {task_id}")
    return match.group(1).strip()


def last_issue_status(issues: str, issue_id: str) -> str:
    match = re.search(
        rf"^##+ {re.escape(issue_id)}\s*$([\s\S]*?)(?=^##+ ISSUE-|\Z)",
        issues,
        re.MULTILINE,
    )
    if not match:
        raise AssertionError(f"missing issue: {issue_id}")
    statuses = re.findall(r"^- 状态：(.+)$", match.group(1), re.MULTILINE)
    if not statuses:
        raise AssertionError(f"missing status: {issue_id}")
    return statuses[-1].strip()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()

    tasks = read("tasks.md")
    issues = read("control/issues.md")
    spec = read("docs/work/B_TECH/spec.md")
    rtm = read("docs/work/C_REQ/rtm_v1.md")
    nfr = read("docs/deliverables/18-非功能设计与工程规则.md")
    adr1 = read("docs/deliverables/16-ADR-001-应用拆分与部署单元.md")
    adr2 = read("docs/deliverables/17-ADR-002-事务Outbox与持久任务.md")
    openapi_path = ROOT / "docs/deliverables/15-接口契约-openapi_v1.yaml"
    openapi_text = openapi_path.read_text(encoding="utf-8")

    task_states = {task: task_status(tasks, task) for task in [f"G3-{n:02}" for n in range(2, 15)]}
    upstream_done = all("DONE" in task_states[f"G3-{n:02}"] for n in range(2, 10))
    management_plans_started = all(
        any(token in task_states[f"G3-{n:02}"] for token in ("REVIEW", "DONE"))
        for n in range(11, 15)
    )

    counts = {
        "fr": len(set(re.findall(r"G2-FR-(\d{3})", spec))),
        "ac": len(set(re.findall(r"AC-G2-FR-\d{3}-\d{2}", spec))),
        "rclr": len(set(re.findall(r"G2-RCLR-(\d{3})", spec))),
        "pe": len(set(re.findall(r"PE-(\d{2})", nfr))),
        "eng": len(set(re.findall(r"ENG-(\d{3})", nfr))),
        "star_fr": len(set(re.findall(r"G2-FR-(\d{3})\s*/\s*★", rtm))),
        "tc": 3 * len(set(re.findall(r"TC-G2-FR-(\d{3})-01—03", rtm))),
        "dbd_tr": len(set(re.findall(r"DBD-TR-(\d{3})", rtm))),
        "dld_tr": len(set(re.findall(r"DLD-TR-(\d{3})", rtm))),
        "api_tr": len(set(re.findall(r"API-TR-(\d{3})", rtm))),
    }
    operation_ids = re.findall(r"^\s*operationId:\s*['\"]?([^'\"\s]+)", openapi_text, re.MULTILINE)
    counts["openapi_operations"] = len(operation_ids)
    counts["openapi_unique_operations"] = len(set(operation_ids))

    count_expectations = {
        "fr": 39,
        "ac": 117,
        "rclr": 10,
        "pe": 12,
        "eng": 20,
        "star_fr": 34,
        "tc": 117,
        "dbd_tr": 39,
        "dld_tr": 39,
        "api_tr": 39,
        "openapi_operations": 57,
        "openapi_unique_operations": 57,
    }
    count_checks = {key: counts[key] == value for key, value in count_expectations.items()}

    candidate_files = [
        "docs/work/A_PM/plan.md",
        "constitution.md",
        "docs/work/A_PM/overview_design.md",
        "docs/work/B_TECH/database_design.md",
        "docs/work/B_TECH/detailed_design.md",
        "docs/work/B_TECH/interface_design.md",
        "docs/work/C_REQ/rtm_v1.md",
        "docs/work/C_REQ/test_plan.md",
        "docs/deliverables/10-需求追踪矩阵RTMv1.docx",
        "docs/deliverables/11-概要设计说明书.docx",
        "docs/deliverables/12-详细设计说明书.docx",
        "docs/deliverables/13-数据库设计说明书.docx",
        "docs/deliverables/14-接口设计说明书.docx",
        "docs/deliverables/15-接口契约-openapi_v1.yaml",
        "docs/deliverables/16-ADR-001-应用拆分与部署单元.md",
        "docs/deliverables/17-ADR-002-事务Outbox与持久任务.md",
        "docs/deliverables/18-非功能设计与工程规则.md",
        "docs/deliverables/19-测试计划.docx",
    ]
    manifest = []
    for relative in candidate_files:
        path = ROOT / relative
        manifest.append(
            {
                "path": relative,
                "exists": path.exists(),
                "size": path.stat().st_size if path.exists() else None,
                "sha256": sha256(path) if path.exists() else None,
            }
        )

    review_patterns = {
        "G3-02": "logs/reviews/*G3-02*review*.md",
        "G3-03": "logs/reviews/*G3-03*review*.md",
        "G3-04": "logs/reviews/*G3-04*review*.md",
        "G3-05": "logs/reviews/*G3-05*review*.md",
        "G3-06": "logs/reviews/*G3-06*review*.md",
        "G3-07": "logs/reviews/*G3-07*review*.md",
        "G3-08": "logs/reviews/*G3-08*review*.md",
        "G3-09": "logs/reviews/*G3-09*review*.md",
    }
    review_evidence = {
        task: sorted(str(path.relative_to(ROOT)).replace("\\", "/") for path in ROOT.glob(pattern))
        for task, pattern in review_patterns.items()
    }

    formal_paths = [ROOT / item for item in candidate_files if item.endswith(".docx")]
    forbidden_terms = ("澜图", "遥感影像智能解译", "通关实训第3组", "PostGIS", "MinIO", "Vue3")
    case_residue = {}
    for path in formal_paths:
        text = docx_text(path)
        hits = [term for term in forbidden_terms if term in text]
        if hits:
            case_residue[str(path.relative_to(ROOT)).replace("\\", "/")] = hits

    blocker_status = last_issue_status(issues, "ISSUE-G3-01-001")
    adr_status = {
        "ADR-001": bool(re.search(r"状态[：:]\s*Accepted", adr1, re.IGNORECASE)),
        "ADR-002": bool(re.search(r"状态[：:]\s*Accepted", adr2, re.IGNORECASE)),
    }
    readme = read("docs/deliverables/README.md")
    readme_missing = [Path(item).name for item in candidate_files if item.startswith("docs/deliverables/") and Path(item).name not in readme]

    consistency_pass = all(count_checks.values()) and all(item["exists"] for item in manifest)
    consistency_pass = consistency_pass and all(review_evidence.values()) and not case_residue and all(adr_status.values())
    blockers = []
    if not blocker_status.startswith("CLOSED"):
        blockers.append("ISSUE-G3-01-001 remains OPEN and explicitly blocks G3-10/M3 freeze")
    if not management_plans_started:
        blockers.append("plan.md requires G3-02—G3-14 REVIEW, but G3-11—G3-14 remain TODO")

    result = {
        "task": "G3-10",
        "audit_date": "2026-09-21",
        "task_states": task_states,
        "upstream_g3_02_to_09_done": upstream_done,
        "management_plans_g3_11_to_14_review_or_done": management_plans_started,
        "issue_g3_01_001_status": blocker_status,
        "counts": counts,
        "count_expectations": count_expectations,
        "count_checks": count_checks,
        "adr_status": adr_status,
        "manifest": manifest,
        "review_evidence": review_evidence,
        "case_residue": case_residue,
        "deliverables_readme_missing": readme_missing,
        "consistency_checks_pass": consistency_pass,
        "freeze_ready": consistency_pass and upstream_done and not blockers,
        "decision": "FREEZE" if consistency_pass and upstream_done and not blockers else "DO_NOT_FREEZE",
        "blockers": blockers,
        "observations": (
            (["docs/deliverables/README.md does not yet inventory all current G3 deliverables"] if readme_missing else [])
            + ["G3-06 A review contains a non-blocking column-width transcription observation recorded by C"]
        ),
    }

    payload = json.dumps(result, ensure_ascii=False, indent=2)
    print(payload)
    if args.out:
        out = args.out if args.out.is_absolute() else ROOT / args.out
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(payload + "\n", encoding="utf-8")

    if not consistency_pass:
        raise SystemExit(2)


if __name__ == "__main__":
    main()
