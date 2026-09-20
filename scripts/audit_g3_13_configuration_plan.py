from __future__ import annotations

import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "docs/work/A_PM/configuration_management_plan.md"
DOCX = ROOT / "docs/deliverables/22-配置管理计划.docx"
TASKS = ROOT / "tasks.md"
ALLOWED = {"TODO", "DOING", "REVIEW", "DONE", "BLOCKED"}


def main() -> None:
    problems = []
    work = WORK.read_text(encoding="utf-8")
    tasks = TASKS.read_text(encoding="utf-8")
    doc = Document(DOCX)
    visible = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)

    required = [
        "配置项识别、登记与状态", "版本、分支、提交与制品规则", "基线建立、冻结与解除",
        "变更、CR 与 CCB 流程", "权限、安全与配置审计", "发布、部署与回滚",
        "配置库、备份与归档", "AC-G2-FR-013-01 配置抽查",
        "ISSUE-G3-01-001", "ISSUE-G3-10-001", "39 条 FR", "34 条★", "117 条 AC",
    ]
    for marker in required:
        if marker not in work or marker not in visible:
            problems.append({"code": "MISSING_MARKER", "detail": marker})
    for pollution in ("澜图", "遥感影像", "Alembic", "通关实训第 3 组", "openapi_lantu"):
        if pollution in visible:
            problems.append({"code": "TEACHING_POLLUTION", "detail": pollution})

    status_rows = []
    for line in tasks.splitlines():
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if len(cells) == 8 and cells[0] not in {"Task ID", "---"}:
            status_rows.append((cells[0], cells[-1]))
            if cells[-1] not in ALLOWED:
                problems.append({"code": "TASK_STATUS", "detail": [cells[0], cells[-1]]})
    g313 = dict(status_rows).get("G3-13")
    if g313 != "REVIEW":
        problems.append({"code": "G3_13_STATUS", "detail": g313})

    heading1 = sum(1 for p in doc.paragraphs if p.text.strip() and p.style and p.style.name == "Heading 1")
    heading2 = sum(1 for p in doc.paragraphs if p.text.strip() and p.style and p.style.name == "Heading 2")
    if (heading1, heading2) != (11, 18):
        problems.append({"code": "HEADINGS", "detail": [heading1, heading2]})
    if len(doc.tables) != 10:
        problems.append({"code": "TABLES", "detail": len(doc.tables)})

    expected_captions = [
        "表 1-1 配置管理依据及用途", "表 2-1 配置管理角色职责与权限边界",
        "表 3-1 配置项类别与最低完整性检查", "表 4-1 配置对象版本标识规则",
        "表 5-1 基线类型、建立门禁与变更权限", "表 6-1 CR 与 CCB 处理流程",
        "表 7-1 配置审计类型、时点与证据", "表 9-1 配置库逻辑区域与控制要求",
        "表 10-1 AC-G2-FR-013-01 配置证据抽查",
    ]
    actual_captions = [p.text.strip() for p in doc.paragraphs if re.match(r"^表\s+\d+-\d+\s+", p.text.strip())]
    if actual_captions != expected_captions:
        problems.append({"code": "TABLE_CAPTIONS", "detail": {"expected": expected_captions, "actual": actual_captions}})
    work_captions = [line.strip() for line in work.splitlines() if re.match(r"^表\s+\d+-\d+\s+", line.strip())]
    if work_captions != expected_captions:
        problems.append({"code": "WORK_TABLE_CAPTIONS", "detail": {"expected": expected_captions, "actual": work_captions}})
    body_children = list(doc.element.body)
    adjacent_caption_count = 0
    for idx, child in enumerate(body_children[:-1]):
        if child.tag != qn("w:p"):
            continue
        text = "".join(t.text or "" for t in child.xpath(".//w:t")).strip()
        if text in expected_captions and body_children[idx + 1].tag == qn("w:tbl"):
            adjacent_caption_count += 1
    if adjacent_caption_count != len(expected_captions):
        problems.append({"code": "TABLE_CAPTION_ADJACENCY", "detail": adjacent_caption_count})

    with ZipFile(DOCX) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
        comments = "word/comments.xml" in zf.namelist()
    if comments:
        problems.append({"code": "COMMENTS_PRESENT"})
    if re.search(r"<w:(?:ins|del)(?:\s|>)", xml):
        problems.append({"code": "TRACKED_CHANGES_PRESENT"})
    toc_pages = dict(re.findall(r"(?:^|>)(\d{1,2} [^<]+)</w:t>.*?<w:tab/>.*?<w:t>(\d+)</w:t>", xml, re.S))

    result = {
        "task": "G3-13",
        "pass": not problems,
        "problems": problems,
        "counts": {"heading1": heading1, "heading2": heading2, "tables": len(doc.tables), "business_table_captions": len(actual_captions), "adjacent_caption_table_pairs": adjacent_caption_count, "task_rows": len(status_rows)},
        "task_status": g313,
        "toc_last": {k: v for k, v in toc_pages.items() if k.startswith(("9 ", "10 ", "11 "))},
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["pass"] else 1)


if __name__ == "__main__":
    main()
