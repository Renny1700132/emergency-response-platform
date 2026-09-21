from __future__ import annotations
import json
import re
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
WORK=ROOT/"docs/work/A_PM/risk_management_plan.md"
DOCX=ROOT/"docs/deliverables/23-风险管理计划与风险登记册v2.docx"
TASKS=ROOT/"tasks.md"

def main() -> None:
    problems=[]; work=WORK.read_text(encoding="utf-8"); doc=Document(DOCX)
    expected=["表 2-1 风险治理角色职责与边界","表 3-1 风险概率与影响评分定义","表 4-1 当前极高与高风险登记","表 4-2 当前中风险登记","表 6-1 AC-G2-FR-006-01 风险控制核对","表 7-1 风险与 Issue、变更、测试及配置闭环","表 9-1 G3-14 主责整改自检"]
    work_caps=[x.strip() for x in work.splitlines() if re.match(r"^表\s+\d+-\d+\s+",x.strip())]
    doc_caps=[p.text.strip() for p in doc.paragraphs if re.match(r"^表\s+\d+-\d+\s+",p.text.strip())]
    if work_caps!=expected: problems.append({"code":"WORK_CAPTIONS","actual":work_caps})
    if doc_caps!=expected: problems.append({"code":"DOCX_CAPTIONS","actual":doc_caps})
    children=list(doc.element.body); adjacent=0
    for i,ch in enumerate(children[:-1]):
        if ch.tag==qn("w:p"):
            text="".join(t.text or "" for t in ch.xpath(".//w:t")).strip()
            if text in expected and children[i+1].tag==qn("w:tbl"): adjacent+=1
    if adjacent!=7: problems.append({"code":"CAPTION_ADJACENCY","actual":adjacent})
    rows={}
    for line in work.splitlines():
        if line.startswith("| RR-"):
            cells=[c.strip() for c in line.strip().strip("|").split("|")]
            rid=cells[0].split()[0]; score=cells[2] if len(cells)>=7 else cells[2]
            rows[rid]=score
    expected_ids={f"RR-{i:02d}" for i in range(1,17)}
    if set(rows)!=expected_ids: problems.append({"code":"RISK_IDS","missing":sorted(expected_ids-set(rows)),"extra":sorted(set(rows)-expected_ids)})
    for rid,field in rows.items():
        m=re.search(r"(\d)/(\d)/(\d+)\s*(极高|高|中|低)",field)
        if not m: problems.append({"code":"RISK_SCORE_PARSE","id":rid,"field":field}); continue
        p,i,score,level=int(m.group(1)),int(m.group(2)),int(m.group(3)),m.group(4)
        expected_level="低" if score<=4 else "中" if score<=9 else "高" if score<=16 else "极高"
        if p*i!=score or level!=expected_level: problems.append({"code":"RISK_SCORE_LEVEL","id":rid,"field":field,"expected":f"{p*i} {expected_level}"})
    toc_texts=[p.text.strip() for p in doc.paragraphs]
    toc_expected={"6 AC-G2-FR-006-01 专项核对\t5","9 主责自检与准出\t6"}
    if not toc_expected.issubset(set(toc_texts)):
        problems.append({"code":"TOC_FINAL_PAGES","expected":sorted(toc_expected)})
    task_line=next((line for line in TASKS.read_text(encoding="utf-8").splitlines() if line.startswith("| G3-14 |")),"")
    task_status=task_line.rstrip().rstrip("|").split("|")[-1].strip() if task_line else "MISSING"
    result={"task":"G3-14","pass":not problems,"problems":problems,"counts":{"tables_total":len(doc.tables),"business_tables":len(doc.tables)-1,"captions":len(doc_caps),"adjacent_caption_table_pairs":adjacent,"risk_ids":len(rows),"toc_final_entries":2},"task_status":task_status}
    print(json.dumps(result,ensure_ascii=False,indent=2)); raise SystemExit(0 if not problems else 1)

if __name__=="__main__": main()
