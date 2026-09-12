from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "reference" / "02-澄清与质询记录（教学样例）.docx"
OUTPUT = ROOT / "docs" / "deliverables" / "02-澄清与质询记录.docx"

PROJECT = "某自然博物馆智能运营中心建设项目——应急管理子系统"

CLARIFICATIONS = [
    ("CLR-001", "历史录像回放及不少于30天保存责任如何划分？", "历史录像回放按强制要求处理；既有视频系统负责录像存储及不少于30天保存。本项目负责 GB/T 28181 接入、实时调阅、历史回放和验证；既有系统不满足保存条件时，不转化为本项目新建录像存储责任。", "方案保留视频接入、调阅、回放和验证；不将既有录像存储或改造扩入范围。"),
    ("CLR-002", "亚米级定位与≤2秒刷新由谁提供和保障？", "甲方提供既有定位基础设施、接口及验收环境中的亚米级源数据；乙方负责接口适配、人员关联、楼层/坐标映射和≤2秒刷新展示，且不得降低源精度。", "方案以既有定位源为输入，增加源/展示时间戳、坐标映射和刷新验证设计。"),
    ("CLR-003", "疏散门禁联动如何满足安全控制要求？", "疏散门禁联动按强制“应支持”处理；仅由授权人员确认后下发，服从既有门禁安全联锁；完整记录操作人、时间和结果；失败时告警并人工降级，不绕过安全机制。", "方案采用授权确认、幂等下发、回执审计、失败告警与人工处置衔接。"),
    ("CLR-004", "消息通道、并行发送和到达率的责任如何划分？", "甲方提供既有统一消息/APP/SMS通道、账号、签名、配额及联调条件；乙方负责业务编排、≥20路并行、重试、回执、留痕和降级；≥99%到达率在正常甲方通道下验证。", "方案保留消息编排、重试、回执和留痕；验收以正常既有通道为前提。"),
    ("CLR-005", "统一中台能力与本项目建设边界是什么？", "甲方提供统一身份、组织/用户、权限、消息、工作流、文件、门户和数据汇聚等通用能力；乙方只做应急业务适配，不重复建设中台通用服务。", "方案仅建设应急业务适配、映射、容错和降级，不重复建设通用中台。"),
    ("CLR-006", "地图资料、三维场景与测绘建模的责任如何划分？", "甲方提供合法可用的二维/楼层地图及验收需要的既有三维地图、场景、地图服务和楼层/坐标关系；乙方负责加载、展示、标绘、拾取和业务叠加。新增测绘、从零三维建模及底图版权采购不在范围内。", "方案落实地图加载、标绘、拾取和业务图层；不承诺测绘、建模或版权采购。"),
    ("CLR-007", "移动端采用何种交付形态，发布责任归谁？", "移动端冻结为 H5 嵌入甲方现有智慧管理 APP；不单独建设原生 Android/iOS APP。乙方交付 H5、接口和集成说明，甲方负责宿主 APP、签名和发布。", "方案按 Web+H5 实施，验收覆盖宿主内登录、消息、定位、扫码、拍照和上传。"),
    ("CLR-008", "离线部署的基础环境和软件部署责任如何划分？", "甲方提供机房、服务器、网络、域名/证书、备份介质及后备电源等基础设施；乙方提供最低配置、端口/网络要求、容器化部署、初始化和恢复方案。不得据此虚构未确认的具体硬件配置。", "方案提供环境前置清单、容器化部署和恢复设计；不虚构硬件规格。"),
    ("CLR-009", "初始化数据的真实性、处理和确认责任如何划分？", "甲方提供真实业务源数据并负责业务正确性确认；乙方负责模板、映射、清洗、导入、技术校验和示例初始化；不得自行编造缺失业务数据，最终初始化结果双方确认。", "方案增加模板、映射、清洗日志、导入校验和双方确认机制。"),
    ("CLR-010", "既有外部系统接口资料、环境和联调责任如何划分？", "甲方负责协调视频、信息发布、入侵、门禁、消防、物联网、中台、消息等既有系统，提供文档、版本、账号、权限、测试环境和协调窗口；乙方负责适配、异常处理、联调和验证。视频与消息须在 M2 前完成真实连通验证。", "计划将外部前置纳入接口台账；视频、消息在 M2 前形成连通验证证据。"),
    ("CLR-011", "等保相关应用控制、测试和第三方测评责任如何划分？", "本项目按不低于等保二级相关应用控制要求设计；乙方负责应用侧安全控制、SCA、漏洞扫描和 Web 渗透复测，高危漏洞清零。正式定级备案、基础设施合规及正式第三方等保测评由甲方组织，第三方测评采购默认不在本项目范围。", "方案保留应用安全控制和复测证据；不将定级、基础设施合规或第三方采购扩入范围。"),
    ("CLR-012", "项目和招标人应使用何种正式称谓？", f"统一采用匿名名称：“{PROJECT}”；招标人/甲方统一写“某自然博物馆”，不得虚构真实单位名称。", "封面、正文和交付物统一采用匿名名称，删除或避免任何现实主体推断。"),
]

QNAS = [
    ("Q-01", "为什么不新建视频、定位、门禁等系统，如何保证联动成功？", "这些均为甲方既有系统能力或基础设施。本项目承担接口适配、业务编排、授权控制、异常处置和联调验证，不替代既有系统建设。以接口契约、账号权限、测试环境和联调窗口为前置，按逐接口正常、异常、回执和人工降级用例形成证据。", "Baseline CLR-001/002/003/010；G1-04 技术方案接口与联动设计；G1-07 外部依赖风险。"),
    ("Q-02", "预案启动≤3秒、消息≥20路且到达率≥99%如何验证？", "预案启动以服务端关键路径、状态写入、审计和结果回执为对象进行性能测试；消息以甲方提供的正常验收通道为前提，执行不少于20路并行发送、回执关联、失败重试和到达率统计。测试数据、时间戳、通道回执和异常记录形成验证证据，不能以模拟口头结论替代。", "Baseline 关键数字与 CLR-004/010；G1-04 性能与消息设计；G1-06 M2 联调安排。"),
    ("Q-03", "亚米级定位并非乙方建设，如何证明系统没有降低精度？", "甲方提供验收环境中的亚米级源数据，乙方不对源定位设施作替代承诺。验证时比对源数据与展示数据的人员关联、楼层/坐标映射、时间戳和刷新间隔，确认展示链路≤2秒且无不合理精度截断、坐标偏移或错误转换。", "Baseline CLR-002；G1-04 定位与地图设计；G1-07 数据质量与外部依赖风险。"),
    ("Q-04", "视频回放与不少于30天保存的责任如何划分？", "历史回放为强制响应；既有视频系统负责录像存储及不少于30天保存，本项目负责 GB/T 28181 接入、实时调阅、历史检索、回放和验证。验收中验证目录、鉴权、检索和跨30天回放能力；既有系统不满足保存前提不构成本项目新建存储责任。", "Baseline CLR-001；G1-04 视频接入与验证设计；G1-06 接口联调计划。"),
    ("Q-05", "门禁远程开启如何避免形成安全风险？", "系统只在授权人员确认后请求既有门禁执行，服从既有门禁安全联锁，不绕过现场安全机制。全过程记录操作人、时间、指令和结果；超时或失败立即告警，并转人工处置。验证覆盖权限、二次确认、幂等、回执、联锁和失败降级。", "Baseline CLR-003；G1-04 门禁联动与安全设计；G1-07 联动风险。"),
    ("Q-06", "外部系统或消息通道故障时，应急处置是否还能继续？", "核心事件、预案、任务和处置状态在本项目边界内保持可追踪；外部调用采用超时、重试、幂等、回执和失败告警。消息或外部联动不可用时，系统保留待处理状态、提示人工联系或现场处置，不伪造已送达或已执行结果；恢复后按记录进行补偿或复核。", "Baseline CLR-003/004/005/010；G1-04 异常与降级设计；G1-07 风险登记。"),
    ("Q-07", "H5 而非原生 APP 是否满足 Android/iOS 和现场处置需求？", "冻结交付形态是嵌入甲方现有智慧管理 APP 的 H5，不另建原生 APP。乙方交付 H5、业务接口和集成说明，甲方负责宿主、签名和发布；以甲方提供的宿主版本和终端矩阵验证登录、消息、定位、扫码、拍照、上传等现场流程。", "Baseline CLR-007；G1-04 H5 方案；G1-06 联调与试运行计划。"),
    ("Q-08", "AI 辅助生成代码和文档，如何防止幻觉、错误承诺和质量问题？", "课程项目中的 AI 辅助仅用于提高分析和编制效率，不替代需求、范围、关键数字和验收结论的人工核验。方案以冻结 Baseline 为约束，关键内容需回指需求、澄清或控制文件；未确认事实保留待确认标识，实施阶段仍以代码审查、测试、扫描和验收证据判定质量。", "Baseline 的变更与证据规则；G1-04 安全/验证设计；G1-06 质量活动；G1-07 风险控制。"),
]

def set_cell_shading(cell, fill="D9E2F3"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), "7F8FA4")
        borders.append(item)
    tbl_pr.append(borders)

def set_cell_text(cell, text, bold=False, size=9.0, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

def set_column_widths(table, widths_cm):
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)

def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    return p

def add_body(doc, text, indent=True):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(10.5)
    return p

def add_cover_line(doc, text, before=0, after=0, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "宋体"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r.font.size = Pt(size)
    return p

def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def set_border_bottom(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '000000')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_flow(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("澄清与质询闭环：事实可追溯、责任可核验")
    r.bold = True
    r.font.name = "宋体"; r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体"); r.font.size = Pt(10.5)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Normal Table"
    set_table_borders(table)
    labels = ["识别需求歧义", "课程模拟书面裁决", "冻结 Baseline 解释", "方案/计划/风险落实", "实施验证与归档"]
    for i, label in enumerate(labels):
        set_cell_text(table.cell(0, i), label, bold=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.cell(0, i), "D9E2F3" if i % 2 == 0 else "F2E9D4")
    set_column_widths(table, [3.05] * 5)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = note.add_run("说明：本流程为补录文件的证据闭环，不表示存在真实现场会议、签署或甲方公文。")
    rr.font.name = "宋体"; rr._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体"); rr.font.size = Pt(8.5); rr.font.color.rgb = RGBColor(166, 87, 0)
    add_caption(doc, "图 1-1  澄清与质询闭环流程（课程模拟/补录）")

def remove_body_content(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

def main():
    doc = Document(TEMPLATE)
    remove_body_content(doc)
    section = doc.sections[0]
    section.top_margin = Cm(2.54); section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54); section.right_margin = Cm(2.54)

    # Cover, preserving template page settings, headers and footer/page numbering.
    add_cover_line(doc, "文档编号：G1-05-BL-001        版本号：V1.0", before=20, size=11)
    add_cover_line(doc, "密    级：课程模拟 · 补录", before=32, size=11, align=WD_ALIGN_PARAGRAPH.RIGHT)
    title = add_cover_line(doc, PROJECT, before=105, after=8, size=21, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_border_bottom(title)
    add_cover_line(doc, "澄清与质询记录", before=36, after=12, size=25, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_cover_line(doc, "（课程模拟/补录）", size=13, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_cover_line(doc, "编制单位：项目组【待人工确认】", before=120, size=12)
    add_cover_line(doc, "编    制：A（补录）", before=10, size=12)
    add_cover_line(doc, "审    核：B（技术确认/复核）【待人工确认】", before=10, size=12)
    add_cover_line(doc, "核    验：C（合规核验）【待人工确认】", before=10, size=12)
    add_cover_line(doc, "补录日期：2026 年 9 月 12 日", before=10, size=12)
    add_page_break(doc)

    add_caption(doc, "修订记录")
    t = doc.add_table(rows=2, cols=5); t.style = "Normal Table"; set_table_borders(t)
    headers = ["版本", "日期", "修订章节", "修订说明", "编制/修订人"]
    for i, h in enumerate(headers):
        set_cell_text(t.cell(0, i), h, bold=True, size=10.0, align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_shading(t.cell(0, i))
    row = ["V1.0", "2026-09-12", "全文", "【补录】依据已归档的课程模拟书面裁决及冻结 Baseline，补齐缺失的正式候选交付物；不回溯虚构会议、签署或现场发言。", "A（补录）"]
    for i, val in enumerate(row): set_cell_text(t.cell(1, i), val, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER if i in (0,1,4) else WD_ALIGN_PARAGRAPH.LEFT)
    set_column_widths(t, [2.2, 2.5, 2.5, 6.9, 2.4])
    add_body(doc, "注：本文件补录的是课程项目的澄清与质询候选记录。表内“甲方/招标人（课程模拟）书面裁决”是当前 Baseline 的解释依据，不冒充现实单位签章、正式函件或法律文件。", indent=False)
    add_page_break(doc)

    doc.add_paragraph("1 说明", style="Heading 1")
    add_body(doc, "本记录包括两部分：第一部分为课程模拟甲方/招标人书面裁决的补录归档；第二部分为依据已冻结方案形成的评委会质询问答（课程模拟/补录）。其用途是明确课程项目投标响应边界、责任分工和后续验证前提，供编制、复核与归档使用。")
    add_body(doc, "补录依据为已归档的课程模拟澄清记录、问题台账、冻结 Baseline、变更记录及 G1-04、G1-06、G1-07、技术标等工作成果。历史材料未提供完整的真实现场质询纪要，因此第3章不作为现实会议记录，也不附会议信息、真实人员、签字或现场发言。", indent=False)
    add_flow(doc)

    doc.add_paragraph("2 书面澄清记录", style="Heading 1")
    add_body(doc, "以下12项均为【课程模拟甲方/招标人书面裁决】，用于解决课程项目需求歧义并构成当前 Baseline 的解释依据。", indent=False)
    add_caption(doc, "表 2-1  书面澄清记录")
    t = doc.add_table(rows=1, cols=4); t.style = "Normal Table"; set_table_borders(t)
    heads = ["编号", "澄清问题", "甲方/招标人（课程模拟）书面裁决", "对投标方案的影响"]
    for i, h in enumerate(heads):
        set_cell_text(t.cell(0, i), h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_shading(t.cell(0, i))
    set_repeat_table_header(t.rows[0])
    for num, question, decision, impact in CLARIFICATIONS:
        cells = t.add_row().cells
        for i, value in enumerate((num, question, decision, impact)):
            set_cell_text(cells[i], value, size=8.3, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_column_widths(t, [1.7, 3.4, 7.0, 4.4])

    doc.add_paragraph("3 评委会质询问答（课程模拟/补录）", style="Heading 1")
    add_body(doc, "本章系基于冻结 Baseline 和现有方案的课程模拟质询，不是现实评标现场的原始问答。回答仅说明设计边界、验证路径和证据要求；未形成的接口、环境、性能或现场证据仍须在实施阶段取得。", indent=False)
    add_caption(doc, "表 3-1  评委会质询问答（课程模拟/补录）")
    t = doc.add_table(rows=1, cols=4); t.style = "Normal Table"; set_table_borders(t)
    heads = ["编号", "质询", "答复", "依据"]
    for i, h in enumerate(heads):
        set_cell_text(t.cell(0, i), h, bold=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER); set_cell_shading(t.cell(0, i))
    set_repeat_table_header(t.rows[0])
    for num, question, answer, evidence in QNAS:
        cells = t.add_row().cells
        for i, value in enumerate((num, question, answer, evidence)):
            set_cell_text(cells[i], value, size=8.3, align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_column_widths(t, [1.5, 4.0, 7.0, 4.0])
    add_body(doc, "记录人：＿＿＿＿＿＿（【待人工确认】）", indent=False)
    add_body(doc, "投标人代表确认：＿＿＿＿＿＿（【待人工确认】）", indent=False)
    add_body(doc, "评委会代表确认：＿＿＿＿＿＿（【待人工确认】）", indent=False)
    add_body(doc, "日期：＿＿＿＿＿＿（如有正式确认文件后补充；本补录不虚构日期）", indent=False)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)

if __name__ == "__main__":
    main()
