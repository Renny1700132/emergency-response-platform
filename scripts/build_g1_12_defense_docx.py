"""Build the G1-12 candidate Word document from the controlled Markdown draft."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".tmp" / "g1_11_python_libs"))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REFERENCE = ROOT / "docs" / "reference" / "06-述标答辩讲稿与策略（教学样例）.docx"
SOURCE = ROOT / "docs" / "work" / "A_PM" / "presentation_and_defense.md"
TARGET = ROOT / "docs" / "deliverables" / "06-述标答辩讲稿与策略.docx"


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag.endswith("sectPr"):
            continue
        body.remove(child)


def add_cover(doc):
    p = doc.add_paragraph("文档编号：03-G1-12　　版本号：V0.1")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p = doc.add_paragraph("密　　级：课程模拟使用")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph()
    for text, size, bold in [
        ("某自然博物馆智能运营中心建设项目", 22, True),
        ("——应急管理子系统", 22, True),
        ("述标答辩讲稿与模拟质询策略", 20, True),
        ("（课程模拟）", 14, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.font.size = Pt(size)
    for text in [
        "编制单位：项目组【待人工确认】", "主 述 标 人：A（项目经理 PM）", "审　　核：B（技术口径）、C（合规/数字）", "编制日期：2026 年 9 月 12 日",
    ]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def table_from_lines(doc, lines):
    rows = [[part.strip() for part in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(set(cell) <= {"-", ":"} for cell in row)]
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Normal Table"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        cell.text = value
    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = value


def build_content(doc):
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("## 第一部分"))
    lines = lines[start:]
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            table_from_lines(doc, table_lines)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Paragraph")
        else:
            doc.add_paragraph(line)
        i += 1


doc = Document(REFERENCE)
clear_body(doc)
add_cover(doc)
doc.add_heading("修订记录", level=1)
table = doc.add_table(rows=1, cols=5)
table.style = "Normal Table"
for cell, value in zip(table.rows[0].cells, ["版本", "日期", "修订章节", "修订说明", "编制/修订人"]):
    cell.text = value
for row in [("V0.1", "2026-09-12", "全文", "首次编制；依据冻结基线、技术标和已通过复核成果形成课程模拟述标稿。", "A")]:
    cells = table.add_row().cells
    for cell, value in zip(cells, row):
        cell.text = value
doc.add_paragraph("注：本文件用于课程模拟述标与模拟质询，现场事实、接口条件和验收证据按项目里程碑另行形成。")
build_content(doc)
TARGET.parent.mkdir(parents=True, exist_ok=True)
doc.save(TARGET)
print(TARGET)
