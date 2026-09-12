"""Generate G1-13 from the teaching sample's layout shell, not its business content."""
import sys, re
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'.tmp'/'g1_11_python_libs'))
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC=ROOT/'docs/work/A_PM/requirements_confirmation.md'
REF=ROOT/'docs/reference/05-需求确认书（教学样例）.docx'
OUT=ROOT/'docs/deliverables/05-需求确认书.docx'
BLUE='1F4E78'; DARK='1F1F1F'
def font(run,size=10.5,bold=False):
    run.font.name='宋体';run._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');run.font.size=Pt(size);run.bold=bold;run.font.color.rgb=RGBColor.from_string(DARK)
def shade(cell,color):
    x=OxmlElement('w:shd');x.set(qn('w:fill'),color);cell._tc.get_or_add_tcPr().append(x)
def cell(c,text,bold=False,size=8):
    c.text='';p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
    r=p.add_run(text);font(r,size,bold);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def borders(t):
    e=OxmlElement('w:tblBorders')
    for k in ['top','left','bottom','right','insideH','insideV']:
        z=OxmlElement('w:'+k);z.set(qn('w:val'),'single');z.set(qn('w:sz'),'4');z.set(qn('w:color'),'A6A6A6');e.append(z)
    t._tbl.tblPr.append(e)
def add_table(d,rows,fs=7.5):
    t=d.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;borders(t)
    for ri,row in enumerate(rows):
        cs=t.rows[0].cells if ri==0 else t.add_row().cells
        for i,x in enumerate(row):
            cell(cs[i],x,ri==0,fs)
            if ri==0:shade(cs[i],BLUE);cs[i].paragraphs[0].runs[0].font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def p(d,text='',size=10.5,align=None,bold=False,after=5):
    q=d.add_paragraph();q.paragraph_format.space_after=Pt(after);q.paragraph_format.line_spacing=1.22
    if align is not None:q.alignment=align
    font(q.add_run(text),size,bold);return q
def heading(d,text,level):
    q=d.add_heading(text,level);q.paragraph_format.space_before=Pt(8);q.paragraph_format.space_after=Pt(5)
    for r in q.runs:
        r.font.name='黑体';r._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体');r.font.size=Pt(14 if level==1 else 11);r.font.color.rgb=RGBColor.from_string(BLUE)
    return q
def parse_table(lines):
    group=[]
    while lines and lines[0].startswith('|'):group.append(lines.pop(0).strip())
    return [[v.strip() for v in x.strip('|').split('|')] for x in group if not re.match(r'^\|[ -|]+\|$',x)]
d=Document(REF);body=d._element.body
for x in list(body):
    if not x.tag.endswith('sectPr'):body.remove(x)
s=d.sections[0];s.top_margin=Cm(2.25);s.bottom_margin=Cm(2.1);s.left_margin=Cm(2.25);s.right_margin=Cm(2.25)
# cover from the sample's design vocabulary
p(d,'文档编号：【待人工确认】    版本号：V0.1',9.5,after=2);p(d,'密    级：课程模拟使用',9.5,after=36)
p(d,'某自然博物馆智能运营中心建设项目',18,WD_ALIGN_PARAGRAPH.CENTER,True,6);p(d,'——应急管理子系统',18,WD_ALIGN_PARAGRAPH.CENTER,True,26)
p(d,'需求确认书',26,WD_ALIGN_PARAGRAPH.CENTER,True,5);p(d,'（课程模拟候选版）',13,WD_ALIGN_PARAGRAPH.CENTER,False,38)
for x in ['编制单位：【项目组/编制单位正式名称或组号待人工确认】','编    制：A（项目经理 PM / 总编）【真实姓名待人工确认】','技术审核：B（技术负责人 / 架构师）【真实姓名待人工确认】','合规审核：C（需求 / 质量 / 合规负责人）【真实姓名待人工确认】','批    准：甲乙双方授权代表确认','编制日期：【正式签署日期待人工确认】']:p(d,x,11,WD_ALIGN_PARAGRAPH.CENTER,False,7)
d.add_page_break();p(d,'修订记录',16,WD_ALIGN_PARAGRAPH.CENTER,True,8)
add_table(d,[['版本','日期','修订章节','修订说明','编制/修订人'],['V0.1','【待人工确认】','全部','课程模拟候选版：确认MVP优先顺序、全项目边界和验收基线，待复核。','A（项目经理 PM / 总编）']],8.5)
p(d,'注：MVP仅定义优先形成可验证闭环的顺序，不删除《用户需求书》的任何功能或验收要求。',8.5,after=9)
lines=SRC.read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    raw=lines[i].strip();i+=1
    if not raw or raw.startswith('# ') or raw.startswith('## 文档控制信息') or raw.startswith('- 项目名称') or raw.startswith('- 文档性质') or raw.startswith('- 主编') or raw.startswith('- 对外正式稿') or raw.startswith('内部追溯：'):continue
    if raw.startswith('## '):
        if raw.startswith('## 2 '):
            heading(d,raw[3:],1);p(d,'图 2-1  MVP范围与边界一览',9,WD_ALIGN_PARAGRAPH.CENTER,True,3)
            fig=add_table(d,[['优先形成可验证闭环','接口与质量前置验证','非MVP但仍属本期'],['预案与事件处置\n指挥态势与现场执行','视频、定位、消息、物联、中台\n性能、安全、部署','其余功能、数据初始化、文档培训\n全量39条FR与PE验收']],8.5)
            for c in fig.rows[1].cells:shade(c,'EAF2F8')
            p(d,'说明：五项MVP定义优先闭环；非MVP功能仍按本期项目计划实施与验收。',8.2,after=5)
        elif raw.startswith('## 4 '):d.add_page_break();heading(d,raw[3:],1)
        elif raw.startswith('## 5 '):
            p(d,'表 4-1  非功能验收与责任基线',9,WD_ALIGN_PARAGRAPH.CENTER,True,3)
            add_table(d,[['类别','验收基线与可判定证据'],['性能','PE-01≤3秒；PE-02≤2秒；PE-03≤3分钟；PE-04≥20路、≥99%；PE-05≤2秒/次、亚米级；PE-06≤3秒；PE-07的95%请求≤3秒；PE-08≤2秒；PE-09≤1秒；PE-10安防≤30秒、应急信息≤60秒；PE-11≥100；PE-12≤5秒。'],['可靠性与安全','试运行可用率≥99.5%（计划内维护除外）；核心模块单元测试覆盖率≥70%；高危漏洞为0。'],['部署与兼容','容器化一键部署；干净环境至可访问≤2小时；非乙方人员按手册独立部署一次成功。'],['接口责任','甲方提供既有系统、资料、账号、环境、合法数据和协调窗口；乙方负责适配、联调、容错、降级、验证和交付。']],8.0)
            heading(d,raw[3:],1)
        elif raw.startswith('## 6 '):
            d.add_page_break();heading(d,raw[3:],1)
        else:heading(d,raw[3:],1)
    elif raw.startswith('### '):heading(d,raw[4:],2)
    elif raw.startswith('|'):
        group=[raw]
        while i<len(lines) and lines[i].strip().startswith('|'):group.append(lines[i].strip());i+=1
        rows=[[v.strip() for v in x.strip('|').split('|')] for x in group if not re.match(r'^\|[ -|]+\|$',x)]
        add_table(d,rows,7.05 if len(rows)>4 else 8)
    elif raw.startswith('- '):p(d,'• '+raw[2:],9.4,after=3)
    else:p(d,raw,10.0,after=4)
p(d,'表 6-1  签署信息表',9,WD_ALIGN_PARAGRAPH.CENTER,True,3)
add_table(d,[['甲方（课程模拟）：某自然博物馆','乙方：项目组/编制单位【正式名称或组号待人工确认】'],['授权代表：教师/评委代表【姓名待人工确认】\n签字/签章：【待人工确认】\n日期：【待人工确认】','项目经理：A【真实姓名待人工确认】\n技术复核：B【真实姓名待人工确认】\n合规复核：C【真实姓名待人工确认】\n签字/签章：【待人工确认】\n日期：【待人工确认】']],9)
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT);print(OUT)
