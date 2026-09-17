from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pathlib import Path
p=Presentation(); p.slide_width=Inches(10); p.slide_height=Inches(5.2); s=p.slides.add_slide(p.slide_layouts[6])
def box(x,y,w,h,text,color):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(*color); sh.line.color.rgb=RGBColor(80,80,80); tf=sh.text_frame; tf.text=text
 for para in tf.paragraphs: para.font.size=Pt(13); para.font.name='SimSun'; para.alignment=1
 return sh
box(.4,2,1.2,.8,'Web\n大屏',(220,230,240)); box(.4,3.2,1.2,.8,'H5',(220,230,240)); box(2.0,2.5,1.4,1,'MOD-PLATFORM\nAPI/审计',(210,225,240)); box(4,1,1.5,.7,'领域模块\nPLAN/EVENT/TASK',(230,240,225)); box(4,2,1.5,.7,'RESOURCE/DUTY\nDRILL/KNOW',(230,240,225)); box(4,3,1.5,.7,'SITUATION\n读投影',(230,240,225)); box(6.1,2.2,1.5,1,'MOD-INTEGRATION\n独立端口',(245,235,215)); box(8.1,1.0,1.4,.6,'视频/发布',(240,240,240)); box(8.1,1.8,1.4,.6,'入侵/门禁/消防',(240,240,240)); box(8.1,2.6,1.4,.6,'IoT/中台/消息',(240,240,240)); box(8.1,3.4,1.4,.6,'地图/定位',(240,240,240))
for a,b in [(1.6,2.4),(1.6,3.6),(3.4,3),(5.5,1.35),(5.5,2.35),(5.5,3.35),(7.6,2.7)]:
 ln=s.shapes.add_connector(1, Inches(a), Inches(2.9), Inches(b), Inches(2.9)); ln.line.color.rgb=RGBColor(90,90,90)
Path('docs/work/A_PM/figures').mkdir(parents=True,exist_ok=True); p.save('docs/work/A_PM/figures/g3_03_architecture.pptx'); p.save('docs/work/A_PM/figures/g3_03_architecture.pptx')
# export unavailable; use libreoffice conversion later
