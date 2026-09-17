from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import copy
ref=Path('docs/reference/11-概要设计说明书（教学样例）.docx')
out=Path('docs/deliverables/11-概要设计说明书.docx')
doc=Document(ref)
body=doc._element.body
sect=body.sectPr
for child in list(body):
    if child is not sect: body.remove(child)
styles=doc.styles
for style in ['Normal','Heading 1','Heading 2','Heading 3']:
    try:
        styles[style].font.name='Times New Roman';styles[style]._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体' if style=='Normal' else '黑体')
    except: pass
sec=doc.sections[0]
# cover
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('某自然博物馆智能运营中心建设项目\n应急管理子系统');r.bold=True;r.font.size=Pt(18)
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('概要设计说明书');r.bold=True;r.font.size=Pt(22)
for x in ['文档编号：EM-G3-03','版本：V0.1（评审稿）','编制：A【待人工确认】','复核：C【待人工确认】','日期：2026年09月17日']:
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(x)
doc.add_page_break()
p=doc.add_paragraph('修订记录',style='Heading 1')
t=doc.add_table(rows=1, cols=4);
for c,x in zip(t.rows[0].cells,['版本','日期','修改说明','责任人']): c.text=x
for row in [('V0.1','2026-09-17','首次形成概要设计评审稿','A【待人工确认】')]:
 cells=t.add_row().cells
 for c,x in zip(cells,row): c.text=x
doc.add_page_break()
def h(s,l=1): doc.add_paragraph(s,style=f'Heading {l}')
def para(s): doc.add_paragraph(s)
def table(headers, rows):
 t=doc.add_table(rows=1,cols=len(headers));
 for c,x in zip(t.rows[0].cells,headers): c.text=x
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row): c.text=x
 return t
h('1 引言');para('本说明书定义应急管理子系统的高层逻辑架构、数据边界、部署边界、外部系统边界及非功能设计约束，为后续数据库、详细设计、接口契约、ADR、RTM 和测试计划提供受控输入。')
h('1.1 设计范围',2);para('本设计覆盖 Web、大屏和 H5 渠道接入的应急业务逻辑架构。它不定义类、状态机细节、物理表、字段长度、索引、数据库产品或完整接口字段。原型 Mock、localStorage 和 prototypeStore 不构成正式架构依据。')
h('1.2 架构状态',2);para('ARC-A（模块化单体、端口适配与持久化异步任务）仅为当前候选；最终拆分粒度、异步一致性和部署编排仍须由后续 ADR 及人工决策确认。')
h('2 技术架构');para('整体关系如图 2-1 所示。渠道经统一 API 边界调用领域模块；领域模块拥有业务规则和写入责任，外部平台只经独立适配端口访问。')
table(['层次','设计元素'], [('渠道层','Web、大屏、H5；均经统一 API'),('业务层','MOD-PLAN、EVENT、TASK、RESOURCE、DUTY、DRILL、KNOWLEDGE'),('投影与适配层','MOD-SITUATION 只读投影；MOD-INTEGRATION 独立外部端口'),('公共能力层','MOD-PLATFORM：鉴权上下文、traceId、审计、幂等、错误、健康检查')])
para('表 2-1 逻辑技术架构分层')
h('2.1 模块职责与依赖',2)
table(['模块','主责与依赖'], [('MOD-PLAN','预案、版本、流程/任务模板；依赖中台工作流与文件'),('MOD-EVENT','事件、续报、核实、关闭；依赖预案、任务、IoT'),('MOD-TASK','任务快照、派发、反馈、回执补偿；依赖消息与组织权限'),('MOD-SITUATION','可重建态势读投影；依赖事件、资源、任务及地图/视频'),('MOD-RESOURCE','人员、物资、盘点快照/差异；依赖地图、组织、文件'),('MOD-DUTY / MOD-DRILL / MOD-KNOWLEDGE','值班打卡、演练评估、知识检索；分别依赖地图/消息/文件'),('MOD-MOBILE / MOD-INTEGRATION / MOD-PLATFORM','H5 渠道、外部适配、统一 API/审计/幂等/健康检查')])
h('3 数据架构');para('数据按领域模块唯一写入。事件、任务、资源、值班、演练、知识及审计保留业务编号、外部标识、traceId 和多时间语义。MOD-SITUATION 仅维护可重建投影，不能反向成为事实源。物理模型留待 G3-04。')
h('4 部署架构与故障域');para('系统部署于甲方内部环境，采用容器制品、配置外置和受控日志。渠道/API、领域业务、外部适配、数据和可观测性为逻辑故障域。建议配置 KN-070—078 尚待容量、网络、存储和恢复验证。')
table(['故障域','隔离与恢复约束'], [('外部适配','超时/失败不回滚已提交业务；有限重试、审计和人工降级'),('领域业务','先持久化业务状态，再派生通知、外部调用和读投影'),('控制调用','授权、二次确认、单次下发、回执对账；不自动重放'),('数据与可观测性','健康检查、关联日志、备份恢复演练和干净环境部署作为后续验证输入')])
h('5 外部边界与主要 NFR');para('EXT-VIDEO、EXT-PUBLISH、EXT-INTRUSION、EXT-ACCESS、EXT-FIRE、EXT-IOT、EXT-MIDDLE、EXT-MESSAGE 保持独立语义。甲方提供既有能力、账号、协议和测试环境；本系统负责业务适配、关联、审计和失败处理。')
para('ISSUE-G3-01-001 仍为 OPEN：视频/消息字段、认证、回执和性能未验证，接口不冻结且不得宣称连通。KN-064 以同一 traceId/eventId 验证视频、信息发布、物联网、中台的正常、无权、超时/失败联动。')
table(['设计预算','约束（待测）'], [('时效','预案下发≤3秒；告警接入≤2秒；定位≤2秒；视频首帧≤3秒'),('性能容量','95%页面≤3秒；峰值≥100人；事件/任务在线≥10年、≥2万/20万'),('可靠性安全','7×24；试运行≥99.5%；MTTR≤2小时；审计≥180天；高危漏洞0'),('部署','甲方内网私有化；干净环境一次成功且≤2小时')])
h('6 需求与设计追踪');table(['设计元素','需求/AC 锚点','后续挂接'], [('ARCH-01 统一 API 边界','AC-G2-FR-013-01、013-03','G3-06/G3-08'),('ARCH-02 事件任务编排','AC-G2-FR-003-01、014-01、015-01','G3-05/G3-08'),('ARCH-03 资源与快照','AC-G2-FR-009-01、025-02','G3-04/G3-08'),('ARCH-04 外部端口与降级','AC-G2-FR-006-01、020-02、027-01、029-03','G3-06/G3-08'),('ARCH-05 联动与可观测性','AC-G2-FR-006-01、027-01；KN-064','G3-08/G3-09')])
para('G3-08 将完成39条 FR、34条★FR、117条 AC 的双向设计挂接；本表不替代 RTM。')
h('7 设计限制与后续输入');para('本概要设计不替代后续 DBD、DLD、接口说明书、OpenAPI、ADR 或测试计划。任何外部字段、部署资源、坐标体系和中台能力未有受控证据时，标记【待人工确认】并按 Issue/变更流程处理。')
doc.save(out)
