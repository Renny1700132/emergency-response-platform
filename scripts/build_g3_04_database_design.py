from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = next((ROOT / "docs/reference").glob("13-*.docx"))
MARKDOWN = ROOT / "docs/work/B_TECH/database_design.md"
OUTPUT = ROOT / "docs/deliverables/13-数据库设计说明书.docx"
FIGURE_DIR = ROOT / "docs/deliverables/figures"
FIGURE_DIR.mkdir(parents=True, exist_ok=True)
PPTX = FIGURE_DIR / "13-图3-1-核心数据域关系图.pptx"
PNG = FIGURE_DIR / "13-图3-1-核心数据域关系图.png"


def font_path(bold=False):
    candidates = [
        Path("C:/Windows/Fonts/simhei.ttf") if bold else Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc") if bold else Path("C:/Windows/Fonts/msyh.ttc"),
    ]
    return str(next(p for p in candidates if p.exists()))


def build_figure():
    boxes = [
        (0.35, 0.45, 2.1, 0.75, "预案域\n版本  流程  模板", (225, 235, 245)),
        (3.05, 0.45, 2.1, 0.75, "事件域\n事件  续报  核实", (225, 240, 230)),
        (5.75, 0.45, 2.1, 0.75, "任务域\n任务  指派  反馈", (245, 235, 215)),
        (8.45, 0.45, 2.1, 0.75, "资源域\n人员  定位  物资", (235, 230, 245)),
        (11.0, 0.45, 1.95, 0.75, "值班演练\n打卡  评估", (235, 240, 225)),
        (2.05, 3.0, 2.35, 0.9, "业务事实层\n单一写入  状态版本\n四类时间", (218, 232, 244)),
        (5.5, 3.0, 2.35, 0.9, "可靠性层\nOutbox  幂等\n消息回执  调用日志", (243, 229, 204)),
        (8.95, 3.0, 2.35, 0.9, "横切数据层\n附件引用  审计\n可重建态势投影", (226, 235, 224)),
        (1.1, 5.55, 2.55, 0.85, "统一中台\n身份  组织  权限  文件", (240, 240, 240)),
        (5.35, 5.55, 2.55, 0.85, "既有业务系统\n视频  消息  门禁  IoT", (240, 240, 240)),
        (9.65, 5.55, 2.55, 0.85, "运维与数据治理\n备份  归档  恢复  迁移", (240, 240, 240)),
    ]
    # Deterministic PNG export matching the editable PPTX content.
    im = Image.new("RGB", (1920, 1080), "white")
    draw = ImageDraw.Draw(im)
    f1, f2 = ImageFont.truetype(font_path(True), 30), ImageFont.truetype(font_path(False), 23)
    sx, sy = 1920 / 13.333, 1080 / 7.5
    for x, y, w, h, txt, color in boxes:
        xy = (int(x*sx), int(y*sy), int((x+w)*sx), int((y+h)*sy))
        draw.rounded_rectangle(xy, radius=14, fill=color, outline=(80,80,80), width=2)
        lines = txt.split("\n")
        total = sum((f1 if i == 0 else f2).getbbox(line)[3] + 6 for i, line in enumerate(lines))
        yy = (xy[1] + xy[3] - total) / 2
        for i, line in enumerate(lines):
            ft = f1 if i == 0 else f2
            bb = draw.textbbox((0,0), line, font=ft)
            draw.text(((xy[0]+xy[2]-bb[2])/2, yy), line, font=ft, fill="black")
            yy += bb[3] + 6
    for x1, y1, x2, y2 in [
        (2.45,0.83,3.05,0.83),(5.15,0.83,5.75,0.83),(7.85,0.83,8.45,0.83),(10.55,0.83,11.0,0.83),
        (3.2,1.2,3.2,3.0),(6.8,1.2,6.8,3.0),(9.5,1.2,9.5,3.0),
        (4.4,3.45,5.5,3.45),(7.85,3.45,8.95,3.45),
        (3.2,3.9,2.4,5.55),(6.7,3.9,6.65,5.55),(10.1,3.9,10.9,5.55),
    ]:
        draw.line((int(x1*sx),int(y1*sy),int(x2*sx),int(y2*sy)), fill=(90,90,90), width=3)
    im.save(PNG)


def set_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:color", "w:highlight", "w:shd"):
        for node in rpr.findall(qn(tag)):
            rpr.remove(node)


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc; tcpr = tc.get_or_add_tcPr(); mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar"); tcpr.append(mar)
    for m, v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        n = mar.find(qn("w:"+m))
        if n is None: n = OxmlElement("w:"+m); mar.append(n)
        n.set(qn("w:w"), str(v)); n.set(qn("w:type"), "dxa")


def parse_blocks(text):
    lines = text.splitlines(); i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line or line.startswith("- 任务：") or line.startswith("- 文档编号：") or line.startswith("- 版本：") or line.startswith("- 主责") or line.startswith("- 状态：") or line.startswith("- 项目：") or line.startswith("- 受控输入："):
            i += 1; continue
        if line.startswith("# "):
            i += 1; continue
        if line.startswith("## "):
            yield ("h1", line[3:]); i += 1; continue
        if line.startswith("### "):
            yield ("h2", line[4:]); i += 1; continue
        if line == "[[FIGURE:core_domain_model]]":
            yield ("figure", None); i += 1; continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*[-:]+", lines[i+1]):
            rows=[]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([re.sub(r"<br\s*/?>", "\n", re.sub(r"[`*]", "", c.strip()), flags=re.I) for c in lines[i].strip().strip("|").split("|")]); i += 1
            yield ("table", [rows[0]] + rows[2:]); continue
        if re.match(r"^\d+\.\s+", line):
            yield ("list", re.sub(r"^\d+\.\s+", "", line)); i += 1; continue
        if line.startswith("-"):
            yield ("list", line.lstrip("- ")); i += 1; continue
        para=[line]; i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#","|","[[FIGURE:")) and not re.match(r"^\d+\.\s+", lines[i]):
            para.append(lines[i].strip()); i += 1
        yield ("p", " ".join(para))


def add_inline(paragraph, text):
    # Preserve inline code and bold semantics without leaking markdown marks.
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part: continue
        if part.startswith("`"):
            r=paragraph.add_run(part[1:-1]); r.font.name="Consolas"; r.font.size=Pt(10.5)
        elif part.startswith("**"):
            r=paragraph.add_run(part[2:-2]); r.bold=True
        else: r=paragraph.add_run(part)
        set_black(r)


def copy_para_format(dst, src):
    if src._p.pPr is not None:
        if dst._p.pPr is not None: dst._p.remove(dst._p.pPr)
        dst._p.insert(0, copy.deepcopy(src._p.pPr))


def build_docx():
    shutil.copy2(REFERENCE, OUTPUT)
    d = Document(OUTPUT)
    body = d._element.body; sectpr = body.sectPr
    # Keep cover, revision heading, revision table and the blank after it; remove sample body.
    start = d.paragraphs[18]._p
    remove = False
    for child in list(body):
        if child is start: remove = True
        if remove and child is not sectpr: body.remove(child)

    cover = {
        0: "文档编号：YJGL-G3-04　　版本号：V0.2（整改评审稿）",
        1: "密　　级：内部 · 评审用",
        3: "某自然博物馆智能运营中心建设项目——应急管理子系统",
        4: "数据库设计说明书",
        5: "（G3-04 整改评审稿）",
        7: "编制单位：020202项目组【待人工确认】",
        8: "编　　制：B（技术主责）【待人工确认】",
        9: "审　　核：A、C【待复核】",
        10: "批　　准：【待人工确认】",
        11: "编制日期：2026 年 9 月 17 日",
        14: "注：本表记录数据库设计受控版本沿革；本版进入 A/C 复核，不代表现场联调、压测或验收已通过。",
    }
    for idx, text in cover.items():
        p=d.paragraphs[idx]; p.clear(); p.add_run(text)
    rev=d.tables[0]
    while len(rev.rows)>2: rev._tbl.remove(rev.rows[-1]._tr)
    vals=["V0.1","2026-09-17","全部","形成概念、逻辑、物理三级数据库设计评审稿","B【待人工确认】"]
    for c,v in zip(rev.rows[1].cells,vals): c.text=v
    row = rev.add_row()
    vals=["V0.2","2026-09-17","第4、8章","补齐39 FR、34★、117 AC逐条追踪及支撑实体","B【待人工确认】"]
    for c,v in zip(row.cells,vals): c.text=v

    ref = Document(REFERENCE)
    sample_table = ref.tables[1]
    caption_sample = ref.paragraphs[27]
    chapter_no=0; table_no=0; current_heading=""
    for kind, payload in parse_blocks(MARKDOWN.read_text(encoding="utf-8")):
        if kind == "h1":
            chapter_no += 1; table_no=0; current_heading=payload
            p=d.add_paragraph(style="Heading 1"); add_inline(p,payload)
        elif kind == "h2":
            current_heading=payload
            p=d.add_paragraph(style="Heading 2"); add_inline(p,payload)
        elif kind == "p":
            p=d.add_paragraph(); add_inline(p,payload)
        elif kind == "list":
            p=d.add_paragraph(style="List Paragraph"); add_inline(p,"• "+payload)
        elif kind == "figure":
            p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(PNG), width=Inches(6.1))
            cap=d.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            add_inline(cap,"图 3-1 核心数据域关系图")
            copy_para_format(cap, caption_sample)
        elif kind == "table":
            table_no += 1
            cap=d.add_paragraph(); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
            title=re.sub(r"^\d+(?:\.\d+)?\s*", "", current_heading)
            add_inline(cap,f"表 {chapter_no}-{table_no} {title}")
            copy_para_format(cap, caption_sample)
            rows=payload; t=d.add_table(rows=1, cols=len(rows[0]))
            t.autofit = False
            if sample_table._tbl.tblPr is not None:
                t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,copy.deepcopy(sample_table._tbl.tblPr))
            for ci,val in enumerate(rows[0]):
                t.rows[0].cells[ci].text=val; shade(t.rows[0].cells[ci],"D9E2F3")
            for row in rows[1:]:
                cells=t.add_row().cells
                for ci,val in enumerate(row): cells[ci].text=val
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for ri,row in enumerate(t.rows):
                for ci,cell in enumerate(row.cells):
                    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; set_cell_margin(cell)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.15
                        for r in p.runs:
                            r.font.name="SimSun"; r.font.size=Pt(10.5); r.bold=(ri==0); set_black(r)
                    if ri and ri%2==0: shade(cell,"F7F9FC")
            width_map = {
                2: [5.0, 10.2],
                3: [3.0, 6.0, 6.2],
                4: [2.5, 4.2, 3.4, 5.1],
                5: [2.0, 3.1, 2.5, 4.5, 3.1],
            }
            widths = width_map.get(len(rows[0]), [15.2/len(rows[0])] * len(rows[0]))
            if rows[0][0].startswith("设计 ID / FR"):
                widths = [2.4, 4.1, 5.4, 3.3]
            tbl_pr = t._tbl.tblPr
            tbl_layout = tbl_pr.find(qn("w:tblLayout"))
            if tbl_layout is None:
                tbl_layout = OxmlElement("w:tblLayout"); tbl_pr.append(tbl_layout)
            tbl_layout.set(qn("w:type"), "fixed")
            tbl_w = tbl_pr.find(qn("w:tblW"))
            if tbl_w is None:
                tbl_w = OxmlElement("w:tblW"); tbl_pr.append(tbl_w)
            tbl_w.set(qn("w:w"), str(int(sum(widths)*567))); tbl_w.set(qn("w:type"), "dxa")
            grid_cols = t._tbl.tblGrid.gridCol_lst
            for ci, width in enumerate(widths):
                if ci < len(grid_cols): grid_cols[ci].set(qn("w:w"), str(int(width*567)))
                t.columns[ci].width = Cm(width)
            for row in t.rows:
                for ci, cell in enumerate(row.cells):
                    cell.width = Cm(widths[ci])

    # Ensure all text is black, preserve template fonts and section/header/footer.
    for p in d.paragraphs:
        for r in p.runs: set_black(r)
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: set_black(r)
    d.core_properties.title="数据库设计说明书"
    d.core_properties.subject="某自然博物馆智能运营中心建设项目——应急管理子系统 G3-04"
    d.core_properties.author="020202项目组"
    d.save(OUTPUT)


if __name__ == "__main__":
    build_figure()
    ascii_pptx = FIGURE_DIR / "g3_04_core_domain_model.pptx"
    if ascii_pptx.exists():
        shutil.copy2(ascii_pptx, PPTX)
        ascii_pptx.unlink()
    build_docx()
    print(OUTPUT)
    print(PPTX)
    print(PNG)
