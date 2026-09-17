from __future__ import annotations

import argparse
import copy
import os
import re
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor as PptRGB
from pptx.util import Inches as PptInches, Pt as PptPt


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = next((ROOT / "docs/reference").glob("12-*.docx"))
MARKDOWN = ROOT / "docs/work/B_TECH/detailed_design.md"
OUTPUT = ROOT / "docs/deliverables/12-详细设计说明书.docx"
FIGURE_DIR = ROOT / "docs/deliverables/figures"
PPTX = FIGURE_DIR / "12-详细设计图源.pptx"
FIGURES = {
    "component_structure": ("图 2-1 详细设计组件结构", "12-图2-1-详细设计组件结构.png"),
    "incident_flow": ("图 4-1 事件核实与响应启动流程", "12-图4-1-事件核实与响应启动流程.png"),
    "task_message_flow": ("图 5-1 任务与可靠消息协作流程", "12-图5-1-任务与可靠消息协作流程.png"),
    "inventory_attendance_flow": ("图 6-1 盘点快照与移动打卡控制流程", "12-图6-1-盘点快照与移动打卡控制流程.png"),
}


REFERENCE_PACKAGE_PARTS = {
    "word/styles.xml", "word/stylesWithEffects.xml", "word/numbering.xml",
    "word/settings.xml", "word/fontTable.xml", "word/webSettings.xml",
    "word/theme/theme1.xml", "word/header1.xml", "word/footer1.xml",
}


def font_path(bold=False):
    candidates = [
        Path("C:/Windows/Fonts/simhei.ttf") if bold else Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc") if bold else Path("C:/Windows/Fonts/msyh.ttc"),
    ]
    return str(next(p for p in candidates if p.exists()))


def draw_flow(path: Path, title: str, nodes: list[tuple[str, str]], edges: list[tuple[int, int, str]], snake=True):
    im = Image.new("RGB", (1800, 880), "white")
    d = ImageDraw.Draw(im)
    title_font = ImageFont.truetype(font_path(True), 38)
    body_font = ImageFont.truetype(font_path(False), 25)
    small_font = ImageFont.truetype(font_path(False), 19)
    d.text((900, 50), title, font=title_font, fill="black", anchor="mm")
    cols = 4
    boxes = []
    for i, (head, body) in enumerate(nodes):
        row, col = divmod(i, cols)
        if snake and row == 1: col = 3 - col
        x1, y1 = 80 + col * 430, 155 + row * 280
        x2, y2 = x1 + 330, y1 + 150
        boxes.append((x1, y1, x2, y2))
        fill = [(224, 235, 247), (226, 239, 218), (255, 242, 204), (234, 226, 243)][col]
        d.rounded_rectangle((x1, y1, x2, y2), 18, fill=fill, outline=(70, 70, 70), width=3)
        d.text(((x1+x2)//2, y1+40), head, font=body_font, fill="black", anchor="mm")
        d.multiline_text(((x1+x2)//2, y1+98), body, font=small_font, fill="black", anchor="mm", align="center", spacing=5)
    for a, b, label in edges:
        ax1, ay1, ax2, ay2 = boxes[a]; bx1, by1, bx2, by2 = boxes[b]
        if abs((ax1+ax2)-(bx1+bx2)) > abs((ay1+ay2)-(by1+by2)):
            start = (ax2 if bx1 > ax1 else ax1, (ay1+ay2)//2)
            end = (bx1 if bx1 > ax1 else bx2, (by1+by2)//2)
        else:
            start = ((ax1+ax2)//2, ay2 if by1 > ay1 else ay1)
            end = ((bx1+bx2)//2, by1 if by1 > ay1 else by2)
        d.line((*start, *end), fill=(80,80,80), width=4)
        vx, vy = end[0]-start[0], end[1]-start[1]
        mag = max((vx*vx+vy*vy) ** .5, 1)
        ux, uy = vx/mag, vy/mag
        px, py = -uy, ux
        tip = end
        p1 = (end[0]-18*ux+8*px, end[1]-18*uy+8*py)
        p2 = (end[0]-18*ux-8*px, end[1]-18*uy-8*py)
        d.polygon([tip,p1,p2], fill=(80,80,80))
        if label:
            d.text(((start[0]+end[0])//2, (start[1]+end[1])//2-12), label, font=small_font, fill=(30,30,30), anchor="mm")
    im.save(path)


def add_ppt_slide(prs, title, nodes, edges, snake=True):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(PptInches(.4), PptInches(.2), PptInches(12.5), PptInches(.5))
    p = tb.text_frame.paragraphs[0]; p.text = title; p.alignment = PP_ALIGN.CENTER
    p.runs[0].font.name = "黑体"; p.runs[0].font.size = PptPt(22); p.runs[0].font.bold = True
    boxes=[]
    for i,(head,body) in enumerate(nodes):
        row,col=divmod(i,4)
        if snake and row == 1: col=3-col
        x=.55+col*3.2; y=1.25+row*2.25
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PptInches(x),PptInches(y),PptInches(2.45),PptInches(1.25))
        sh.fill.solid(); sh.fill.fore_color.rgb=[PptRGB(224,235,247),PptRGB(226,239,218),PptRGB(255,242,204),PptRGB(234,226,243)][col]
        sh.line.color.rgb=PptRGB(80,80,80); sh.text=f"{head}\n{body}"
        for j,p in enumerate(sh.text_frame.paragraphs):
            p.alignment=PP_ALIGN.CENTER
            for r in p.runs:
                r.font.name="宋体"; r.font.size=PptPt(13 if j else 15); r.font.bold=(j==0); r.font.color.rgb=PptRGB(0,0,0)
        boxes.append(sh)
    for a,b,label in edges:
        sa,sb=boxes[a],boxes[b]
        if abs(sa.top-sb.top) < PptInches(.2):
            if sb.left > sa.left: x1,y1=sa.left+sa.width,sa.top+sa.height//2; x2,y2=sb.left,sb.top+sb.height//2
            else: x1,y1=sa.left,sa.top+sa.height//2; x2,y2=sb.left+sb.width,sb.top+sb.height//2
            kind=1
        else:
            x1,y1=sa.left+sa.width//2,sa.top+sa.height; x2,y2=sb.left+sb.width//2,sb.top; kind=2
        line=slide.shapes.add_connector(kind,x1,y1,x2,y2); line.line.color.rgb=PptRGB(90,90,90); line.line.width=PptPt(1.5)
    return slide


def build_figures():
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    specs = [
        ("component_structure", "详细设计组件结构", [("Web / 大屏 / H5","渠道与查询"),("应用服务","用例编排"),("领域组件","状态与规则"),("仓储 / Outbox","事务与可靠事件"),("态势投影","可重建读模型"),("审计 / 可观测","traceId 与指标"),("外部端口","协议隔离"),("既有系统","中台与专业平台")], [(0,1,"命令/查询"),(1,2,"调用"),(2,3,"持久化"),(3,4,"事件"),(4,5,"留痕"),(5,6,"端口"),(6,7,"适配")], True),
        ("incident_flow", "事件核实与响应启动流程", [("告警/上报","可靠 ID 去重"),("事件待核实","状态与四类时间"),("人工核实","超时只升级"),("响应启动","绑定预案版本"),("任务工厂","原子生成任务"),("Outbox","提交后派生"),("消息与态势","失败不回滚事实"),("关闭/重开","材料与审计")], [(0,1,"接入"),(1,2,"待办"),(2,3,"已核实"),(3,4,"同事务"),(4,5,"同事务"),(5,6,"异步"),(6,7,"闭环")], True),
        ("task_message_flow", "任务与可靠消息协作流程", [("任务命令","鉴权与幂等"),("任务聚合","状态迁移"),("事务提交","事实+Outbox"),("消息调度","解析收件人"),("统一消息","发送/回执"),("投递记录","尝试与终态"),("迟到回执","幂等关联"),("人工降级","失败处置")], [(0,1,"校验"),(1,2,"保存"),(2,3,"派生"),(3,4,"调用"),(4,5,"回执"),(5,6,"迟到"),(6,7,"失败")], True),
        ("inventory_attendance_flow", "盘点快照与移动打卡控制流程", [("盘点发布","冻结账面快照"),("移动实盘","只追加记录"),("差异计算","含期间变动"),("授权复核","调整台账"),("扫码请求","身份/二维码"),("时空校验","时段/范围/新鲜度"),("幂等写入","返回首次结果"),("缺卡扫描","告警+Outbox")], [(0,1,"采集"),(1,2,"计算"),(2,3,"确认"),(4,5,"校验"),(5,6,"成功"),(6,7,"截止扫描")], False),
    ]
    prs=Presentation(); prs.slide_width=PptInches(13.333); prs.slide_height=PptInches(7.5)
    for key,title,nodes,edges,snake in specs:
        draw_flow(FIGURE_DIR / FIGURES[key][1], title, nodes, edges, snake)
        add_ppt_slide(prs,title,nodes,edges,snake)
    prs.save(PPTX)


def set_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    for tag in ("w:color", "w:highlight", "w:shd"):
        for node in rpr.findall(qn(tag)): rpr.remove(node)


def shade(cell, fill):
    tcpr=cell._tc.get_or_add_tcPr(); shd=tcpr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcpr.append(shd)
    shd.set(qn("w:fill"),fill)


def parse_blocks(text):
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].rstrip()
        if not line or line.startswith(("- 任务：","- 文档编号：","- 版本：","- 主责","- 状态：","- 项目：","- 受控输入：")): i+=1; continue
        if line.startswith("# "): i+=1; continue
        if line.startswith("## "): yield "h1",line[3:]; i+=1; continue
        if line.startswith("### "): yield "h2",line[4:]; i+=1; continue
        m=re.fullmatch(r"\[\[FIGURE:([a-z_]+)\]\]",line)
        if m: yield "figure",m.group(1); i+=1; continue
        if line.startswith("|") and i+1<len(lines) and re.match(r"^\|\s*[-:]+",lines[i+1]):
            rows=[]
            while i<len(lines) and lines[i].startswith("|"):
                rows.append([re.sub(r"<br\s*/?>","\n",re.sub(r"[`*]","",c.strip()),flags=re.I) for c in lines[i].strip().strip("|").split("|")]); i+=1
            yield "table",[rows[0]]+rows[2:]; continue
        if re.match(r"^\d+\.\s+",line): yield "list",re.sub(r"^\d+\.\s+","",line); i+=1; continue
        para=[line]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(("#","|","[[FIGURE:")) and not re.match(r"^\d+\.\s+",lines[i]):
            para.append(lines[i].strip()); i+=1
        yield "p"," ".join(para)


def copy_run_format(dst, src):
    if src is not None and src._r.rPr is not None:
        if dst._r.rPr is not None: dst._r.remove(dst._r.rPr)
        dst._r.insert(0, copy.deepcopy(src._r.rPr))


def replace_text_with_format(paragraph, text, template_paragraph):
    paragraph.clear()
    run = paragraph.add_run(text)
    if template_paragraph.runs: copy_run_format(run, template_paragraph.runs[0])
    return run


def add_inline(p,text,run_template=None):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)",text):
        if not part: continue
        if part.startswith("`"): r=p.add_run(part[1:-1]); copy_run_format(r,run_template); r.font.name="Consolas"; r.font.size=Pt(10.5)
        elif part.startswith("**"): r=p.add_run(part[2:-2]); copy_run_format(r,run_template); r.bold=True
        else: r=p.add_run(part)
        if not part.startswith(("`","**")): copy_run_format(r,run_template)


def copy_para_format(dst,src):
    if src._p.pPr is not None:
        if dst._p.pPr is not None: dst._p.remove(dst._p.pPr)
        dst._p.insert(0,copy.deepcopy(src._p.pPr))


def copy_cell_visual_format(dst, src):
    dst_pr = dst._tc.get_or_add_tcPr()
    src_pr = src._tc.tcPr
    for child in list(dst_pr):
        if child.tag != qn("w:tcW"): dst_pr.remove(child)
    if src_pr is not None:
        for child in src_pr:
            if child.tag != qn("w:tcW"): dst_pr.append(copy.deepcopy(child))


def format_cell_text(cell, template_cell, bold=False):
    src_p = template_cell.paragraphs[0]
    for p in cell.paragraphs:
        copy_para_format(p, src_p)
        for r in p.runs:
            copy_run_format(r, src_p.runs[0] if src_p.runs else None)
            if bold: r.bold=True


def restore_reference_package_parts(output, reference):
    temp = output.with_suffix(".pair-tmp.docx")
    with zipfile.ZipFile(output, "r") as out_zip, zipfile.ZipFile(reference, "r") as ref_zip:
        out_names = out_zip.namelist(); ref_names = set(ref_zip.namelist())
        with zipfile.ZipFile(temp, "w", zipfile.ZIP_DEFLATED) as new_zip:
            for name in out_names:
                data = ref_zip.read(name) if name in REFERENCE_PACKAGE_PARTS and name in ref_names else out_zip.read(name)
                new_zip.writestr(name, data)
    os.replace(temp, output)


def build_docx():
    shutil.copy2(REFERENCE,OUTPUT)
    d=Document(OUTPUT); body=d._element.body; sectpr=body.sectPr; start=d.paragraphs[16]._p; removing=False
    ref=Document(REFERENCE)
    for child in list(body):
        if child is start: removing=True
        if removing and child is not sectpr: body.remove(child)
    cover={0:"文档编号：YJGL-G3-05　　版本号：V0.1（评审稿）",1:"密　　级：内部 · 评审用",3:"某自然博物馆智能运营中心建设项目——应急管理子系统",4:"详细设计说明书",5:"（G3-05 评审稿）",7:"编制单位：020202项目组【待人工确认】",8:"编　　制：B（技术主责）【待人工确认】",9:"审　　核：A、C【待复核】",10:"批　　准：【待人工确认】",11:"编制日期：2026 年 9 月 17 日",14:"注：本表记录详细设计受控版本沿革；本版进入 A/C 复核，不代表编码、联调、压测或验收已通过。"}
    for idx,text in cover.items():
        replace_text_with_format(d.paragraphs[idx], text, ref.paragraphs[idx])
    rev=d.tables[0]
    while len(rev.rows)>2: rev._tbl.remove(rev.rows[-1]._tr)
    row=rev.rows[1]; vals=["V0.1","2026-09-17","全部","形成模块职责、逻辑类、关键流程、异常与追踪评审稿","B【待人工确认】"]
    for ci,(c,v) in enumerate(zip(row.cells,vals)):
        replace_text_with_format(c.paragraphs[0],v,ref.tables[0].rows[1].cells[ci].paragraphs[0])
    sample_table=ref.tables[1]; figure_para_sample=ref.paragraphs[20]; figure_caption_sample=ref.paragraphs[21]; table_caption_sample=ref.paragraphs[48]
    h1_sample=ref.paragraphs[16]; h2_sample=ref.paragraphs[42]; body_sample=ref.paragraphs[17]; list_sample=ref.paragraphs[44]
    chapter=0; table_no=0; current=""
    for kind,payload in parse_blocks(MARKDOWN.read_text(encoding="utf-8")):
        if kind=="h1": chapter+=1; table_no=0; current=payload; p=d.add_paragraph(style="Heading 1"); copy_para_format(p,h1_sample); add_inline(p,payload,h1_sample.runs[0])
        elif kind=="h2": current=payload; p=d.add_paragraph(style="Heading 2"); copy_para_format(p,h2_sample); add_inline(p,payload,h2_sample.runs[0])
        elif kind=="p": p=d.add_paragraph(); copy_para_format(p,body_sample); add_inline(p,payload,body_sample.runs[0])
        elif kind=="list": p=d.add_paragraph(style="List Paragraph"); copy_para_format(p,list_sample); add_inline(p,"• "+payload,list_sample.runs[0])
        elif kind=="figure":
            p=d.add_paragraph(); copy_para_format(p,figure_para_sample); figure_run=p.add_run(); copy_run_format(figure_run,figure_para_sample.runs[0]); figure_run.add_picture(str(FIGURE_DIR/FIGURES[payload][1]),width=ref.inline_shapes[0].width)
            cap=d.add_paragraph(); copy_para_format(cap,figure_caption_sample); add_inline(cap,FIGURES[payload][0],figure_caption_sample.runs[0])
        elif kind=="table":
            table_no+=1; cap=d.add_paragraph()
            table_title = re.sub(r"^\d+(?:\.\d+)?\s*", "", current)
            copy_para_format(cap,table_caption_sample); cap.paragraph_format.keep_with_next=True; add_inline(cap, f"表 {chapter}-{table_no} {table_title}",table_caption_sample.runs[0])
            rows=payload; t=d.add_table(rows=1,cols=len(rows[0]))
            if sample_table._tbl.tblPr is not None: t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,copy.deepcopy(sample_table._tbl.tblPr))
            for ci,val in enumerate(rows[0]): t.rows[0].cells[ci].text=val; shade(t.rows[0].cells[ci],"D9E2F3")
            for values in rows[1:]:
                cells=t.add_row().cells
                for ci,val in enumerate(values): cells[ci].text=val
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            widths={2:[5,10.2],3:[3.2,5.7,6.3],4:[2.7,4.0,4.3,4.2],5:[2.1,3.2,3.2,3.7,3.0]}.get(len(rows[0]),[15.2/len(rows[0])]*len(rows[0]))
            if rows[0][0].startswith("设计 ID"): widths=[2.6,4.0,4.4,4.2]
            label_cell=sample_table.rows[0].cells[0]; value_cell=sample_table.rows[1].cells[1]
            for ri,row in enumerate(t.rows):
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for ci,cell in enumerate(row.cells):
                    copy_cell_visual_format(cell,label_cell if ri==0 else value_cell)
                    cell.width=Cm(widths[ci])
                    format_cell_text(cell,label_cell if ri==0 else value_cell,bold=(ri==0))
    d.core_properties.title="详细设计说明书"; d.core_properties.subject="某自然博物馆智能运营中心建设项目——应急管理子系统 G3-05"; d.core_properties.author="020202项目组"
    d.save(OUTPUT)
    restore_reference_package_parts(OUTPUT, REFERENCE)


if __name__=="__main__":
    ap=argparse.ArgumentParser(); ap.add_argument("--figures-only",action="store_true"); ap.add_argument("--docx-only",action="store_true"); args=ap.parse_args()
    if not args.docx_only: build_figures()
    if not args.figures_only: build_docx()
    print(PPTX); print(OUTPUT)
