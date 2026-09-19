from __future__ import annotations

import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SPEC = ROOT / "docs/work/B_TECH/spec.md"
WORK = ROOT / "docs/work/C_REQ/test_plan.md"
REFERENCE = ROOT / "docs/reference/21-测试计划（教学样例）.docx"
OUTPUT = ROOT / "docs/deliverables/19-测试计划.docx"
FIGURE = ROOT / "docs/deliverables/figures/00-图9-2-四级测试与验收证据链.png"


def clean(text: str) -> str:
    return re.sub(r"\*\*|`", "", text).strip()


def parse_requirements() -> list[dict[str, object]]:
    text = SPEC.read_text(encoding="utf-8")
    heading_re = re.compile(r"^### (G2-FR-(\d{3}))｜([^｜]+)｜(★|非★)｜(.+)$", re.M)
    matches = list(heading_re.finditer(text))
    result = []
    for idx, match in enumerate(matches):
        block = text[match.end() : matches[idx + 1].start() if idx + 1 < len(matches) else len(text)]
        acs = []
        for ac_id, ac_text in re.findall(r"- (AC-G2-FR-\d{3}-\d{2})：(.+)", block):
            then = re.search(r"\*\*Then\*\*\s*(.+?)(?:。|$)", ac_text)
            acs.append((ac_id, clean(then.group(1) if then else ac_text)))
        if len(acs) != 3:
            raise ValueError(f"{match.group(1)} has {len(acs)} ACs")
        result.append({"id": match.group(1), "num": int(match.group(2)), "orig": match.group(3), "star": match.group(4), "title": match.group(5), "acs": acs})
    if len(result) != 39:
        raise ValueError(f"Expected 39 FRs, found {len(result)}")
    return result


def level_for(n: int) -> str:
    if n in {3, 5, 6, 11, 20, 26, 27, 28, 29, 31}:
        return "集成/契约/系统"
    if n in {14, 15, 16, 21, 22, 23, 24, 25}:
        return "集成/系统"
    if n in {1, 2, 8, 9, 10, 12, 17, 18, 19, 30, 32, 33, 34, 35, 36, 37, 38, 39}:
        return "单元/系统"
    return "系统"


def build_markdown(reqs: list[dict[str, object]]) -> str:
    lines = [
        "# 某自然博物馆智能运营中心建设项目——应急管理子系统测试计划",
        "",
        "- 任务：G3-09",
        "- 版本：V0.2 评审候选",
        "- 主责 / 复核：C / A（OVR-025 单人复核）",
        "- 状态：REVIEW（计划与测试设计完成，尚未执行测试）",
        "- 日期：2026-09-19",
        "- 依据：G3-08 RTM、spec 39 FR/117 AC、G3-03—07 设计集、GB/T 15532-2008、GB/T 25000.10 质量特性映射。",
        "",
        "## 1 引言",
        "",
        "### 1.1 编写目的",
        "",
        "本计划在设计完成时锁定测试范围、级别、策略、进度、准入准出、环境和证据口径，使 39 条 FR、34 条★、117 条 AC 以及 PE/NFR/接口边界可在后续实现阶段直接转化为可重复测试。G3-09 只形成测试计划和用例设计，不把尚未执行的测试预填为通过。",
        "",
        "### 1.2 测试依据",
        "",
        "- `docs/work/B_TECH/spec.md`：39 条 FR、117 条 Given/When/Then AC 与 G2-RCLR-001—010。",
        "- `docs/work/C_REQ/rtm_v1.md`：FR—AC—ARCH/DLD/DBD/API—测试占位双向追踪。",
        "- G3-03—G3-07：概要、详细、数据库、接口/OpenAPI、ADR 与非功能设计。",
        "- `control/facts.md`、`control/key_numbers.md`、`control/issues.md`：事实、数字、责任边界和开放阻断。",
        "- GB/T 15532-2008：测试级别、计划、用例、执行、缺陷与结果记录组织依据。",
        "- GB/T 25000.10 质量特性映射：性能效率、可靠性、安全性、兼容性、易用性、可维护性和可移植性测试分类依据。",
        "",
        "### 1.3 证据纪律",
        "",
        "每次执行必须记录版本、环境、数据集、脚本或步骤、起止时间、操作者、原始结果、缺陷号和结论。Mock、桩、静态扫描、文档检查和渲染 QA 只能证明其对应层级，不得替代真实接口、目标环境性能、安全、恢复或验收证据。`ISSUE-G3-01-001` 保持 OPEN，视频和统一消息通道的真实字段、认证、回执与连通结论不得预填。",
        "",
        "## 2 测试范围",
        "",
        "### 2.1 测试对象与基线",
        "",
        "表 2-1 测试对象与基线",
        "",
        "| 测试对象 | 受控规模 | 主要依据 | 本计划状态 |",
        "| --- | --- | --- | --- |",
        "| 功能需求 | 39 FR、117 AC、34★ | spec、SRS、RTM | 117 个 AC 级测试设计已编号，待实现/执行 |",
        "| 设计与数据 | 39 DLD-TR、39 DBD-TR、39 API-TR | 11—15 号交付物 | 设计可追踪，代码与数据库未实现 |",
        "| 接口契约 | OpenAPI 57 项操作、8 个外部端口 | 14、15 号交付物 | 静态契约可检查，真实联调待环境 |",
        "| 非功能 | PE-01—12、容量/可靠性/安全/兼容/易用/维护/移植 | 18 号交付物 | 方法与证据口径已定义，结果待测 |",
        "| 澄清裁决 | G2-RCLR-001—010 | spec §3.6、RTM | 纳入相关 AC 的正例、反例和异常场景 |",
        "",
        "### 2.2 测试级别",
        "",
        "表 2-2 测试级别、责任与目标",
        "",
        "| 级别 | 目标 | 主要范围 | 主责 / 复核 |",
        "| --- | --- | --- | --- |",
        "| 单元测试 | 验证状态机、校验、幂等、权限和计算规则 | 领域服务、规则、值对象、映射器 | B / A、C |",
        "| 组件与集成测试 | 验证模块、数据库、Outbox、文件和外部适配器协作 | MOD-*、DBD-TR、DLD-TR、EXT-* | B / C |",
        "| 契约测试 | 阻断 OpenAPI、错误码和事件结构漂移 | 57 项操作、API-TR-001—039 | B / A、C |",
        "| 系统测试 | 按业务角色执行 39 FR/117 AC 端到端场景 | Web、H5、应用服务、数据与地图 | C / A |",
        "| 专项测试 | 验证性能、可靠性、安全、兼容、易用、恢复和部署 | PE/NFR/ENG 与关键数字 | C 组织，B 技术支持，A 协调 |",
        "| 验收支持 | 整理可由甲方复核的测试包和偏差 | ★、关键数字、外部接口、交付物 | A 组织，B/C 支持 |",
        "",
        "### 2.3 覆盖规则",
        "",
        "1. 每条 AC 对应一个稳定测试编号 `TC-G2-FR-nnn-xx`，共 117 个；一个设计用例可由多条自动化函数或人工步骤实现。",
        "2. 34 条★的 102 个 AC 必须全部执行并通过；任何失败或缺证据均阻断发布候选。",
        "3. 每条 FR 至少覆盖正常路径和权限、边界或失败路径；性能型 AC 必须记录环境、数据规模、分位值、错误率和资源使用。",
        "4. 每条测试必须反向回指 AC、设计 ID、执行结果和缺陷；无需求来源的孤儿用例不得纳入覆盖率。",
        "5. 核心单元测试覆盖率目标为 ≥70%，39 FR 功能覆盖与验收用例通过率目标为 100%；这些均是后续执行门禁，不是 G3-09 已测结论。",
        "",
        "### 2.4 不测与延期范围",
        "",
        "- 不验证甲方既有视频系统的录像存储实现，只验证本项目接口调阅、回放及甲方提供的 ≥30 天保存证据。",
        "- 不重新测绘、不制作三维模型，不测试宿主 APP 原生功能；只测试甲方服务加载、叠加和 H5 嵌入适配。",
        "- 不把第三方中台、定位、门禁、消防、入侵和物联网内部实现纳入白盒测试，只验证适配契约、授权边界、错误、降级和审计。",
        "- `ISSUE-G3-01-001` 关闭前，PE-04、PE-06 以及视频/消息真实联调只能保持 BLOCKED/PENDING，不能用桩测结果替代。",
        "",
        "## 3 测试策略",
        "",
        "### 3.1 测试左移与证据链",
        "",
        "[FIGURE]",
        "",
        "图 3-1 四级测试与验收证据链",
        "",
        "设计评审即检查 AC 的可测试性、设计 ID 落点和失败语义；开发时先建立可失败的单元、契约或集成断言；系统测试按 RTM 运行端到端场景；验收阶段只接收具有版本、环境、原始结果和缺陷闭环的证据。",
        "",
        "### 3.2 分级策略",
        "",
        "表 3-1 测试策略与执行规则",
        "",
        "| 策略 | 执行方法 | 证据 | 失败处置 |",
        "| --- | --- | --- | --- |",
        "| 单元与属性测试 | 对状态转移、幂等键、时限、权限和统计规则覆盖正常、边界、非法输入 | 测试日志、覆盖率、断言明细 | 阻断合并，修复后回归 |",
        "| 契约测试 | 解析 OpenAPI，校验 schema、必填字段、错误码、traceId、版本兼容和破坏性差异 | 校验报告、契约快照、差异报告 | 破坏性变更进入 Change/ADR |",
        "| 数据与迁移测试 | 在空库、升级库和回滚路径执行迁移，核对约束、索引、审计和引用完整性 | 迁移日志、校验 SQL、摘要 | 阻断部署候选 |",
        "| 集成与故障注入 | 对 8 个 EXT 端口覆盖成功、无权、超时、断连、重复和未知载荷 | 请求响应/回执、适配器日志、告警 | 进入重试、降级或人工处置并建缺陷 |",
        "| 系统与回归 | 按角色从 Web/H5 执行 117 个 AC 级场景，变更影响集纳入回归 | 截图/录像索引、业务记录、审计、结果表 | 失败 AC 标红并阻断准出 |",
        "| 性能与容量 | 采用受控数据、预热、重复轮次，记录 P50/P95/P99、吞吐、错误率、CPU/内存/连接 | 脚本、原始数据、资源曲线、报告 | 未达指标不得以平均值掩盖 |",
        "| 安全测试 | 执行 SCA、漏洞扫描、鉴权、越权、输入、导出、日志脱敏和高危复测 | 扫描报告、请求证据、复测记录 | 高危未清零禁止发布 |",
        "| 恢复与兼容 | 演练进程/数据库/适配器故障和备份恢复；按甲方确认矩阵测试浏览器/H5 宿主 | 演练记录、恢复摘要、兼容矩阵 | 恢复失败或矩阵未确认不得准出对应项 |",
        "",
        "### 3.3 测试环境与数据",
        "",
        "表 3-2 环境、数据与隔离规则",
        "",
        "| 环境 | 用途 | 数据 | 约束 |",
        "| --- | --- | --- | --- |",
        "| CI | 单元、静态检查、OpenAPI 和迁移检查 | 合成最小数据、桩外部端口 | 不产生真实联调或性能结论 |",
        "| 集成环境 | 数据库、Outbox、模块和适配器故障注入 | 可重复生成的脱敏/合成数据 | 每次执行固定版本和初始化脚本 |",
        "| 系统测试环境 | 117 AC、Web/H5、地图和业务闭环 | 受控角色、事件、任务、站点、物资、打卡与演练数据 | 与目标配置差异必须记录 |",
        "| 目标/验收环境 | PE、外部联调、兼容和验收支持 | 甲方确认的数据与账号 | 账号、窗口、接口版本和合法数据权由甲方提供/协调 |",
        "",
        "### 3.4 缺陷与回归",
        "",
        "缺陷至少记录严重度、影响 AC/设计 ID、环境、步骤、期望、实际、证据、责任人、修复版本和复测结论。致命或严重缺陷、★失败、RTM 断链、高危安全问题、恢复失败均阻断候选；修复必须回归直接 AC、同模块关联 AC 和受影响契约。",
        "",
        "## 4 准入与准出",
        "",
        "表 4-1 分级准入与准出门禁",
        "",
        "| 阶段 | 准入条件 | 准出条件 | 阻断项 |",
        "| --- | --- | --- | --- |",
        "| 设计测试准备 | G3-03—08 DONE，117 AC 可追踪 | 测试计划 REVIEW，TC 编号无缺失 | AC/设计断链 |",
        "| 单元/契约 | 代码、迁移、OpenAPI 有版本；测试可重复 | 用例全绿，核心覆盖率≥70%，OpenAPI 零 schema error | 核心失败、契约破坏 |",
        "| 集成 | 单元/契约准出，依赖版本和桩口径明确 | 跨模块链路全绿，重试/降级/审计可观察 | 数据不一致、不可恢复副作用 |",
        "| 系统 | 候选可部署，39 FR/117 AC 数据就绪 | 117 AC 100% 有结果；34★全部通过；致命/严重缺陷清零 | 缺证据、★失败、严重缺陷 |",
        "| 专项 | 主链路稳定，环境/规模/矩阵已确认 | PE/NFR 逐项有原始证据；安全高危清零；恢复演练通过 | 环境不等价、指标失败、高危、恢复失败 |",
        "| 验收支持 | 系统与专项准出，文档/RTM/缺陷一致 | 甲方签认或明确遗留与偏差，不静默豁免 | ISSUE-G3-01-001 未按真实证据或书面裁决关闭 |",
        "",
        "## 5 进度、角色与交付",
        "",
        "### 5.1 进度安排",
        "",
        "表 5-1 测试活动进度",
        "",
        "| 活动 | 工作日 | 主责 | 产出 |",
        "| --- | --- | --- | --- |",
        "| 计划、TC 编号与门禁冻结 | D7—D8 | C，A 单人复核 | 本计划、RTM 测试挂接候选 |",
        "| 单元、契约和迁移测试 | D9—D15 随开发 | B/功能实现人 | CI、覆盖率、契约和迁移报告 |",
        "| 集成与系统测试 | D15—D16 | C 组织，B 支持 | 117 AC 结果、接口与数据证据 |",
        "| 性能、安全、兼容与恢复专项 | D16—D17 | C 组织，A/B 支持 | 专项报告、扫描和演练记录 |",
        "| 回归、缺陷收口与验收包 | D17 | 全组 | 回归结果、缺陷看板、RTM v4 候选 |",
        "",
        "### 5.2 角色分工",
        "",
        "表 5-2 测试职责",
        "",
        "| 角色 | 职责 | 不得替代的复核 |",
        "| --- | --- | --- |",
        "| C | 维护计划、TC/RTM、测试数据、执行组织、结果与符合性 | A 对 G3-09 准出与计划一致性的单人复核 |",
        "| B | 提供环境、脚本、桩、可观测性，支持性能/接口/恢复并修复缺陷 | 本轮提供技术支持，不作为 G3-09 复核人 |",
        "| A | 协调版本、窗口、资源、评审和验收包，裁定计划偏差 | 按 OVR-025 独立完成 G3-09 正式复核 |",
        "| 甲方/教师模拟角色 | 提供或确认外部环境、账号、数据、兼容矩阵和验收窗口 | 乙方不能代造其真实环境证据 |",
        "",
        "## 6 P0 场景与外部联动",
        "",
        "表 6-1 P0 场景设计",
        "",
        "| 场景 | 正常路径 | 无权/非法路径 | 超时/失败路径 | 主要证据 |",
        "| --- | --- | --- | --- | --- |",
        "| P0-01 事件核实至任务下达 | 核实、启动预案、创建任务和通知 | 无启动权限被拒绝并审计 | Outbox/消息失败进入重试或人工处置 | 同一 eventId/traceId、状态与时间戳 |",
        "| P0-02 消息可靠推送 | 正常验收通道 ≥20 路且到达率 ≥99% | 非授权对象不下发 | 超时、迟到和重复回执幂等，失败转人工 | 请求、尝试、回执、终态统计 |",
        "| P0-03 KN-064 四系统联动 | 视频、信息发布、物联网、中台沿同一事件链完成调用 | 无权视频/发布/中台访问分别拒绝 | 任一系统超时/失败均隔离、告警、降级 | 同一 eventId/traceId 的四系统证据 |",
        "| P0-04 视频调阅与回放 | 有权实时调阅、历史回放和首帧计时 | 无视频权限拒绝 | 接口超时明确失败，不承担录像存储 | 脱敏请求/响应、首帧、保存责任证据 |",
        "| P0-05 门禁安全联动 | 授权确认后发指令并接收联锁结果 | 无权或缺少二次确认不下发 | 拒绝/超时/联锁失败告警并人工降级 | 指令、令牌、回执和审计 |",
        "| P0-06 H5 事件与任务 | 事件、附件、接收、反馈和完成闭环 | 越权事件/任务不可见或不可操作 | 网络/附件失败可恢复且不伪报成功 | H5 状态、服务记录、附件引用 |",
        "| P0-07 扫码打卡 | 有效二维码、时段、半径和身份写入一次 | 越界、过期、错误小组拒绝 | 重复扫码幂等，位置源不可用明确降级 | 打卡记录、幂等键、耗时和审计 |",
        "| P0-08 盘点快照 | 按发布快照比对并复核差异 | 无权更新拒绝 | 并发出入库单列，失败不覆盖原快照 | 快照、现场值、期间变动和调整记录 |",
        "",
        "## 7 非功能与关键数字验证",
        "",
        "表 7-1 PE-01—PE-12 测试设计",
        "",
        "| 指标 | 目标 | 方法与统计口径 | 当前门禁 |",
        "| --- | --- | --- | --- |",
        "| PE-01 | 预案启动通知与任务下发≤3秒 | 同一 traceId，从授权启动到可追踪下发记录 | 待实现/待测 |",
        "| PE-02 | 告警接入响应≤2秒 | 源产生、接收、持久化时间戳和异常率 | 待目标环境实测 |",
        "| PE-03 | 事件确认至任务下达≤3分钟 | 核实通过到任务下发记录，含失败样本 | 待实现/待测 |",
        "| PE-04 | 消息并行≥20路、到达率≥99% | 批量发送、终态回执、迟到/重复/失败分类 | ISSUE-G3-01-001 阻断实测 |",
        "| PE-05 | 定位刷新≤2秒且不降低源精度 | 比对源、接收、投影、展示时间与精度 | 待甲方定位源 |",
        "| PE-06 | 视频首帧≤3秒、既有保存≥30天 | 授权调阅到首帧；保存期查甲方既有证据 | ISSUE-G3-01-001 阻断实测 |",
        "| PE-07 | Web 普通页面 P95≤3秒 | 目标规模下记录 P50/P95/P99、错误率和资源 | 待压测 |",
        "| PE-08 | 地图常规操作≤2秒 | 典型缩放、平移、点选端到端计时 | 待合法地图服务 |",
        "| PE-09 | 扫码写入/回传≤1秒 | 有效、过期、越界、重复分别计时 | 待压测 |",
        "| PE-10 | 安防刷新≤30秒、信息刷新≤60秒 | 源事件到投影可见时间和配置核对 | 待联调/待测 |",
        "| PE-11 | 峰值在线用户≥100 | ≥100 用户下成功率、P95/P99、错误与资源 | 待压测 |",
        "| PE-12 | 一年期统计报表≤5秒 | 一年期受控数据、授权过滤和结果一致性 | 待压测 |",
        "",
        "表 7-2 其他质量特性测试",
        "",
        "| 类别 | 受控对象 | 方法 | 准出证据 |",
        "| --- | --- | --- | --- |",
        "| 容量 | NFR-CAP-01—04 | 生成受控规模数据，执行导入、查询、统计、归档和恢复 | 数据生成脚本、规模摘要、结果与资源 |",
        "| 可靠性 | NFR-REL-01—03 | 可用率口径、进程/数据库/适配器故障、备份恢复 | 监控、演练记录、恢复数量/关系/摘要 |",
        "| 安全 | NFR-SEC-01—02 | SCA、漏洞、鉴权、越权、导出、脱敏、高危复测 | 扫描/渗透结果，高危为 0 |",
        "| 兼容/易用 | NFR-COMP-01、NFR-USE-01 | 甲方确认浏览器/H5 宿主矩阵，关键任务走查和培训考核 | 兼容矩阵、步骤/跳转数、考核记录 |",
        "| 维护/移植 | NFR-MNT-01—02、NFR-PORT-01 | 覆盖率、干净环境部署≤2小时、配置/代码分离和回滚 | CI、部署计时、配置清单、回滚记录 |",
        "",
        "## 8 117 个 AC 级测试设计索引",
        "",
        "本章锁定测试编号、目标和证据类型。详细前置数据、步骤和断言在实现阶段进入测试用例库；编号不得脱离 AC 单独重命名。",
        "",
        "表 8-1 FR—AC—测试设计索引",
        "",
        "| FR / ★ | TC 与 AC 目标 | 建议级别 | 计划证据 |",
        "| --- | --- | --- |",
    ]
    for req in reqs:
        acs = req["acs"]
        lines.append(
            f"| {req['id']} / {req['star']} | "
            + "<br>".join(f"{ac.replace('AC-', 'TC-')}：{summary}" for ac, summary in acs)
            + f" | {level_for(req['num'])} | 结果记录、业务/审计数据、必要的请求回执 |"
        )
    lines += [
        "",
        "## 9 风险与应对",
        "",
        "表 9-1 测试风险与处置",
        "",
        "| 风险 | 触发条件 | 应对 | 对准出的影响 |",
        "| --- | --- | --- | --- |",
        "| 外部环境或账号不到位 | 视频、消息、中台等无法取得真实调用条件 | 提前登记窗口和所需证据；使用桩只验证内部逻辑；保留 OPEN Issue | 对应联调/PE 不得通过，阻断 M3 |",
        "| 测试数据与目标规模不等价 | 数据量、角色或空间服务不足 | 记录生成规则和差异；补测目标环境 | 性能/容量结论无效 |",
        "| 用例与需求漂移 | FR/AC/设计变更未同步 TC | CI 检查孤儿/缺失 ID，变更评审更新 RTM | 追踪断链阻断发布 |",
        "| 缺陷积压压缩回归 | 致命/严重缺陷临近准出仍开放 | 按严重度停线，优先修复并执行影响集回归 | 致命/严重未清零不准出 |",
        "| 兼容矩阵未确认 | Android/iOS 或浏览器版本未由甲方确认 | 先测现有开发环境并标注参考；确认后重跑 | 不冻结兼容结论 |",
        "| 性能脚本或监控不可信 | 结果缺少时间、资源或原始数据 | 脚本版本化，校准计时，保留原始输出 | 无原始证据的指标不判通过 |",
        "",
        "## 10 配置、报告与准出结论",
        "",
        "测试配置项包括计划、用例、数据生成脚本、环境清单、OpenAPI 快照、执行结果、覆盖率、性能原始数据、安全扫描、故障演练、缺陷和 RTM。测试报告必须按执行版本汇总通过、失败、阻断和未执行项，禁止只给总通过率而省略★、PE、外部接口和开放 Issue。",
        "",
        "**G3-09 主责结论**：已形成测试级别、范围、策略、进度、环境、准入准出、P0 场景、PE/NFR 与 117 个 AC 级测试编号，可按 OVR-025 提交 A 单人复核。当前状态保持 REVIEW；本结论不表示任何测试已经执行或通过。`ISSUE-G3-01-001` 保持 OPEN，并继续阻断视频/消息真实联调结论和 G3-10/M3 最终冻结。",
        "",
        "## 附录 A 证据包最小字段",
        "",
        "表 A-1 单次执行证据字段",
        "",
        "| 字段 | 要求 |",
        "| --- | --- |",
        "| 标识 | TC、AC、FR、设计 ID、缺陷 ID 可回指 |",
        "| 版本 | commit、构建、数据库迁移、OpenAPI/配置版本 |",
        "| 环境 | 主机/容器、浏览器/H5 宿主、外部接口版本、时间同步 |",
        "| 数据 | 数据集版本、规模、账号/角色、脱敏与清理方式 |",
        "| 执行 | 操作者、开始/结束时间、步骤或脚本、重复轮次 |",
        "| 结果 | 期望、实际、原始输出、统计口径、截图/日志/回执索引 |",
        "| 结论 | PASS/FAIL/BLOCKED/NOT_RUN 之一；失败关联缺陷与复测 |",
    ]
    return "\n".join(lines) + "\n"


def set_font(run, east: str = "宋体", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr, value in (("eastAsia", east), ("ascii", "Times New Roman"), ("hAnsi", "Times New Roman")):
        fonts.set(qn(f"w:{attr}"), value)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_paragraph(p, text: str, size: float | None = None, bold: bool | None = None) -> None:
    ppr = deepcopy(p._p.pPr) if p._p.pPr is not None else None
    for child in list(p._p):
        p._p.remove(child)
    if ppr is not None:
        p._p.insert(0, ppr)
    set_font(p.add_run(text), size=size, bold=bold)


def remove_reference_body(doc: Document) -> None:
    start = next((p for p in doc.paragraphs if p.text.strip().startswith("1 引言")), None)
    if start is None:
        raise RuntimeError("Reference body start not found")
    body = doc.element.body
    children = list(body)
    start_idx = children.index(start._p)
    for child in children[start_idx:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def replace_toc(doc: Document) -> None:
    body = doc.element.body
    sdts = list(body.xpath("./w:sdt"))
    if not sdts:
        raise RuntimeError("Reference TOC content control not found")
    toc = sdts[0]
    entries = [
        (1, "1 引言", 1), (2, "1.1 编写目的", 1), (2, "1.2 测试依据", 1), (2, "1.3 证据纪律", 1),
        (1, "2 测试范围", 1), (2, "2.1 测试对象与基线", 1), (2, "2.2 测试级别", 1),
        (2, "2.3 覆盖规则", 2), (2, "2.4 不测与延期范围", 2),
        (1, "3 测试策略", 2), (2, "3.1 测试左移与证据链", 2), (2, "3.2 分级策略", 3),
        (2, "3.3 测试环境与数据", 3), (2, "3.4 缺陷与回归", 3),
        (1, "4 准入与准出", 3), (1, "5 进度、角色与交付", 4),
        (2, "5.1 进度安排", 4), (2, "5.2 角色分工", 4),
        (1, "6 P0 场景与外部联动", 4), (1, "7 非功能与关键数字验证", 5),
        (1, "8 117 个 AC 级测试设计索引", 5), (1, "9 风险与应对", 7),
        (1, "10 配置、报告与准出结论", 8), (1, "附录 A 证据包最小字段", 8),
    ]
    title = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr"); style = OxmlElement("w:pStyle"); style.set(qn("w:val"), "TOCHeading"); ppr.append(style)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc); title.append(ppr)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "目 录"; r.append(t); title.append(r); toc.addprevious(title)
    for level, text, page in entries:
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle"); style.set(qn("w:val"), f"TOC{level}"); ppr.append(style)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot"); tab.set(qn("w:pos"), "8500")
        tabs.append(tab); ppr.append(tabs); p.append(ppr)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = text; r.append(t); p.append(r)
        rt = OxmlElement("w:r"); rt.append(OxmlElement("w:tab")); p.append(rt)
        rp = OxmlElement("w:r"); tp = OxmlElement("w:t"); tp.text = str(page); rp.append(tp); p.append(rp)
        toc.addprevious(p)
    body.remove(toc)


def set_cell_text(cell, text: str, size: float, bold: bool = False) -> None:
    cell.text = ""
    parts = text.split("<br>")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for i, part in enumerate(parts):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(clean(part)), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0


def format_table(table, rows: list[list[str]]) -> None:
    cols = len(rows[0])
    size = 8.0 if len(rows) > 20 else (7.5 if cols >= 5 else 8.5)
    table.autofit = False
    total = 9072
    weights = {2: [1, 3], 3: [1, 2, 3], 4: ([1.0, 5.5, 1.5, 2.0] if len(rows) > 20 else [1.2, 2.2, 2.5, 2.3]), 5: [1.2, 1.5, 3.6, 1.5, 2.2]}.get(cols, [1] * cols)
    widths = [int(total * w / sum(weights)) for w in weights]
    grid = table._tbl.tblGrid
    for gc, width in zip(grid.gridCol_lst, widths):
        gc.set(qn("w:w"), str(width))
    for ri, row in enumerate(table.rows):
        if ri == 0:
            trpr = row._tr.get_or_add_trPr(); rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); trpr.append(rep)
        for ci, cell in enumerate(row.cells):
            tcpr = cell._tc.get_or_add_tcPr(); tcw = tcpr.first_child_found_in("w:tcW")
            if tcw is None: tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
            tcw.set(qn("w:w"), str(widths[ci])); tcw.set(qn("w:type"), "dxa")
            set_cell_text(cell, rows[ri][ci], size, bold=ri == 0)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(text), east="黑体", size=9, bold=True)


def parse_table(lines: list[str], i: int) -> tuple[list[list[str]], int]:
    rows = []
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [x.strip() for x in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", x) for x in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    table_style = doc.tables[0].style
    pending_caption = None
    i = 1
    while i < len(lines):
        raw = lines[i].strip()
        if raw.startswith("- ") and i < 10:
            i += 1; continue
        if raw.startswith("## "):
            p = doc.add_paragraph(style="Heading 1"); p.paragraph_format.keep_with_next = True
            set_font(p.add_run(clean(raw[3:])), east="黑体", bold=True); i += 1; continue
        if raw.startswith("### "):
            p = doc.add_paragraph(style="Heading 2"); p.paragraph_format.keep_with_next = True
            set_font(p.add_run(clean(raw[4:])), east="黑体", bold=True); i += 1; continue
        if re.match(r"^表\s+[0-9A-Z]+-\d+\s+", raw):
            pending_caption = raw; i += 1; continue
        if raw == "[FIGURE]":
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str(FIGURE), width=Cm(15.2))
            shape._inline.docPr.set("title", "图 3-1 四级测试与验收证据链")
            shape._inline.docPr.set("descr", "展示单元测试、集成测试、系统测试到验收测试的四级证据递进关系。")
            i += 1; continue
        if re.match(r"^图\s+\d+-\d+\s+", raw):
            add_caption(doc, raw); i += 1; continue
        if raw.startswith("|"):
            rows, i = parse_table(lines, i)
            if pending_caption:
                add_caption(doc, pending_caption); pending_caption = None
            table = doc.add_table(rows=1, cols=len(rows[0])); table.style = table_style
            for ci, text in enumerate(rows[0]): table.rows[0].cells[ci].text = text
            for values in rows[1:]:
                cells = table.add_row().cells
                for ci, text in enumerate(values): cells[ci].text = text
            format_table(table, rows); continue
        if not raw or raw.startswith("# "):
            i += 1; continue
        style = "List Paragraph" if re.match(r"^(?:\d+\.|- )", raw) else "Normal"
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.25
        if style == "Normal": p.paragraph_format.first_line_indent = Pt(24); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        set_font(p.add_run(clean(re.sub(r"^- ", "", raw))), size=10.5)
        i += 1


def build_docx(markdown: str) -> None:
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    remove_reference_body(doc)
    replace_toc(doc)
    cover = {
        0: "文档编号：YJGL-G3-09　　版本号：V0.2",
        1: "密　　级：内部 · 教学用",
        3: "某自然博物馆智能运营中心建设项目——应急管理子系统",
        4: "测试计划",
        5: "（G3-09 评审候选）",
        7: "编制单位：020202项目组【待人工确认】",
        8: "编　　制：C（需求与测试负责人）【待人工确认姓名】",
        9: "审　　核：A（OVR-025 单人复核）",
        10: "批　　准：【待人工确认】",
        11: "编制日期：2026 年 9 月 19 日",
        13: "修订记录",
        14: "注：本表只记录可核验版本事件；执行结果将在后续测试报告和 RTM v4 中形成。",
    }
    for idx, text in cover.items(): replace_paragraph(doc.paragraphs[idx], text)
    rev = doc.tables[0]
    while len(rev.rows) > 1: rev._tbl.remove(rev.rows[-1]._tr)
    revisions = [
        ["V0.1", "2026-09-19", "全文", "建立 G3-09 测试计划评审候选，锁定 117 个 AC 级测试编号和测试左移门禁。", "C"],
        ["V0.2", "2026-09-19", "封面、目录、5、8、10", "按 A 审核修正复核主体、目录导航、长表可读性和图 3-1 替代文本。", "C"],
    ]
    for vals in revisions:
        row = rev.add_row().cells
        for cell, val in zip(row, vals): cell.text = val
    format_table(rev, [["版本", "日期", "修订章节", "修订说明", "编制/修订人"], *revisions])
    add_body(doc, markdown)
    body = doc.element.body
    sect = body.sectPr
    body.remove(sect); body.append(sect)
    doc.core_properties.title = "测试计划"
    doc.core_properties.subject = "某自然博物馆智能运营中心建设项目——应急管理子系统 G3-09"
    doc.core_properties.author = "020202项目组"
    doc.save(OUTPUT)


def main() -> int:
    reqs = parse_requirements()
    markdown = build_markdown(reqs)
    WORK.write_text(markdown, encoding="utf-8")
    build_docx(markdown)
    print(WORK)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
