from __future__ import annotations

import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_g3_09_test_plan as base

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "docs/work/A_PM/project_management_plan.md"
REFERENCE = ROOT / "docs/reference/22-项目管理计划（教学样例）.docx"
OUTPUT = ROOT / "docs/deliverables/20-项目管理计划.docx"
FIGURE = ROOT / "docs/deliverables/figures/20-图4-1-G3至M3里程碑与门禁.png"


def replace_toc(doc: Document) -> None:
    body = doc.element.body
    sdts = list(body.xpath("./w:sdt"))
    if not sdts:
        raise RuntimeError("Reference TOC content control not found")
    toc = sdts[0]
    entries = [
        (1, "1 引言", 1), (2, "1.1 编写目的", 1), (2, "1.2 项目概况", 1), (2, "1.3 管理原则", 1),
        (1, "2 范围管理", 1), (2, "2.1 范围基线", 1), (2, "2.2 范围控制", 1), (2, "2.3 交付物与工作分解", 1),
        (1, "3 组织与资源管理", 2), (2, "3.1 组织与职责", 2), (2, "3.2 资源与 WIP", 2),
        (1, "4 进度与里程碑管理", 2), (2, "4.1 里程碑与依赖", 2), (2, "4.2 进度控制", 3), (2, "4.3 进度缓冲", 3),
        (1, "5 沟通与决策管理", 3), (1, "6 质量、配置与交付管理", 3),
        (2, "6.1 质量门禁", 3), (2, "6.2 配置与基线", 3), (2, "6.3 验收与交付", 4),
        (1, "7 偏差、问题与变更管理", 4), (2, "7.1 偏差分级", 4), (2, "7.2 偏差处置流程", 4),
        (2, "7.3 任务催办与临时任务专项规则", 4), (1, "8 风险与问题管理", 4),
        (1, "9 监控、报告与收尾", 5), (1, "10 本计划准出检查", 5),
    ]
    title = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr"); style = OxmlElement("w:pStyle"); style.set(qn("w:val"), "TOCHeading"); ppr.append(style)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc); title.append(ppr)
    r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = "目 录"; r.append(t); title.append(r); toc.addprevious(title)
    for level, text, page in entries:
        p = OxmlElement("w:p"); ppr = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle"); style.set(qn("w:val"), f"TOC{level}"); ppr.append(style)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot"); tab.set(qn("w:pos"), "8500"); tabs.append(tab); ppr.append(tabs); p.append(ppr)
        r = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = text; r.append(t); p.append(r)
        rt = OxmlElement("w:r"); rt.append(OxmlElement("w:tab")); p.append(rt)
        rp = OxmlElement("w:r"); tp = OxmlElement("w:t"); tp.text = str(page); rp.append(tp); p.append(rp)
        toc.addprevious(p)
    body.remove(toc)


def build() -> None:
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    base.remove_reference_body(doc)
    replace_toc(doc)
    cover = {
        0: "文档编号：YJGL-G3-11　　版本号：V0.1",
        1: "密　　级：内部 · 教学用",
        3: "某自然博物馆智能运营中心建设项目——应急管理子系统",
        4: "项目管理计划",
        5: "（G3-11 评审候选）",
        7: "编制单位：020202项目组【待人工确认】",
        8: "编　　制：何思源（项目经理 / 总编）",
        9: "审　　核：任俊强（需求与质量负责人）",
        10: "批　　准：【待人工确认】",
        11: "编制日期：2026 年 9 月 19 日",
        13: "修订记录",
        14: "注：本文件为评审候选；通过 C 复核后方可由 REVIEW 转为 DONE。",
    }
    for idx, text in cover.items():
        base.replace_paragraph(doc.paragraphs[idx], text)
    rev = doc.tables[0]
    while len(rev.rows) > 1:
        rev._tbl.remove(rev.rows[-1]._tr)
    revisions = [["V0.1", "2026-09-19", "全文", "建立 G3-11 项目管理计划候选，形成范围、资源、里程碑、偏差与 AC-G2-FR-015-03 管理抽查规则。", "何思源"]]
    row = rev.add_row().cells
    for cell, value in zip(row, revisions[0]): cell.text = value
    base.format_table(rev, [["版本", "日期", "修订章节", "修订说明", "编制/修订人"], *revisions])

    base.FIGURE = FIGURE
    base.add_body(doc, WORK.read_text(encoding="utf-8"))
    body = doc.element.body; sect = body.sectPr; body.remove(sect); body.append(sect)
    doc.core_properties.title = "项目管理计划"
    doc.core_properties.subject = "某自然博物馆智能运营中心建设项目——应急管理子系统 G3-11"
    doc.core_properties.author = "020202项目组"
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header.paragraphs:
                base.replace_paragraph(header.paragraphs[0], "某自然博物馆智能运营中心建设项目——应急管理子系统 · 项目管理计划", size=9)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
