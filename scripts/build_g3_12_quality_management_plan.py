from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_g3_09_test_plan as base


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "docs/work/C_REQ/quality_management_plan.md"
REFERENCE = ROOT / "docs/reference/23-质量管理计划（教学样例）.docx"
OUTPUT = ROOT / "docs/deliverables/21-质量管理计划.docx"


def remove_reference_body(doc: Document) -> None:
    start = next((p for p in doc.paragraphs if p.text.strip().startswith("1 ") and p.style and p.style.name == "Heading 1"), None)
    if start is None:
        raise RuntimeError("Reference body start not found")
    body = doc.element.body
    children = list(body)
    start_idx = children.index(start._p)
    for child in children[start_idx:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def replace_toc(doc: Document) -> None:
    body = doc.element.body
    sdts = list(body.xpath("./w:sdt"))
    if not sdts:
        raise RuntimeError("Reference TOC content control not found")
    toc = sdts[0]
    entries = [
        (1, "1 引言", 1), (2, "1.1 编写目的", 1), (2, "1.2 适用范围", 1), (2, "1.3 依据与优先级", 1),
        (1, "2 质量目标与受控对象", 1), (1, "3 质量组织与职责", 2),
        (1, "4 质量保证活动（QA）", 2), (1, "5 质量控制与缺陷管理（QC）", 2),
        (2, "5.1 检查与测试控制", 2), (2, "5.2 缺陷分级", 2), (2, "5.3 缺陷生命周期", 2),
        (1, "6 评审与质量门禁", 3), (1, "7 质量度量与报告", 3),
        (1, "8 PE-01—PE-12 质量证据控制", 3), (1, "9 文档、AI 与证据质量", 4),
        (2, "9.1 正式文档质量", 4), (2, "9.2 AI 辅助质量", 4), (2, "9.3 证据最小字段", 4),
        (1, "10 不符合、纠正与持续改进", 4), (1, "11 开放事项与本计划准出", 4),
    ]
    for level, text, page in entries:
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle")
        style.set(qn("w:val"), f"TOC{level}")
        ppr.append(style)
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:leader"), "dot")
        tab.set(qn("w:pos"), "8500")
        tabs.append(tab)
        ppr.append(tabs)
        p.append(ppr)
        run = OxmlElement("w:r")
        value = OxmlElement("w:t")
        value.text = text
        run.append(value)
        p.append(run)
        tab_run = OxmlElement("w:r")
        tab_run.append(OxmlElement("w:tab"))
        p.append(tab_run)
        page_run = OxmlElement("w:r")
        page_text = OxmlElement("w:t")
        page_text.text = str(page)
        page_run.append(page_text)
        p.append(page_run)
        toc.addprevious(p)
    body.remove(toc)


def build() -> None:
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    remove_reference_body(doc)
    replace_toc(doc)

    cover = {
        0: "文档编号：YJGL-G3-12-QMP　　版本号：V0.1",
        1: "密　　级：内部 · 教学用",
        3: "某自然博物馆智能运营中心建设项目——应急管理子系统",
        4: "质量管理计划",
        5: "（G3-12 评审候选）",
        7: "编制单位：020202项目组【待人工确认】",
        8: "编　　制：任俊强（需求与质量负责人）",
        9: "审　　核：严宇（技术负责人）",
        10: "批　　准：【待人工确认】",
        11: "编制日期：2026 年 9 月 20 日",
        13: "修订记录",
        14: "注：本表只记录可核验的版本事件；B 独立复核通过前，本文件保持评审候选。",
    }
    for index, text in cover.items():
        base.replace_paragraph(doc.paragraphs[index], text)

    revision = doc.tables[0]
    while len(revision.rows) > 1:
        revision._tbl.remove(revision.rows[-1]._tr)
    revisions = [[
        "V0.1", "2026-09-20", "全文",
        "建立 G3-12 质量管理计划评审候选，规定评审、缺陷、度量、门禁及 PE-01—PE-12 证据口径。",
        "任俊强",
    ]]
    row = revision.add_row().cells
    for cell, value in zip(row, revisions[0]):
        cell.text = value
    base.format_table(revision, [["版本", "日期", "修订章节", "修订说明", "编制/修订人"], *revisions])

    base.add_body(doc, WORK.read_text(encoding="utf-8"))
    body = doc.element.body
    sect = body.sectPr
    body.remove(sect)
    body.append(sect)

    doc.core_properties.title = "质量管理计划"
    doc.core_properties.subject = "某自然博物馆智能运营中心建设项目——应急管理子系统 G3-12"
    doc.core_properties.author = "020202项目组"
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header.paragraphs:
                base.replace_paragraph(header.paragraphs[0], "某自然博物馆智能运营中心建设项目——应急管理子系统 · 质量管理计划", size=9)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
