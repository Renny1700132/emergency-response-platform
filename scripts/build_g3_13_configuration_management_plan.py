from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import build_g3_09_test_plan as base


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "docs/work/A_PM/configuration_management_plan.md"
REFERENCE = ROOT / "docs/reference/24-配置管理计划（教学样例）.docx"
OUTPUT = ROOT / "docs/deliverables/22-配置管理计划.docx"


def remove_reference_body(doc: Document) -> None:
    start = next((p for p in doc.paragraphs if p.text.strip().startswith("1 配置管理范围与组织")), None)
    if start is None:
        raise RuntimeError("Reference body start not found")
    body = doc.element.body
    children = list(body)
    start_idx = children.index(start._p)
    for child in children[start_idx:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def replace_toc(doc: Document) -> None:
    body = doc.element.body
    sdts = list(body.xpath("./w:sdt"))
    if not sdts:
        raise RuntimeError("Reference TOC content control not found")
    toc = sdts[0]
    entries = [
        (1, "1 目的、范围与依据", 1), (2, "1.1 目的", 1), (2, "1.2 适用范围", 1), (2, "1.3 依据与优先级", 1),
        (1, "2 组织、职责与职责分离", 2), (1, "3 配置项识别、登记与状态", 2),
        (2, "3.1 配置项类别与清单", 2), (2, "3.2 配置项状态", 3),
        (1, "4 版本、分支、提交与制品规则", 3), (2, "4.1 版本标识", 3), (2, "4.2 Git 与提交", 3), (2, "4.3 制品完整性", 3),
        (1, "5 基线建立、冻结与解除", 3), (2, "5.1 基线类型与内容", 3), (2, "5.2 冻结步骤", 3),
        (1, "6 变更、CR 与 CCB 流程", 3), (2, "6.1 CR 最小字段", 3), (2, "6.2 流程与裁决", 3),
        (1, "7 权限、安全与配置审计", 4), (2, "7.1 权限控制", 4), (2, "7.2 审计类型", 4),
        (1, "8 发布、部署与回滚", 4), (2, "8.1 发布准入", 4), (2, "8.2 发布与回滚步骤", 4),
        (1, "9 配置库、备份与归档", 4), (2, "9.1 逻辑结构", 4), (2, "9.2 备份、保留与归档", 5),
        (1, "10 AC-G2-FR-013-01 配置抽查", 5), (1, "11 配置状态报告、例外与本计划准出", 5),
    ]
    for level, text, page in entries:
        p = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle"); style.set(qn("w:val"), f"TOC{level}"); ppr.append(style)
        tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right"); tab.set(qn("w:leader"), "dot"); tab.set(qn("w:pos"), "8500")
        tabs.append(tab); ppr.append(tabs); p.append(ppr)
        run = OxmlElement("w:r"); value = OxmlElement("w:t"); value.text = text; run.append(value); p.append(run)
        tab_run = OxmlElement("w:r"); tab_run.append(OxmlElement("w:tab")); p.append(tab_run)
        page_run = OxmlElement("w:r"); page_text = OxmlElement("w:t"); page_text.text = str(page); page_run.append(page_text); p.append(page_run)
        toc.addprevious(p)
    body.remove(toc)


def build() -> None:
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    remove_reference_body(doc)
    replace_toc(doc)
    cover = {
        0: "文档编号：YJGL-G3-13-CMP　　版本号：V0.1",
        1: "密　　级：内部 · 教学用",
        3: "某自然博物馆智能运营中心建设项目——应急管理子系统",
        4: "配置管理计划",
        5: "（G3-13 评审候选）",
        7: "编制单位：020202项目组【待人工确认】",
        8: "编　　制：何思源（项目负责人 / 配置负责人）",
        9: "审　　核：任俊强（需求与质量负责人）",
        10: "批　　准：【待人工确认】",
        11: "编制日期：2026 年 9 月 20 日",
        14: "修订记录",
        15: "注：本表只记录可核验的版本事件；C 独立复核通过前，本文件保持评审候选。",
    }
    for idx, text in cover.items():
        base.replace_paragraph(doc.paragraphs[idx], text)
    revision = doc.tables[0]
    while len(revision.rows) > 1:
        revision._tbl.remove(revision.rows[-1]._tr)
    revisions = [["V0.1", "2026-09-20", "全文", "建立 G3-13 配置管理计划评审候选，规定配置项、版本、基线、CR/CCB、权限审计、发布回滚与归档，并抽查 AC-G2-FR-013-01。", "何思源"]]
    row = revision.add_row().cells
    for cell, value in zip(row, revisions[0]):
        cell.text = value
    base.format_table(revision, [["版本", "日期", "修订章节", "修订说明", "编制/修订人"], *revisions])
    base.add_body(doc, WORK.read_text(encoding="utf-8"))
    body = doc.element.body; sect = body.sectPr; body.remove(sect); body.append(sect)
    doc.core_properties.title = "配置管理计划"
    doc.core_properties.subject = "某自然博物馆智能运营中心建设项目——应急管理子系统 G3-13"
    doc.core_properties.author = "020202项目组"
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header.paragraphs:
                base.replace_paragraph(header.paragraphs[0], "某自然博物馆智能运营中心建设项目——应急管理子系统 · 配置管理计划", size=9)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
    print(OUTPUT)
