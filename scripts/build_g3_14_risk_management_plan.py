from __future__ import annotations

import shutil
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import build_g3_09_test_plan as base

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "docs/work/A_PM/risk_management_plan.md"
REFERENCE = ROOT / "docs/reference/25-风险管理计划与风险登记册v2（教学样例）.docx"
OUTPUT = ROOT / "docs/deliverables/23-风险管理计划与风险登记册v2.docx"

def remove_reference_body(doc: Document) -> None:
    start = next((p for p in doc.paragraphs if p.text.strip().startswith("1 风险管理方法")), None)
    if start is None:
        raise RuntimeError("Reference body start not found")
    body = doc.element.body
    children = list(body)
    start_idx = children.index(start._p)
    for child in children[start_idx:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)

def replace_toc(doc: Document) -> None:
    body = doc.element.body
    sdts = list(body.xpath("./w:sdt"))
    if not sdts:
        raise RuntimeError("Reference TOC content control not found")
    toc = sdts[0]
    entries = [
        (1,"1 目的、范围与依据",1),(1,"2 风险治理、角色与衔接",2),(1,"3 识别、评价与优先级",2),
        (2,"3.1 识别与分类",2),(2,"3.2 概率与影响",2),(1,"4 风险登记册 v2",3),
        (2,"4.1 当前极高与高风险",3),(2,"4.2 中低风险与持续监测",5),(1,"5 触发、监测、应对与升级",6),
        (1,"6 AC-G2-FR-006-01 专项核对",6),(1,"7 Issue、变更、测试与配置闭环",7),
        (1,"8 报告、复核与关闭",7),(1,"9 主责自检与准出",8)]
    for level, label, page in entries:
        p=OxmlElement("w:p"); ppr=OxmlElement("w:pPr"); style=OxmlElement("w:pStyle"); style.set(qn("w:val"),f"TOC{level}"); ppr.append(style)
        tabs=OxmlElement("w:tabs"); tab=OxmlElement("w:tab"); tab.set(qn("w:val"),"right"); tab.set(qn("w:leader"),"dot"); tab.set(qn("w:pos"),"8500"); tabs.append(tab); ppr.append(tabs); p.append(ppr)
        r=OxmlElement("w:r"); t=OxmlElement("w:t"); t.text=label; r.append(t); p.append(r); tr=OxmlElement("w:r"); tr.append(OxmlElement("w:tab")); p.append(tr); pr=OxmlElement("w:r"); pt=OxmlElement("w:t"); pt.text=str(page); pr.append(pt); p.append(pr); toc.addprevious(p)
    body.remove(toc)

def build() -> None:
    shutil.copy2(REFERENCE, OUTPUT)
    doc=Document(OUTPUT); remove_reference_body(doc); replace_toc(doc)
    cover={0:"文档编号：YJGL-G3-14-RMP　　版本号：V0.1",1:"密　　级：内部 · 教学用",3:"某自然博物馆智能运营中心建设项目——应急管理子系统",4:"风险管理计划与风险登记册 v2",5:"（G3-14 评审候选）",7:"编制单位：020202项目组【待人工确认】",8:"编　　制：何思源（项目负责人）",9:"审　　核：严宇（技术负责人）",10:"批　　准：【待人工确认】",11:"编制日期：2026 年 9 月 20 日",14:"修订记录",15:"注：本表只记录可核验的版本事件；B 独立复核通过前，本文件保持评审候选。"}
    for i,text in cover.items(): base.replace_paragraph(doc.paragraphs[i],text)
    revision=doc.tables[0]
    while len(revision.rows)>1: revision._tbl.remove(revision.rows[-1]._tr)
    vals=["V0.1","2026-09-20","全文","建立 G3-14 风险管理计划与风险登记册 v2，覆盖识别、评价、责任、触发、应对、复核、关闭及 AC-G2-FR-006-01 专项核对。","何思源"]
    row=revision.add_row().cells
    for cell,val in zip(row,vals): cell.text=val
    base.format_table(revision, [["版本","日期","修订章节","修订说明","编制/修订人"],vals])
    base.add_body(doc, WORK.read_text(encoding="utf-8"))
    body=doc.element.body; sect=body.sectPr; body.remove(sect); body.append(sect)
    doc.core_properties.title="风险管理计划与风险登记册 v2"; doc.core_properties.subject="G3-14"; doc.core_properties.author="020202项目组"
    for section in doc.sections:
        for header in (section.header,section.first_page_header,section.even_page_header):
            if header.paragraphs: base.replace_paragraph(header.paragraphs[0],"某自然博物馆智能运营中心建设项目——应急管理子系统 · 风险管理计划",size=9)
    doc.save(OUTPUT)

if __name__ == "__main__":
    build(); print(OUTPUT)
