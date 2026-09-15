"""Mechanical regression checks for the G2-R03 review candidate."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "deliverables" / "08-软件需求规格说明书SRS.docx"
MARKDOWN = ROOT / "docs" / "work" / "A_PM" / "software_requirements_specification_v0.1.md"


def body_blocks(document):
    table_by_element = {table._tbl: table for table in document.tables}
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "paragraph", "".join(child.itertext()).strip()
        elif child.tag == qn("w:tbl"):
            table = table_by_element[child]
            yield "table", table.cell(0, 0).text.strip()


doc = Document(DOCX)
blocks = list(body_blocks(doc))

markers = [
    "3.5 非 MVP",
    "3.6核心业务流程",
    "表 3-6 核心业务流程",
    "3.7核心状态模型",
    "表 3-7 核心业务状态模型",
    "4 外部接口需求",
    "5 数据需求",
]
positions = {}
for marker in markers:
    positions[marker] = next(i for i, (_, value) in enumerate(blocks) if value.startswith(marker))

assert list(positions.values()) == sorted(positions.values()), positions

text = "\n".join(p.text for p in doc.paragraphs)
text += "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
markdown = MARKDOWN.read_text(encoding="utf-8")

doc_frs = set(re.findall(r"G2-FR-(\d{3})", text))
md_frs = set(re.findall(r"G2-FR-(\d{3})", markdown))
pes = set(re.findall(r"PE-(\d{2})", text))
rclrs = set(re.findall(r"G2-RCLR-(\d{3})", text))

assert doc_frs == {f"{i:03}" for i in range(1, 40)}, len(doc_frs)
assert md_frs == doc_frs, (len(md_frs), len(doc_frs))
assert pes == {f"{i:02}" for i in range(1, 13)}, sorted(pes)
assert rclrs == {f"{i:03}" for i in range(1, 11)}, sorted(rclrs)
assert all(term not in text for term in ("M-01 项目与会话管理", "影像数据接入", "模型注册", "分割执行"))

section = doc.sections[0]
print("ORDER", positions)
print("COUNTS", {"G2-FR": len(doc_frs), "PE": len(pes), "G2-RCLR": len(rclrs), "tables": len(doc.tables)})
print(
    "PAGE_CM",
    {
        "width": round(section.page_width.cm, 2),
        "height": round(section.page_height.cm, 2),
        "top": round(section.top_margin.cm, 2),
        "bottom": round(section.bottom_margin.cm, 2),
        "left": round(section.left_margin.cm, 2),
        "right": round(section.right_margin.cm, 2),
    },
)
print("PASS")
