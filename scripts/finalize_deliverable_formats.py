from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DEL = ROOT / "docs" / "deliverables"
REF = ROOT / "docs" / "reference"
FIG = DEL / "figures"

PAIRS = {
    "00": ("00-投标文件技术标.docx", "00-澜图遥感影像智能解译平台-投标文件技术标（教学案例）.docx"),
    "01": ("01-项目建议书.docx", "01-项目建议书（教学样例）.docx"),
    "02": ("02-澄清与质询记录.docx", "02-澄清与质询记录（教学样例）.docx"),
    "03": ("03-项目计划v1（WBS与甘特图）.docx", "03-项目计划v1（WBS与甘特图·教学样例）.docx"),
    "04": ("04-风险登记册v1.docx", "04-风险登记册v1（教学样例）.docx"),
    "05": ("05-需求确认书.docx", "05-需求确认书（教学样例）.docx"),
    "06": ("06-述标答辩讲稿与策略.docx", "06-述标答辩讲稿与策略（教学样例）.docx"),
}

HISTORIES = {
    "01": [
        ("V1.0", "2026-09-12", "全文", "完成技术与符合性复核并纳入第一关冻结成果。", "何思源"),
        ("V1.1", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.2", "2026-09-13", "版式、图 2-1、图 3-1", "按教学样例复核版式并补充业务流程与建设内容分层图。", "何思源"),
    ],
    "02": [
        ("V1.1", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.2", "2026-09-13", "版式、图 1-1、修订记录", "按教学样例复核版式并以可编辑工程图替换流程占位表。", "何思源"),
    ],
    "03": [
        ("V1.3", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.4", "2026-09-13", "版式与修订记录", "按教学样例逐项复核页面、样式、表格及既有 WBS/甘特图。", "何思源"),
    ],
    "04": [
        ("V1.1", "2026-09-11", "风险关联与跟踪机制", "按复核意见补齐计划联动、触发条件和关闭证据口径。", "何思源"),
        ("V1.2", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.3", "2026-09-13", "版式与修订记录", "按教学样例逐项复核页面、表格与风险矩阵呈现。", "何思源"),
    ],
    "05": [
        ("V0.2", "2026-09-12", "第2—5章", "按需求复核修正 MVP 范围、边界、指标与变更控制口径。", "何思源"),
        ("V1.0", "2026-09-12", "全文", "通过技术与符合性复核并纳入第一关冻结成果。", "何思源"),
        ("V1.1", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.2", "2026-09-13", "前置页、图 2-1、修订记录", "按教学样例复核版式，分离修订记录与正文并补充可编辑 MVP 边界图。", "何思源"),
    ],
}

FINAL_VERSION = {"01": "V1.2", "02": "V1.2", "03": "V1.4", "04": "V1.3", "05": "V1.2"}


def replace_text(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            for run in p.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
            if old in p.text:
                text = p.text.replace(old, new)
                p.clear().add_run(text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if old in p.text:
                        for run in p.runs:
                            if old in run.text:
                                run.text = run.text.replace(old, new)


def clone_named_styles(doc, ref):
    names = ["Title", "Heading 1", "Heading 2", "Heading 3", "Normal", "Caption", "List Paragraph"]
    for name in names:
        try:
            dst, src = doc.styles[name], ref.styles[name]
        except KeyError:
            continue
        for tag in ("w:rPr", "w:pPr"):
            old = dst._element.find(qn(tag))
            new = src._element.find(qn(tag))
            if old is not None:
                dst._element.remove(old)
            if new is not None:
                dst._element.append(deepcopy(new))


def align_sections(doc, ref):
    source = ref.sections[0]
    for sec in doc.sections:
        if sec.orientation == source.orientation:
            sec.page_width, sec.page_height = source.page_width, source.page_height
        sec.top_margin, sec.bottom_margin = source.top_margin, source.bottom_margin
        sec.left_margin, sec.right_margin = source.left_margin, source.right_margin
        sec.header_distance, sec.footer_distance = source.header_distance, source.footer_distance


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=80, bottom=80, end=80):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val)); node.set(qn("w:type"), "dxa")


def format_revision_table(table):
    try:
        table.style = "Table Grid"
    except KeyError:
        tbl_pr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            node = OxmlElement(f"w:{edge}")
            node.set(qn("w:val"), "single"); node.set(qn("w:sz"), "6")
            node.set(qn("w:space"), "0"); node.set(qn("w:color"), "7F8C8D")
            borders.append(node)
        tbl_pr.append(borders)
    table.autofit = False
    widths = [Cm(1.7), Cm(2.2), Cm(2.5), Cm(8.0), Cm(2.7)]
    for r_i, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row.height = Cm(0.78)
        if r_i == 0:
            tr_pr = row._tr.get_or_add_trPr()
            rep = OxmlElement("w:tblHeader"); rep.set(qn("w:val"), "true"); tr_pr.append(rep)
        for c_i, cell in enumerate(row.cells):
            if c_i < len(widths): cell.width = widths[c_i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_i == 0: shade_cell(cell, "D9E5F3")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i != 3 else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    run.font.name = "宋体"; run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(9); run.bold = (r_i == 0)


def add_history(doc, code):
    table = doc.tables[0]
    existing = {r.cells[0].text.strip() for r in table.rows[1:]}
    for version, date, section, note, author in HISTORIES[code]:
        if version in existing: continue
        cells = table.add_row().cells
        for cell, value in zip(cells, (version, date, section, note, author)):
            cell.text = value
    format_revision_table(table)


def para_after(doc, anchor, text="", style=None):
    p = doc.add_paragraph(text, style=style)
    anchor._p.addnext(p._p)
    return p


def image_after(doc, anchor, image_name, caption):
    if any(caption in p.text for p in doc.paragraphs):
        return
    cap = para_after(doc, anchor, caption, "Caption" if "Caption" in [s.name for s in doc.styles] else None)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(FIG / image_name), width=Cm(16.0))
    cap._p.addprevious(pic._p)


def image_before_existing_caption(doc, caption_fragment, image_name):
    for p in doc.paragraphs:
        if caption_fragment in p.text:
            previous = p._p.getprevious()
            if previous is not None and previous.xpath('.//a:blip'):
                return
            pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.keep_with_next = True
            pic.add_run().add_picture(str(FIG / image_name), width=Cm(16.0))
            p._p.addprevious(pic._p)
            try: p.style = doc.styles["Caption"]
            except KeyError: pass
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            return


def find_para(doc, fragment):
    for p in doc.paragraphs:
        if fragment in p.text:
            return p
    raise ValueError(f"paragraph not found: {fragment}")


def add_technical_revision_page(doc):
    if any("技术标修订记录" in p.text for p in doc.paragraphs):
        return
    toc = find_para(doc, "目 录")
    heading = doc.add_paragraph("技术标修订记录", style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table = doc.add_table(rows=1, cols=5)
    for cell, text in zip(table.rows[0].cells, ("版本", "日期", "修订章节", "修订说明", "编制/修订人")):
        cell.text = text
    rows = [
        ("V0.1", "2026-09-11", "全文", "形成技术投标书初稿并开始按需求与技术成果整合。", "何思源"),
        ("V1.0", "2026-09-12", "全文", "完成技术、符合性及跨文档复核，形成第一关冻结技术标。", "何思源"),
        ("V1.1", "2026-09-13", "封面及署名", "完成正式交付物实名化和首轮格式规范化。", "何思源"),
        ("V1.2", "2026-09-13", "前置页、第7—9章、附录", "按教学案例逐项复核版式，补充安全、责任与测试图并修复附录代号。", "何思源"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values): cell.text = value
    format_revision_table(table)
    note = doc.add_paragraph("注：版本事件依据历史 Prompt、Review 与 Git 提交恢复；仅记录可核验的重要阶段。")
    note.add_run().add_break(WD_BREAK.PAGE)
    toc._p.addprevious(heading._p); toc._p.addprevious(table._tbl); toc._p.addprevious(note._p)


def set_body_page_break_after_revision(doc, next_heading_fragment):
    p = find_para(doc, next_heading_fragment)
    p.paragraph_format.page_break_before = True


def remove_table(doc, index):
    table = doc.tables[index]
    table._element.getparent().remove(table._element)


for code, (current_name, ref_name) in PAIRS.items():
    path, ref_path = DEL / current_name, REF / ref_name
    doc, ref = Document(path), Document(ref_path)
    clone_named_styles(doc, ref)
    align_sections(doc, ref)

    if code == "00":
        add_technical_revision_page(doc)
        replace_text(doc, "附录何思源至G", "附录 A 至 G")
        replace_text(doc, "附录何思源 术语和缩略语", "附录 A 术语和缩略语")
        replace_text(doc, "附录严宇 接口与报文设计节选", "附录 B 接口与报文设计节选")
        replace_text(doc, "附录任俊强 条款—投标方案索引", "附录 C 条款—投标方案索引")
        replace_text(doc, "附录何思源至G", "附录 A 至 G")
        toc = find_para(doc, "第一章  投标响应总述")
        toc.text = "\n".join([
            "第一章  投标响应总述\t5", "第二章  需求理解\t6", "第三章  总体技术方案\t10",
            "第四章  功能实现方案\t17", "第五章  关键技术实现方案\t26", "第六章  开源组件与许可证合规\t32",
            "第七章  非功能设计\t33", "第八章  项目实施方案\t35", "第九章  质量保障方案\t39",
            "第十章  培训与售后服务方案\t41", "第十一章  技术规格响应与偏离表\t41", "第十二章  项目团队\t57",
            "第十三章  类似项目业绩\t58", "第十四章  合理化建议\t59", "附录 A  术语和缩略语\t60",
            "附录 B  接口与报文设计节选\t61", "附录 C  条款—投标方案索引\t62", "附录 D  核心数据字典\t64",
            "附录 E  质量度量与验收用例框架\t65", "附录 F  评分要点响应对照\t66", "附录 G  交付物清单\t67",
        ])
        image_after(doc, find_para(doc, "7.2 信息安全性设计"), "00-图7-2-应用安全防护体系.png", "图 7-2 应用安全防护体系")
        image_after(doc, find_para(doc, "8.1 项目组织"), "00-图8-5-项目责任关系.png", "图 8-5 项目责任关系")
        image_after(doc, find_para(doc, "9.2 测试策略"), "00-图9-2-四级测试与验收证据链.png", "图 9-2 四级测试与验收证据链")
    elif code != "06":
        add_history(doc, code)
        version = FINAL_VERSION[code]
        for p in doc.paragraphs[:20]:
            if "版本号：" in p.text:
                prefix = p.text.split("版本号：", 1)[0]
                p.text = prefix + "版本号：" + version
            if "文档版本：" in p.text:
                p.text = "文档版本：" + version
        if code in {"03", "04"}:
            replace_text(doc, "成员 何思源", "何思源")
            replace_text(doc, "成员 任俊强", "任俊强")
            replace_text(doc, "成员何思源", "何思源")
        if code == "01":
            image_after(doc, find_para(doc, "2.1 项目定位与建设目标"), "01-图2-1-应急业务全流程与建设目标.png", "图 2-1 应急业务全流程与建设目标")
            image_after(doc, find_para(doc, "3.1 总体思路"), "01-图3-1-建设内容分层结构.png", "图 3-1 建设内容分层结构")
        elif code == "02":
            image_before_existing_caption(doc, "图 1-1", "02-图1-1-澄清与质询闭环流程.png")
            # The old one-row table was only a diagram surrogate; the native-PPT figure replaces it.
            if len(doc.tables) == 4:
                remove_table(doc, 1)
        elif code == "05":
            image_before_existing_caption(doc, "图 2-1", "05-图2-1-MVP范围与边界一览.png")
            if len(doc.tables) == 5:
                remove_table(doc, 1)
            set_body_page_break_after_revision(doc, "1 目的与依据")

    doc.save(path)
    print(f"UPDATED={path.name}")
