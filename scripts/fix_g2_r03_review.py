from copy import deepcopy
import os
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / 'docs' / 'deliverables' / '08-软件需求规格说明书SRS.docx'
OUTPUT = Path(os.environ.get('G2_R03_SRS_OUTPUT', PATH))


def find(doc, prefix):
    return next(p for p in doc.paragraphs if p.text.startswith(prefix))


def text(p, value):
    p.clear()
    p.add_run(value)


def set_cell(cell, value, header=False):
    cell.text = value
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.name = '宋体'
            r.font.size = Pt(7.5)
            r.bold = header
            r._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:tblHeader')
    node.set(qn('w:val'), 'true')
    tr_pr.append(node)


doc = Document(PATH)

text(doc.paragraphs[0], 'YJGL-G2-03    V1.2')
text(doc.paragraphs[11], '编制日期：2026 年 9 月 16 日')

# Record this real Review-driven correction in the controlled revision table.
history = doc.tables[0]
if not any(row.cells[0].text.strip() == 'V1.2' for row in history.rows[1:]):
    cells = history.add_row().cells
    for cell, value in zip(cells, ('V1.2', '2026-09-16', '第3、8章；目录', '按 C Review 修正流程/状态插入顺序、39/39 AC 完整性表述及重复目录续段，并重建渲染证据。', '何思源')):
        set_cell(cell, value)

# Earlier TOC repair passes could leave a duplicate continuation paragraph.
# Retain the first controlled continuation only.
toc_continuations = [
    p for p in doc.paragraphs
    if p.text.startswith('3.6 核心业务流程与异常路径\t10')
]
for duplicate in toc_continuations[1:]:
    duplicate._p.getparent().remove(duplicate._p)

# Keep the implementation-level additions in the same body order as the
# controlled Markdown and the static TOC: 3.5 → 3.6 → 3.7 → 4.
chapter4 = find(doc, '4 外部接口需求')
flow_heading = find(doc, '3.6核心业务流程')
flow_caption = find(doc, '表 3-6')
flow_table = flow_caption._p.getnext()
state_heading = find(doc, '3.7核心状态模型')
state_caption = find(doc, '表 3-7')
state_table = state_caption._p.getnext()
for element in [flow_heading._p, flow_caption._p, flow_table, state_heading._p, state_caption._p, state_table]:
    chapter4._p.addprevious(element)

# Clear stale teaching-case TOC field results and use a controlled static TOC.
toc = doc.paragraphs[15:18]
text(toc[0], '目 录')
toc[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
text(toc[1], '1 引言\t1\n2 总体说明\t1\n3 功能需求\t2\n3.1 数字预案与事件处置闭环\t2\n3.2 指挥态势、视频与定位联动\t3\n3.3 演练、移动执行与值班物资闭环\t5\n3.4 外部联动与中台业务适配\t8\n3.5 非MVP、本期范围功能需求\t9')
text(toc[2], '3.6 核心业务流程与异常路径\t10\n3.7 核心状态模型\t10\n4 外部接口需求\t10\n5 数据需求\t11\n5.1 核心数据对象\t11\n5.2 核心实体字段字典\t11\n5.3 数据质量与来源\t12\n5.4 数据安全与审计\t12\n6 非功能需求\t12\n7 部署、交付与服务约束\t13\n8 验收与追踪\t14\n9 变更控制与待澄清事项\t14\n10 G2规格化阶段补充业务规则\t15')

# Synchronize the task-stage wording with the reviewed working SRS.
text(find(doc, '下表保留39条'), '下表保留39条受控功能需求及既有验收标准。G2-R02已完成反向澄清裁决，G2-R03仅补充需求级流程、状态、字段和量化语义；本轮影响AC的可观察规则均标记为G2-R05对spec.md与RTM的待同步项。')
text(find(doc, 'G2-R05应在'), 'G2-R05应在spec.md中使用本文件相同的G2-FR和AC编号，并将本轮标记为“spec/RTM待同步”的可观察规则补入双镜像。RTM至少追踪原始FR、G2-FR、用户故事、AC、G2-RCLR、设计、实现、测试和结果；任何编号、范围或验收语义差异均阻断M2返工准出。')
text(next(p for p in doc.paragraphs if p.text.startswith(('29 条 G2-FR', '39 条 G2-FR'))), '39 条 G2-FR 均有至少一个 AC-G2-FR-nnn-xx 验收标准。')
text(find(doc, '功能、★属性、数字'), '功能、★属性、数字、责任、期次或验收强度发生变化时，必须先登记Issue、分析影响并取得授权确认，再同步更新SRS、spec.md、RTM和版本记录。G2-R02已形成并记录G2-RCLR-001—010的课程项目人工裁决；其后新发现、且未由冻结事实或现有裁决确定的问题不得静默写入需求，应按Issue/Change流程处理。')

# Correct the chapter order and rebuild Table 5-1 with the same five requirements-level dimensions as Markdown.
data_quality = find(doc, '5.3 数据质量')
field_heading = find(doc, '5.2 核心实体字段')
field_heading._p.getparent().remove(field_heading._p)
data_quality._p.addprevious(field_heading._p)
text(field_heading, '5.2 核心实体字段字典。表 5-1 为需求级语义字典，不是数据库 ER 或物理表设计。')
old = doc.tables[9]
style = old.style
old._element.getparent().remove(old._element)
rows = [
('应急预案', 'planId String/是；名称 String/是；状态 Enum/是；适用场景 String/是', '唯一标识、名称、草稿/已发布/停用和适用条件', 'G2-FR-001—003', '版本、模板、事件'),
('预案版本', 'versionId String/是；版本号 String/是；生效状态 Enum/是', '版本可追溯，发布后作为启动依据', 'G2-FR-001', '预案、模板'),
('预案任务模板', 'templateId String/是；任务名称 String/是；责任组 String/是；时限 Duration/是', '启动时生成任务；缺关键关系不得发布', 'G2-FR-002、003', '版本、任务'),
('应急事件', 'incidentId String/是；类型 Enum/是；状态 Enum/是；发生时间 DateTime/是', '唯一事件；含待核实/核实超时/处置/关闭状态', 'G2-FR-013—016', '续报、任务、附件'),
('事件续报', 'updateId String/是；业务时间 DateTime/是；接收时间 DateTime/是；内容 Text/是', '迟到数据标记且不覆盖历史', 'G2-RCLR-001、004', '事件'),
('处置任务', 'taskId String/是；当前责任人 String/是；状态 Enum/是；截止时间 DateTime/是', '重指派保留历史，改期显式审计', 'G2-FR-015、G2-RCLR-003', '事件、反馈'),
('任务反馈', 'feedbackId String/是；内容 Text/是；反馈时间 DateTime/是', '接收、反馈、催办和完成证据', 'G2-FR-015', '任务'),
('人员/值班', 'personId String/是；职责 Enum/是；值班时段 TimeRange/否', '人员、职责与当前值班关联', 'G2-FR-017、018', '任务、位置'),
('位置状态', 'personId String/是；坐标 Location/是；源时间 DateTime/是；新鲜度 Enum/是', '过期阈值可配置，过期禁自动调派', 'G2-FR-005、G2-RCLR-009', '人员、事件'),
('物资站点/条目', 'siteId String/是；itemId String/是；数量 Decimal/是；有效期 Date/否', '台账、站点与临期状态', 'G2-FR-007—009', '盘点'),
('盘点计划/快照/记录', 'planId String/是；快照数量 Decimal/是；实盘数量 Decimal/否；复核状态 Enum/是', '快照为比对基准，期间变动单独记录', 'G2-FR-009、G2-RCLR-006', '物资条目'),
('演练计划/执行/评估', 'drillId String/是；计划时间 DateTime/是；实际完成 DateTime/否；状态 Enum/是', '取消/逾期不计完成，补演关联原计划', 'G2-FR-010—012、G2-RCLR-007', '任务、评估'),
('打卡点/记录', 'pointId String/是；任务Id String/是；扫码时间 DateTime/是；结果 Enum/是', '码绑定任务/点位/时段，重复扫码幂等', 'G2-FR-017—020、G2-RCLR-005', '人员'),
('消息发送/外部调用', 'messageId String/是；尝试序号 Integer/是；最终状态 Enum/是；回执时间 DateTime/否', '保留超时/重试/降级历史，重复回执幂等', 'G2-FR-020、022、G2-RCLR-008', '任务、事件'),
('附件/审计记录', 'attachmentId String/是；关联对象 String/是；删除状态 Enum/是；操作者 String/是', '逻辑删除、归档引用保护和删除审计', 'G2-FR-013、016、G2-RCLR-010', '事件、任务、报告'),
]
caption = doc.add_paragraph('表 5-1 核心实体字段字典')
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.keep_with_next = True
table = doc.add_table(rows=1, cols=5)
table.style = style
table.autofit = False
field_heading._p.addnext(caption._p)
caption._p.addnext(table._tbl)
headers = ['实体', '字段（业务类型/必填）', '业务含义与约束', '来源', '关联']
for cell, value in zip(table.rows[0].cells, headers):
    set_cell(cell, value, True)
repeat_header(table.rows[0])
for values in rows:
    for cell, value in zip(table.add_row().cells, values):
        set_cell(cell, value)
for row in table.rows:
    for cell, width in zip(row.cells, [Cm(2.25), Cm(5.45), Cm(4.65), Cm(2.4), Cm(2.35)]):
        cell.width = width

doc.core_properties.last_modified_by = '何思源'
doc.save(OUTPUT)
print(OUTPUT)
