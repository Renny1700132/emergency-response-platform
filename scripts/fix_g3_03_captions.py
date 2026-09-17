from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
p='docs/deliverables/11-概要设计说明书.docx'
d=Document(p)
titles={1:'表 2-1 逻辑技术架构分层',2:'表 2-2 模块职责与依赖',3:'表 4-1 部署故障域与恢复约束',4:'表 5-1 主要非功能设计预算',5:'表 6-1 需求与设计元素追踪'}
# remove prior duplicate table 2-1 plain paragraph
for para in list(d.paragraphs):
 if para.text=='表 2-1 逻辑技术架构分层': para._element.getparent().remove(para._element)
for i,title in titles.items():
 t=d.tables[i]; elem=OxmlElement('w:p'); t._tbl.addnext(elem); para=Paragraph(elem,t._parent); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.add_run(title)
d.save(p)
