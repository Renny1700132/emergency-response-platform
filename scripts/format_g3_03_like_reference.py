from docx import Document
from copy import deepcopy
from docx.enum.text import WD_ALIGN_PARAGRAPH
src=Document('docs/reference/11-概要设计说明书（教学样例）.docx')
dst=Document('docs/deliverables/11-概要设计说明书.docx')
# exact section and header/footer properties
for a,b in zip(dst.sections,src.sections):
 for old in list(a._sectPr): a._sectPr.remove(old)
 for x in src.sections[0]._sectPr: a._sectPr.append(deepcopy(x))
# source role paragraphs
cover=[src.paragraphs[i] for i in [0,1,3,4,5,7,8,9,10,11]]
h1=src.paragraphs[16]; h2=src.paragraphs[17]; normal=src.paragraphs[18]; caption=src.paragraphs[35]
def cp(dstp,sp):
 pp=dstp._p.get_or_add_pPr(); [pp.remove(z) for z in list(pp)]; pp.extend(deepcopy(sp._p.pPr))
 for r in dstp.runs:
  if sp.runs: rr=r._r.get_or_add_rPr(); [rr.remove(z) for z in list(rr)]; rr.extend(deepcopy(sp.runs[0]._r.rPr))
# cover exact roles
for p,s in zip(dst.paragraphs[:7],[cover[2],cover[3],cover[0],cover[1],cover[5],cover[6],cover[9]]): cp(p,s)
for p in dst.paragraphs:
 if p.style.name=='Heading 1': cp(p,h1)
 elif p.style.name=='Heading 2': cp(p,h2)
 elif p.text.startswith('表 '): cp(p,caption)
 elif p.text and p not in dst.paragraphs[:7] and p.style.name=='Normal': cp(p,normal)
# direct table formats mapped to like reference tables
for i,t in enumerate(dst.tables):
 ref=src.tables[min(i,len(src.tables)-1)]
 tp=t._tbl.tblPr; [tp.remove(z) for z in list(tp)]; tp.extend(deepcopy(ref._tbl.tblPr))
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs: cp(p,normal)
dst.save('docs/deliverables/11-概要设计说明书.docx')


