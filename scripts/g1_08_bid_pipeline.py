from __future__ import annotations

import argparse
import re
import textwrap
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT / "docs/reference/00-澜图遥感影像智能解译平台-投标文件技术标（教学案例）.docx"
REQUIREMENTS = ROOT / "docs/用户需求书-03-应急管理子系统.docx"
EXISTING_BID = ROOT / "docs/deliverables/00-投标文件技术标.docx"
AUDIT = ROOT / "docs/work/A_PM/G1-08_granularity_alignment_audit.md"
BID_MD = ROOT / "docs/work/A_PM/technical_bid.md"
FIGURES_PPTX = ROOT / "docs/work/A_PM/technical_bid_figures.pptx"


def dump_docx(path: Path, output: Path) -> None:
    """Dump every paragraph and table cell in document order for source review."""
    doc = Document(path)
    lines = [
        f"# SOURCE: {path.relative_to(ROOT)}",
        f"# paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} "
        f"sections={len(doc.sections)} inline_shapes={len(doc.inline_shapes)}",
        "",
    ]
    table_index = 0
    for block in doc.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if text:
                lines.append(f"[P|{block.style.name}] {text}")
        elif isinstance(block, Table):
            table_index += 1
            lines.append(f"[TABLE {table_index}|rows={len(block.rows)}|cols={len(block.columns)}]")
            for row in block.rows:
                values = [" ".join(cell.text.split()) for cell in row.cells]
                lines.append(" | ".join(values))
            lines.append(f"[/TABLE {table_index}]")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines), encoding="utf-8")


def analyze() -> None:
    out_dir = ROOT / ".tmp/g1_08"
    dump_docx(REFERENCE, out_dir / "reference_full.txt")
    dump_docx(REQUIREMENTS, out_dir / "requirements_full.txt")
    dump_docx(EXISTING_BID, out_dir / "existing_bid_full.txt")
    print(out_dir)


def _number(text: str) -> str:
    match = re.match(r"^((?:\d+\.)*\d+|[A-G]\.\d+)", text)
    return match.group(1) if match else ""


def map_reference_heading(text: str) -> tuple[str, str, str]:
    """Return project mapping, observed gap and planned treatment."""
    chapter_map = {
        "使用说明（教学案例）": "编制说明与证据边界",
        "第一章 投标响应总述": "第一章 投标响应总述",
        "第二章 需求理解": "第二章 需求理解",
        "第三章 总体技术方案": "第三章 总体技术方案",
        "第四章 功能实现方案": "第四章 功能实现方案",
        "第五章 关键技术实现方案": "第五章 关键技术实现方案",
        "第六章 开源组件清单及许可证合规说明": "第六章 开源组件与许可证合规",
        "第七章 非功能设计": "第七章 非功能设计",
        "第八章 项目实施方案": "第八章 项目实施方案",
        "第九章 质量保障方案": "第九章 质量保障方案",
        "第十章 培训与售后服务方案": "第十章 培训与售后服务方案",
        "第十一章 技术规格偏离表": "第十一章 技术规格响应与偏离表",
        "第十二章 项目团队": "第十二章 项目团队",
        "第十三章 类似项目业绩": "第十三章 类似项目业绩",
        "第十四章 合理化建议": "第十四章 合理化建议",
        "附录A 术语和缩略语": "附录A 术语和缩略语",
        "附录B 接口报文示例": "附录B 接口与报文设计节选",
        "附录C 招标条款—投标方案索引表": "附录C 条款—投标方案索引",
        "附录D 数据字典（核心表节选）": "附录D 核心数据字典",
        "附录E 质量度量目标与验收用例框架": "附录E 质量度量与验收用例框架",
        "附录F 评分要点响应对照表": "附录F 评分要点响应对照",
        "附录G 交付文档清单与国标对照": "附录G 交付物清单",
    }
    if text in chapter_map:
        target = chapter_map[text]
        gap = "结构存在但粒度不足" if text != "使用说明（教学案例）" else "缺少独立说明"
        return target, gap, "保留同级结构，按本项目事实充分展开"
    num = _number(text)
    if text.startswith("2.5."):
        names = ["秒级启动", "可靠消息", "空间态势与定位", "安全联动", "中台复用与合规"]
        idx = min(int(num.rsplit(".", 1)[1]), len(names))
        return f"2.5.{idx} 难点：{names[idx - 1]}", "案例业务不适用", "以本项目等价难点替换，并在第五章深化"
    if text.startswith("2."):
        names = {
            "2.1": "项目背景与现状问题", "2.2": "建设目标与可验证结果",
            "2.3": "用户角色与业务场景", "2.4": "十一个功能域总体分析",
            "2.5": "重点难点分析", "2.6": "技术路线与边界对比",
            "2.7": "典型应急场景走查", "2.8": "需求优先级与里程碑映射",
            "2.9": "十二项模拟书面澄清及效力边界",
            "2.10": "同类系统常见失效模式与项目对策",
            "2.11": "建设价值与度量方法",
        }
        return f"{num} {names.get(num, text)}", "缺失或被压缩", "按本项目需求建立等价小节"
    if text.startswith("3."):
        return text.replace("原型验证计划", "验证计划"), "未按参考16节展开", "以G1-04 V0.6为主源逐节整合"
    if re.match(r"4\.\d+ F-", text):
        idx = int(re.match(r"4\.(\d+)", text).group(1))
        names = {
            1: "应急预案管理", 2: "应急指挥中心", 3: "应急资源管理",
            4: "应急演练管理", 5: "应急事件管理", 6: "应急值班管理",
            7: "应急基础管理", 8: "应急管理移动端", 9: "应急信息视图",
            10: "综合安防视图",
        }
        return f"4.{idx} F-{idx:02d} {names[idx]}", "现有功能域被合并", "按本项目功能域替换；另增4.11 F-11数据接口"
    if re.match(r"4\.\d+\.\d+", text):
        return num + " 工程细化", "缺少三级论证", "每个F域扩展需求、组成、流程、数据、接口、异常、权限、验证"
    if text.startswith("4."):
        extras = {
            "4.11": "4.12 界面清单与导航结构",
            "4.12": "4.13 异常处理与降级设计",
            "4.13": "4.14 浏览器、H5兼容与界面规范",
            "4.14": "4.15 数据初始化与预置内容",
            "4.15": "4.16 角色—功能—数据权限矩阵",
        }
        return extras.get(num, text), "缺失或内容不足", "保留等价横向设计并增加F-11"
    if text.startswith("5."):
        techs = [
            "秒级预案启动", "可靠消息推送", "任务编排与事件状态机",
            "事件态势一张图", "定位链路", "视频GB/T 28181接入与回放",
            "门禁安全联动", "统一中台适配", "H5宿主集成",
            "配置化流程", "接口容错与降级", "性能容量", "备份恢复与连续运行",
        ]
        idx = min(int(num.split(".")[1]), len(techs))
        return f"5.{idx} {techs[idx - 1]}", "案例算法主题不适用", "用本项目关键技术替换并扩展至13项"
    if text.startswith(("6.", "7.", "8.", "9.", "10.", "11.")):
        treatment = {
            "6": "补齐准入、SBOM、SCA、许可证和知识产权流程",
            "7": "逐项写指标、机制、异常和验证证据",
            "8": "把G1-06组织、WBS、里程碑、配置、沟通和试运行整合入正文",
            "9": "补齐四级测试、缺陷、评审、环境和验收组织",
            "10": "按培训、SLA、应急和知识转移条款展开",
            "11": "由168行合规矩阵逐条生成要求、响应、位置、证据和偏离",
        }[num.split(".")[0]]
        return text, "现稿概要化或缺失", treatment
    if text.startswith("12."):
        return text.replace("（教学示例）", "【待人工确认】"), "缺少真实人员依据", "保留岗位与简历模板，人员事实全部待人工确认"
    if text.startswith("13."):
        return text.replace("（教学示例）", "【待人工确认】"), "案例业绩不可采用", "保留完整填报结构并标待人工确认"
    if text.startswith("B."):
        names = {
            "B.1": "申请访问令牌", "B.2": "提交事件与启动预案",
            "B.3": "查询任务状态", "B.4": "消息回执、定位与视频请求",
            "B.5": "错误码与重试语义",
        }
        return f"{num} {names.get(num, text)}", "案例报文不适用", "用本项目接口报文替换"
    if text.startswith("D."):
        names = {
            "D.1": "应急事件", "D.2": "预案快照与任务",
            "D.3": "消息与接口调用", "D.4": "资源、打卡与审计",
        }
        return f"{num} {names.get(num, text)}", "案例数据对象不适用", "用本项目核心实体替换并扩展字段"
    if text.startswith(("C.", "E.", "F.")):
        return text, "索引或内容不足", "按本项目条款、指标和评分证据补齐"
    return text, "待核对", "保留结构；无事实依据处标待人工确认"


def reference_heading_metrics(doc: Document) -> list[tuple[str, str, int, int]]:
    metrics = []
    current = None
    tables = images = 0
    for item in doc.iter_inner_content():
        if isinstance(item, Paragraph):
            is_heading = item.text.strip() and ("Heading" in item.style.name or "标题" in item.style.name)
            if is_heading:
                if current:
                    metrics.append((current[0], current[1], tables, images))
                current = (item.text.strip(), item.style.name)
                tables = 0
                images = len(item._p.xpath(".//a:blip"))
            elif current:
                images += len(item._p.xpath(".//a:blip"))
        elif isinstance(item, Table) and current:
            tables += 1
            images += len(item._tbl.xpath(".//a:blip"))
    if current:
        metrics.append((current[0], current[1], tables, images))
    return metrics


def build_audit() -> None:
    metrics = reference_heading_metrics(Document(REFERENCE))
    lines = [
        "# G1-08 技术标粒度对齐审计", "",
        "- Task：G1-08",
        "- 执行角色：A（项目经理 PM / 技术标总编）",
        "- 参考文件：docs/reference/00-澜图遥感影像智能解译平台-投标文件技术标（教学案例）.docx",
        "- 参考规模：14章、附录A—G、59张表、17幅图、717个段落、2个Section。",
        "- 当前稿规模：14章、附录A—G、23张表、4幅图、158个段落；功能域、关键技术、第六至十章和第十一章存在异常合并或证据不足。",
        "- 使用边界：教学案例仅用于结构、信息类型、图表密度和Word格式；案例项目事实、技术栈、人员、业绩、指标均不进入本项目。", "",
        "| 参考章节/小节 | 参考内容类型 | 表格 | 图 | 本项目对应内容 | 当前缺口 | 处理方式 |",
        "| --- | --- | ---: | ---: | --- | --- | --- |",
    ]
    kinds = {"Heading 1": "章/附录", "Heading 2": "主题方案", "Heading 3": "工程细化"}
    for title, style, tables, images in metrics:
        target, gap, treatment = map_reference_heading(title)
        values = [title, kinds.get(style, style), str(tables), str(images), target, gap, treatment]
        lines.append("| " + " | ".join(v.replace("|", "／").replace("\n", " ") for v in values) + " |")
    lines += [
        "", "## 审计结论", "",
        "参考文件全部标题均已逐项映射。重构稿采用同级14章与附录A—G骨架；第四章增加F-11正式功能域，第五章以13项本项目关键技术替换案例算法主题。第十二、十三章保留完整填报结构，缺少真实依据的人员与业绩统一标记【待人工确认】。现有WBS、甘特、里程碑和风险矩阵继续复用，架构与业务机制图按需要补充后进行视觉复核。",
    ]
    AUDIT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {AUDIT.relative_to(ROOT)} with {len(metrics)} mapped headings")


def add_block(lines: list[str], block: str) -> None:
    lines.extend(textwrap.dedent(block).strip().splitlines())
    lines.append("")


def md_table(headers: list[str], rows: list[list[str]]) -> str:
    def clean(value: object) -> str:
        return str(value).replace("|", "／").replace("\n", "<br>")
    output = ["| " + " | ".join(map(clean, headers)) + " |"]
    output.append("| " + " | ".join("---" for _ in headers) + " |")
    output.extend("| " + " | ".join(clean(v) for v in row) + " |" for row in rows)
    return "\n".join(output)


def source_section(source: str, heading: str) -> str:
    pattern = re.compile(rf"^## {re.escape(heading)}.*$", re.M)
    match = pattern.search(source)
    if not match:
        return ""
    start = match.end()
    next_heading = re.search(r"^## (?!#)", source[start:], re.M)
    end = start + next_heading.start() if next_heading else len(source)
    body = source[start:end].strip()
    fence = re.escape(chr(96) * 3)
    body = re.sub(rf"^{fence}mermaid.*?^{fence}\s*", "", body, flags=re.M | re.S)
    return body


def parse_fr_rows() -> dict[int, list[dict[str, str]]]:
    rows: dict[int, list[dict[str, str]]] = {}
    facts = (ROOT / "control/facts.md").read_text(encoding="utf-8")
    pattern = re.compile(r"^\| F-\d+ \| (FR-(\d{2})\.\d+) \| (★|否) \| (.*?) \| (.*?) \|$", re.M)
    for match in pattern.finditer(facts):
        domain = int(match.group(2))
        rows.setdefault(domain, []).append({
            "id": match.group(1), "star": match.group(3),
            "summary": match.group(4), "status": match.group(5),
        })
    return rows


def parse_compliance_rows() -> list[dict[str, str]]:
    path = ROOT / "control/compliance_matrix.md"
    rows = []
    keys = [
        "matrix_id", "clause_id", "source", "requirement", "force", "star",
        "decision", "technical_position", "proposal_position", "verification", "evidence",
        "client", "vendor", "deviation_status", "maturity", "owner",
        "checker", "issue_risk", "notes",
    ]
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("| CM-"):
            continue
        values = [value.strip() for value in line.strip().strip("|").split("|")]
        if len(values) < len(keys):
            values += [""] * (len(keys) - len(values))
        rows.append(dict(zip(keys, values[:len(keys)])))
    return rows


def bid_location(clause_id: str) -> str:
    match = re.search(r"FR-(\d{2})", clause_id)
    if match:
        return f"第4.{int(match.group(1))}节；第5章；第11.1节"
    prefix = clause_id.split("-")[0]
    mapping = {
        "STD": "第3.2节、第6章、第7章",
        "ARCH": "第3章", "DEPLOY": "第3.7、3.13节及第7.4节",
        "TECH": "第3.5节及第5章", "REQ": "第2章及第3.16节",
        "DATA": "第3.6节、第4.15节及附录D", "NFR": "第7章",
        "PE": "第7.1、7.6节及附录E", "INT": "第3.10、3.14节、第4.11节及附录B",
        "IMPL": "第8.2节", "PM": "第8.4至8.8节", "TRAIN": "第10.1节",
        "SLA": "第10.2、10.3、10.6节", "RISK": "第8.3节",
        "ACC": "第9.9节及附录E", "DEL": "第8.6节及附录G",
        "DOC": "第9.6节及附录G", "TB": "附录F",
    }
    return mapping.get(prefix, "第2至10章相关小节；附录C")


DOMAIN_DESIGNS = {
    1: {
        "name": "应急预案管理",
        "purpose": "把纸质预案转化为可发布、可版本化、可执行的数字预案。系统区分综合预案、专项预案和现场处置方案，并按区域、时段和风险情景细化。已发布版本在事件启动时形成不可变快照，防止后续修订改变历史处置依据。",
        "components": "预案基本信息、分层分类、处置流程、任务模板、会议群组、紧急通知人、调派人员、附件、发布审批、版本比较和检索。",
        "flow": "编制人创建草稿并配置流程、任务和人员，提交审核后发布。事件核实通过时，授权指挥人员选择适用版本启动；系统在同一事务边界内生成预案快照、处置实例、任务台账和消息待发送记录。",
        "data": "核心对象为预案、版本、流程定义、节点、任务模板、人员组、附件和启动快照。状态采用草稿、审核中、已发布、已停用；已发布版本只允许通过新版本修订。",
        "interfaces": "依赖中台身份、组织、权限、工作流和文件存储；启动后调用统一消息通道，并向事件、任务和指挥模块提供快照。",
        "exception": "人员或群组失效时阻断发布并给出差异；消息通道故障不回滚已合法启动的事件，待发送记录进入重试和人工联系清单；附件服务不可用时禁止发布缺失强制附件的版本。",
        "permission": "编制、审核、发布、停用和启动分权；发布与启动均记录操作者、时间、版本、理由和关联事件。普通查看者不得取得超出数据范围的联系人或附件。",
        "verify": "以综合、专项和现场三类样例走通编制、审核、发布、修订和启动。核对启动前后快照一致性，现场测量通知与任务入队不超过3秒，并留存回执和审计。",
    },
    2: {
        "name": "应急指挥中心",
        "purpose": "围绕单一事件汇集续报、处置任务、人员位置、周边视频和物资站点，形成带时间含义的一张图和指挥工作台。界面明确区分事实、待确认信息、外部系统异常和人工处置记录。",
        "components": "事件总览、时间线、任务看板、人员定位、视频窗口、物资图层、指令下发、催办、联动状态和异常提示。",
        "flow": "打开事件后加载事件快照和最新续报，再并行请求任务、人员位置、视频目录和物资站点。调度员选择对象下发指令；任务和消息回执持续更新。事件关闭前检查未完成任务和未处理联动异常。",
        "data": "事件、续报、任务、位置快照、视频通道引用、物资站点、指挥指令和联动日志通过event_id关联。位置和设备状态保留来源时间，避免旧数据被误认为实时。",
        "interfaces": "依赖地图、定位、视频、消息、中台组织权限和物资数据；历史录像由既有视频系统保存，本项目只调阅、回放和验证。",
        "exception": "定位超时显示最后更新时间并降级到人工报位；视频断流允许切换备选通道并保留错误码；地图不可用时以列表和楼层文字继续指挥；外部回执失败进入告警。",
        "permission": "指挥人员可下发和催办，安保和值班人员按授权查看，门禁指令需授权确认。所有指令、视频调阅和导出都进入审计。",
        "verify": "使用模拟火情和人员密集事件验证多源并发加载、时间线、定位不超过2秒刷新、视频首帧不超过3秒、物资图层与任务同步；注入断流和超时验证降级。",
    },
    3: {
        "name": "应急资源管理",
        "purpose": "建立人员、小组、职责、值班计划、物资站点和物资台账的可查询底账，使事件启动时能够按职责和空间快速调派，并通过盘点保证数量与有效期可控。",
        "components": "人员档案、应急小组、职责、调派通知、排班、站点、物资类别、物资台账、有效期提醒、盘点计划、盘点任务和差异处理。",
        "flow": "管理员维护人员、小组和值班计划；物资管理员维护站点和物资批次，生成盘点计划。盘点人移动执行后提交实盘数，系统计算账实差异，复核确认后更新数量并保留调整记录。",
        "data": "人员、小组、职责、排班、站点、物资、批次、盘点计划、盘点明细和差异单。物资数量变更使用流水记录，有效期按批次管理。",
        "interfaces": "依赖中台组织人员、地图点位、统一消息和H5盘点；甲方提供源数据并确认正确性，乙方负责模板、清洗映射、导入和技术校验。",
        "exception": "停用人员从新任务候选中排除但保留历史；盘点重复提交以幂等键拒绝；过期物资不得作为可调派库存；数据校验失败形成错误清单而不静默丢弃。",
        "permission": "人员敏感字段按角色和组织范围查看；物资调整由盘点与复核角色分离；导入、导出、数量调整和有效期修改全部审计。",
        "verify": "验证不少于300人、30个站点、1000项物资规模下的检索和调派；走通盘点计划、移动实盘、差异复核、临期提醒和导入对账。",
    },
    4: {
        "name": "应急演练管理",
        "purpose": "把演练频次、计划、任务、执行、评估和改进纳入闭环，支持月、季、年周期，并满足人员密集场所每半年至少一次的监管频次要求。",
        "components": "演练计划、周期规则、参与人员、演练任务、执行记录、附件、评估模板、评估报告、问题项和改进跟踪。",
        "flow": "制定计划并选择预案版本，系统按周期生成任务并通知执行人。执行人按流程反馈，结束后依据条目化模板评分和记录问题，责任人制定改进措施并跟踪关闭。",
        "data": "演练计划、实例、任务、执行证据、评估答案、问题和改进措施均关联预案版本。状态为计划、已下发、执行中、待评估、改进中、已关闭。",
        "interfaces": "依赖预案、任务、消息、H5、文件存储和中台工作流；真实联合演练需要甲方提供现场窗口和参与人员。",
        "exception": "周期任务生成采用唯一周期键防止重复；参与人变更时形成替补或重新指派；移动网络中断允许本地暂存，恢复后按时间顺序补传并提示冲突。",
        "permission": "计划审批、演练执行、评估和改进关闭分角色控制；演练证据和评分修改保留前后值。",
        "verify": "以月度和半年度计划验证自动下发、移动执行、模板评估、问题整改和报告输出；检查重复调度、人员变更和离线补传。",
    },
    5: {
        "name": "应急事件管理",
        "purpose": "实现Web、H5和既有告警来源的统一接报，经过核实审批、预案启动、任务处置、关闭评估和事故调查形成完整生命周期。",
        "components": "事件上报、告警转事件、核实审批、事件列表、续报、预案启动、任务催办、临时任务、关闭评估、事故调查、报告和知识沉淀。",
        "flow": "来源数据先形成待核实事件，审批通过后进入已确认状态；授权人员启动预案并下发任务。处置期间接收续报和反馈，满足关闭条件后评估关闭；需调查的事件生成调查任务和报告。",
        "data": "事件主表保存当前状态，状态历史追加记录每次迁移；上报、告警原文、审批、预案快照、任务、消息、附件、评估和调查通过事件标识关联。",
        "interfaces": "依赖物联网和安防消防告警、中台工作流、消息、文件、预案、任务和H5；报警接入不超过2秒，事件确认到任务下达不超过3分钟。",
        "exception": "重复告警按来源标识去重；审批超时催办；预案启动部分失败保留已完成步骤并补偿；关闭时存在未完成强制任务则阻断并说明原因。",
        "permission": "上报、核实、启动、关闭和调查具有独立权限；自动告警保留原始来源。状态回退、强制关闭和报告导出均审计。",
        "verify": "分别用Web、H5和物联网告警走通全链路；测量告警接入和任务下达；验证非法迁移、重复告警、启动部分失败和关闭门禁。",
    },
    6: {
        "name": "应急值班管理",
        "purpose": "将小组、时段、打卡点位和预警规则配置化，实现扫码打卡、范围校验、就位统计和缺卡提醒，为24小时值班与带班提供证据。",
        "components": "打卡小组、班次、代执行、点位、地图拾取、二维码、有效半径、打卡任务、记录、统计、预警和明细导出。",
        "flow": "管理员配置小组、时段和点位，系统按计划生成任务。值班人员在H5扫码并提交位置，服务端校验身份、时段、点位和范围，写入记录；规则任务统计应到实到并发送提醒。",
        "data": "小组、成员、班次、点位、二维码版本、打卡任务、记录和预警。记录包含服务器时间、源定位、距离、终端和校验结论。",
        "interfaces": "依赖地图点位、H5扫码定位、中台组织身份和统一消息；甲方提供定位和宿主能力，乙方负责校验和留痕。",
        "exception": "二维码过期、身份不匹配、超出时段或范围时拒绝并说明；定位不可用时不得伪造成功，可按授权人工补录并标记来源；重复提交幂等处理。",
        "permission": "普通人员仅查看个人任务和记录；值班管理员配置计划；人工补录和代执行需专门权限并记录原因。",
        "verify": "验证扫码写入回传不超过1秒、点位/小组/人员三个统计维度、缺卡消息和明细导出；测试越界、过期码、重复提交和人工补录。",
    },
    7: {
        "name": "应急基础管理",
        "purpose": "以受控配置支撑预案类型、事件类型、物资站点、评估模板和核实审批流程，使业务变化通过版本化配置落地。",
        "components": "类型字典、事件分级、站点空间配置、评估模板、流程设计、节点时限、启停、版本比较、发布审批和影响分析。",
        "flow": "管理员创建或修订配置，系统校验引用完整性，经审批发布。新业务实例使用新版本，历史实例继续引用原版本；停用前检查仍在使用的对象。",
        "data": "字典、模板、流程定义、节点、规则、版本和引用关系。配置状态采用草稿、审核中、已发布、已停用。",
        "interfaces": "依赖中台工作流、文件、身份权限和地图；向预案、事件、演练、物资模块提供只读发布版本。",
        "exception": "循环流程、缺少处理人、无时限或引用失效时禁止发布；发布失败回滚到上一稳定版；紧急修改须走书面变更和补审。",
        "permission": "配置编辑、审核和发布分离；每次修改保留差异、操作者和批准记录。",
        "verify": "建立四类事件、预案类型、站点、评估模板和多级核实流程，验证发布、引用、修订、停用和历史实例不受影响。",
    },
    8: {
        "name": "应急管理移动端",
        "purpose": "以H5嵌入甲方现有智慧管理APP，为现场人员提供事件、任务、知识库、演练、打卡和盘点六类能力，不另建原生APP。",
        "components": "统一登录桥接、消息跳转、事件拍照上报、任务确认反馈、知识检索、演练执行、扫码定位打卡、物资盘点、离线暂存和上传队列。",
        "flow": "宿主提供登录态和设备能力，H5换取业务会话并加载个人待办。提交时生成客户端请求号，服务端幂等处理；弱网时保存待上传记录，恢复后重传并显示最终回执。",
        "data": "移动会话、设备能力、事件草稿、任务反馈、附件上传、离线队列、打卡和盘点记录。敏感令牌不写入可持久明文存储。",
        "interfaces": "依赖宿主APP登录、推送、定位、扫码、拍照、文件上传和网络；甲方负责宿主签名与发布，乙方提供H5包及适配说明。",
        "exception": "宿主能力缺失时给出兼容提示；上传中断可续传或重试；重复点击不产生重复业务；定位权限被拒时按规则阻断或转人工。",
        "permission": "功能和数据范围与中台身份一致；高风险操作需要再次确认。移动端操作保留用户、设备、时间和结果。",
        "verify": "在Android、iOS及目标宿主矩阵验证六类功能、消息深链、扫码、定位、拍照上传、弱网恢复和会话失效；最终版本【待人工确认】。",
    },
    9: {
        "name": "应急信息视图",
        "purpose": "汇总事件、任务、资源、演练和值班指标，支持筛选、钻取和回到业务明细。统计结果注明时间范围、刷新时间和数据口径。",
        "components": "事件趋势、事件分类、任务状态、处置时长、资源与值班、演练改进、筛选器、钻取和报表导出。",
        "flow": "用户选择时间和组织范围，服务端按统一口径查询聚合数据；点击指标下钻到明细并保留筛选条件。默认刷新不超过60秒，关键告警通过事件驱动局部更新。",
        "data": "使用业务明细、统计快照、指标定义和口径版本。报表保存生成时间、条件和操作者。",
        "interfaces": "依赖事件、任务、资源、值班、演练和中台组织；可向统一数据汇聚接口提供受控指标。",
        "exception": "部分数据源延迟时显示数据时间和缺失范围；大时间跨度采用异步生成；聚合失败不显示误导性零值。",
        "permission": "按组织和角色限制统计与明细；导出需要权限并审计。",
        "verify": "验证筛选、钻取、口径一致性、默认刷新、年度报表不超过5秒和权限隔离；抽样以明细重新计算指标。",
    },
    10: {
        "name": "综合安防视图",
        "purpose": "集中展示安防终端在线状态和告警统计，并从告警定位到地图或既有系统页面。该视图消费既有系统数据，不替代安防子系统。",
        "components": "设备分类统计、在线离线、告警趋势、告警列表、地图定位、事件关联、外部跳转、接口健康和数据时间提示。",
        "flow": "适配层接收或轮询设备与告警数据，统一编码后更新视图。用户从统计下钻到设备或告警，再选择转事件或跳转既有系统。",
        "data": "设备引用、状态快照、告警原文、标准告警、事件关联和接口健康记录。原始来源标识保留用于对账。",
        "interfaces": "依赖物联网、入侵、消防、门禁、视频及地图；综合视图刷新不超过30秒。",
        "exception": "接口断连时保留最后成功时间并告警，恢复后补传或对账；未知设备进入待映射清单；跳转失败保留原告警信息。",
        "permission": "安保和指挥角色按范围查看；告警转事件、门禁联动和导出具有独立权限。",
        "verify": "模拟在线、离线和多类告警，验证不超过30秒刷新、定位、下钻、转事件、断连重试补传和数据对账。",
    },
    11: {
        "name": "数据接口",
        "purpose": "以统一适配层连接中台、物联网、视频、信息发布、消息、定位、门禁和安防消防，隔离厂商差异并形成可监测、可降级、可审计的集成边界。",
        "components": "API网关、认证适配、字段映射、协议适配、幂等、超时重试、熔断、回执、补传、对账、健康检查、版本管理和审计。",
        "flow": "接口先完成契约、账号、环境和窗口登记，再进行连通、正常、异常、性能和恢复验证。运行时请求带关联标识，适配器把外部响应转换为内部统一状态。",
        "data": "接口目录、契约版本、调用记录、请求摘要、响应码、耗时、重试次数、业务关联、回执和对账差异。敏感报文按规则脱敏。",
        "interfaces": "甲方提供或协调资料、账号、授权、环境和窗口；乙方逐接口适配验证。视频和消息在M2前完成真实连通性验证。",
        "exception": "超时采用有限重试和熔断；写操作使用幂等键；断连期间可排队数据保留待补传，控制指令立即告警并转人工。",
        "permission": "服务账号最小权限，密钥与配置分离；高风险门禁指令保留授权人和回执。接口调用和配置变更纳入审计。",
        "verify": "逐接口执行鉴权、字段映射、正常、异常、超时、重复、断连恢复、性能和降级测试；形成契约、日志、联调报告和问题关闭证据。",
    },
}


def append_front_and_needs(lines: list[str]) -> None:
    add_block(lines, """
    # 某自然博物馆智能运营中心建设项目——应急管理子系统

    ## 投标文件技术标

    - 文档版本：V1.0（G1-08重构送审稿）
    - 状态：REVIEW
    - 主责：A（项目经理 PM / 技术标总编）
    - 技术复核：B
    - 合规复核：C
    - 冻结输入：BASELINE-G1-V0.1
    - 编制日期：2026年9月

    ## 编制说明与证据边界

    本文件以真实用户需求书、冻结基线、G1-03至G1-07工作成果、facts、key_numbers和compliance_matrix为内容依据。教学参考技术标只提供章节、信息类型、图表密度和Word格式，不向本项目引入其业务、技术栈、算法、人员、业绩、指标或承诺。

    文中“响应”表示方案已给出设计或管理安排，不等于现场实施和验收已经完成。接口连通、性能、安全、部署、演练和最终验收证据仍按项目计划形成，统一标记为PENDING_EVIDENCE。人员姓名、资历、类似业绩、合同编号等没有真实资料的内容标记【待人工确认】。

    # 第一章 投标响应总述

    ## 1.1 投标响应声明

    我方已按冻结基线理解本项目的范围、39条功能需求、34条带★功能条款、5项非功能或部署类★条款、20项法规标准和全部实施、服务、验收要求。本方案对上述要求均作响应，未提出正偏离或负偏离。历史录像回放、人员定位、疏散门禁、统一消息、中台、地图、H5、部署环境、基础数据、外部接口和应用安全均按12项课程模拟甲方书面澄清执行。

    我方承诺以需求、设计、测试和验收双向追踪组织交付。现阶段承诺边界以用户需求书和冻结基线为准；具体编程语言、数据库、消息中间件和GIS产品尚未冻结，不在投标阶段虚构产品选型。后续设计评审根据甲方环境、接口和运维条件形成可审查选型记录。

    ## 1.2 方案特点

    方案以事件为业务主线，将预案快照、任务、消息、人员、物资、定位、视频、联动和审计放入同一可追踪链路。统一适配层隔离既有系统差异；配置化流程支持事件类型和审批变化；H5复用现有智慧管理APP；消息台账、接口幂等、熔断降级和恢复演练用于保障应急状态下的可用性。

    ## 1.3 方案导读

    本文件包含十四章与附录A至G。第二章解释需求和场景；第三章给出16项总体设计；第四章对F-01至F-11逐域展开；第五章说明13项关键技术；第六至十章覆盖开源合规、非功能、实施、质量、培训与售后；第十一章逐条响应合规矩阵；第十二、十三章保留人员与业绩待确认结构；第十四章提出不扩大范围的建议。附录提供接口、索引、数据字典、用例框架、评分映射和交付清单。

    # 第二章 需求理解

    ## 2.1 项目背景与现状问题

    某自然博物馆同时承担人员密集公共场所安全和不可再生藏品保护责任，突发事件来源覆盖社会安全、自然灾害、火灾与设备运行事故。现状依赖纸质预案、电话通知和分散记录，造成预案调用慢、多源态势分散、资源底数更新困难、演练和值班证据薄弱、事件关闭后难以复盘。

    建设重点是把日常准备与应急处置连接起来。平时维护预案、人员、物资、值班和演练，事发后直接复用这些数据形成处置实例；事后评估再回到预案和知识库，避免形成只在演示时可用的孤立功能。

    ## 2.2 建设目标与可验证结果

    总体目标为“预案一键启动、指令秒级下达、态势一图总览、处置全程留痕、事后评估改进”。方案把目标拆为可验证结果：启动通知与任务下发不超过3秒；告警接入不超过2秒；人员定位刷新不超过2秒且不降低源精度；视频首帧不超过3秒；正常验收通道支持不少于20路并行且到达率不低于99%；峰值在线不少于100人；试运行可用率不低于99.5%。

    指标是否达成只由约定环境中的原始记录、时间戳、回执和测试报告证明。文档中的设计说明不能替代现场证据。

    ## 2.3 用户角色与业务场景
    """)
    roles = [
        ["应急指挥人员", "接报核实、启动预案、调度、监控和调查", "Web指挥端及F-01、F-02、F-05、F-09、F-10"],
        ["值班人员及带班领导", "到岗打卡、接收预警、上报事件", "H5与Web及F-06、F-08"],
        ["应急处置人员", "接收任务、现场处置、续报反馈", "H5任务及F-03、F-05、F-08"],
        ["演练执行人", "执行演练并提交证据", "F-04、F-08"],
        ["安保人员", "安防值守、视频调阅、先期处置和疏散", "F-02、F-08、F-10"],
        ["物资管理员", "物资台账、盘点、有效期和调派保障", "F-03、F-08"],
        ["系统管理员", "配置、接口、权限、日志和运行检查", "F-07、F-11"],
        ["第三方系统", "提供或接收事件、告警、视频、消息和联动信息", "F-11统一适配层"],
    ]
    lines.append(md_table(["角色", "主要职责", "使用入口与功能"], roles)); lines.append("")
    add_block(lines, """
    用户按职责进入管理、指挥、现场处置和系统维护场景。权限模型同时控制功能入口、组织范围、事件范围和敏感字段。第三方系统以服务账号接入，其权限、凭据、调用范围和审计独立管理。

    ## 2.4 十一个功能域总体分析
    """)
    lines.append(md_table(
        ["功能域", "名称", "总体理解"],
        [[f"F-{i:02d}", DOMAIN_DESIGNS[i]["name"], DOMAIN_DESIGNS[i]["purpose"]] for i in range(1, 12)],
    )); lines.append("")
    add_block(lines, """
    十一个功能域不是彼此孤立的菜单。F-07提供受控配置，F-01把配置固化为可执行预案，F-05产生事件，F-02汇集态势，F-03和F-06提供日常资源和值班底账，F-04验证预案，F-08延伸至现场，F-09和F-10形成统计与安防视图，F-11贯穿外部系统。

    ## 2.5 项目重点难点分析

    ### 2.5.1 秒级启动与可靠消息

    预案启动同时涉及快照、任务和多对象通知。同步串行调用外部通道会把不稳定性带入核心事务。方案先在本地完成事件、快照、任务和待发送台账，再异步并发投递，以业务幂等键、回执状态和补偿队列保证不重不漏。

    ### 2.5.2 空间态势与多源时序

    人员、视频、物资和告警来自不同系统，刷新周期和坐标含义不同。一张图必须显示来源时间、楼层和接口健康，过期信息不得伪装成实时。地图不可用时保留列表和楼层文字降级。

    ### 2.5.3 定位精度与刷新链路

    亚米级能力由甲方定位源提供，乙方负责接口、关联、映射和不超过2秒刷新。验收同时比对源坐标与展示坐标及时间戳，证明链路没有降低源精度。

    ### 2.5.4 安全联动与人工降级

    疏散门禁属于高风险控制。系统只允许授权人员确认后下发，服从既有安全联锁，并保留请求、回执和结果。失败立即告警并转人工处置，系统不绕过门禁安全机制。

    ### 2.5.5 中台复用、隔离部署与合规

    身份、组织、权限、消息、工作流、文件、门户和数据汇聚由甲方中台提供。本项目只做业务适配。依赖组件须能在逻辑隔离环境安装，具体产品和版本在设计评审时通过SBOM、许可证和漏洞审查冻结。

    ## 2.6 技术路线与边界对比
    """)
    lines.append(md_table(
        ["比较维度", "重复建设通用平台", "紧耦合既有厂商", "分层业务服务加适配层"],
        [
            ["范围", "扩大本项目范围", "短期接入快但被单一接口绑定", "复用甲方中台并隔离厂商差异"],
            ["变更", "通用能力重复维护", "接口升级影响业务代码", "适配器独立升级，业务模型稳定"],
            ["应急降级", "责任边界模糊", "故障直接扩散", "超时、熔断、补传和人工降级明确"],
            ["验收", "难区分平台与业务责任", "证据分散", "按接口建立契约、日志和端到端证据"],
        ],
    )); lines.append("")
    add_block(lines, """
    本方案采用分层业务服务加适配层。该路线符合中台复用和不重复建设边界，也能针对每个外部系统独立设置超时、重试、熔断、版本与验证规则。

    ## 2.7 典型应急场景走查

    场景一为火灾报警。消防告警在2秒内接入并形成待核实事件，值班人员核实后启动预案，系统在3秒内生成任务与通知；指挥端加载周边视频、人员和物资，授权人员按规程处理门禁与信息发布，处置结束后评估和调查。

    场景二为人员密集活动异常。客流告警转事件，指挥端查看区域人数、安保人员位置和视频，调度疏导小组并持续接收反馈。定位不可用时以人工报位继续，消息未送达对象进入人工联系清单。

    场景三为馆藏区域入侵。入侵告警保留原始来源，关联附近视频和安保人员。未经授权不得下发高风险控制。视频断流时显示故障并切换备选通道。

    场景四为自然灾害。按专项预案调派人员和物资，H5支持弱网暂存和恢复补传。地图服务中断时以楼栋、楼层和列表降级。

    场景五为值班与日常盘点。系统按计划生成打卡和物资盘点任务，对缺卡、临期和账实差异提醒，形成应急准备证据。

    场景六为演练。周期到达自动下发任务，移动端记录执行，结束后按模板评估并跟踪改进，验证预案能够执行。

    ## 2.8 需求优先级与里程碑映射
    """)
    lines.append(md_table(
        ["优先层级", "需求范围", "里程碑安排", "准出证据"],
        [
            ["首先贯通", "预案、事件、任务、消息、身份权限", "M1确认；M2设计；开发前半段", "SRS、RTM、状态机、启动与消息原型"],
            ["高风险集成", "视频、消息、定位、地图、中台、安防消防", "M2前视频与消息连通；M3前全部联动", "契约、连通、异常与降级记录"],
            ["现场闭环", "H5、值班、盘点、演练、态势视图", "M3初验；M4试运行", "终端兼容、现场流程、运行与演练记录"],
            ["最终门禁", "全部FR、PE/NFR、安全、部署、文档", "M5竣工验收", "100%用例通过、高危0、一次独立部署成功"],
        ],
    )); lines.append("")
    add_block(lines, """
    ## 2.9 十二项课程模拟书面澄清及效力边界

    十二项澄清是本课程模拟中已经确定的甲方解释，用于解决真实需求书中的歧义并冻结当前编标口径。它们不代表现实单位签章或法律文件，也不证明接口和指标已现场验收。
    """)
    clarifications = [
        ["DEC-001", "历史回放强制；既有视频系统负责录像与不少于30天保存，本项目负责调阅、回放和验证"],
        ["DEC-002", "甲方提供定位基础设施和亚米级源数据；乙方负责接口、关联、不超过2秒刷新且不降低源精度"],
        ["DEC-003", "门禁联动应支持；授权确认后下发，服从安全联锁，失败告警并人工降级"],
        ["DEC-004", "甲方提供消息通道和账号配额；乙方负责并发、重试、回执与留痕；正常通道验证不少于20路和99%"],
        ["DEC-005", "甲方中台提供八类通用能力；乙方只做业务适配"],
        ["DEC-006", "甲方提供合法二维、楼层和验收所需三维地图；乙方不负责重测绘或三维建模"],
        ["DEC-007", "移动端为H5嵌入现有APP；乙方交付H5包，甲方负责宿主和发布"],
        ["DEC-008", "服务器、网络、证书、备份介质和后备电源由甲方提供；乙方提交最低配置和部署方案"],
        ["DEC-009", "甲方提供源数据并确认业务正确性；乙方清洗映射、导入和技术校验，成果双方确认"],
        ["DEC-010", "甲方协调接口资料、账号、环境和窗口；乙方逐接口适配验证；视频和消息M2前连通"],
        ["DEC-011", "乙方落实应用侧等保二级相关条款和安全测试，高危清零；正式第三方测评不默认在范围内"],
        ["DEC-012", "统一使用匿名项目名称和招标人名称"],
    ]
    lines.append(md_table(["裁决", "冻结口径"], clarifications)); lines.append("")
    add_block(lines, """
    ## 2.10 同类系统常见失效模式与项目对策

    本项目没有可核验的同类项目业绩资料，因此本节不引用案例业绩，而按需求和风险登记册归纳失效模式。预案与任务脱节时采用启动快照；消息只记录“发送”时采用分层回执台账；多源态势时间不一致时显示来源时间；外部接口失败时执行有限重试、熔断和人工降级；配置覆盖历史时使用版本化发布；验收前集中补文档时通过配置库和追踪矩阵同步维护。

    ## 2.11 建设价值与度量方法

    建设价值通过响应速度、处置透明度、资源可用性、演练闭环和值班证据衡量。度量不虚构节省比例：启动、告警、定位、视频、消息、页面、地图、打卡和报表按需求指标实测；事件任务完成率、未送达率、盘点差异、缺卡、演练问题关闭率按试运行记录统计；最终以甲方确认的试运行报告为准。
    """)


def append_architecture(lines: list[str], tech: str) -> None:
    add_block(lines, "# 第三章 总体技术方案\n\n## 3.1 总体设计原则")
    add_block(lines, source_section(tech, "1.1 架构原则"))
    add_block(lines, """
    ## 3.2 执行的标准与规范

    用户需求书列出的20项法律法规、规章和标准均纳入适用性管理。应急预案、演练、安防、视频联网、应用安全、软件质量、文档和测试分别在需求、设计、测试和验收活动中落实。注日期文件使用指定版本；未注日期文件在设计评审时核验现行版本。标准条款冲突按较高要求执行，仍有歧义时提交甲方书面澄清。逐项映射见第十一章。

    ## 3.3 系统总体架构

    总体架构分为展现与接入、应急应用服务、支撑与数据、外部系统与甲方环境四层。Web、H5和大屏通过统一API进入业务服务；业务服务围绕预案、事件、任务、资源、值班、演练和视图组织；地图、消息、视频和外部系统通过适配器接入；业务数据、空间索引和审计记录形成证据底座。

    [图3-1 总体逻辑架构]
    """)


def append_functions_and_key_tech(lines: list[str], tech: str, fr_rows: dict[int, list[dict[str, str]]]) -> None:
    add_block(lines, "# 第四章 功能实现方案")
    for idx in range(1, 12):
        design = DOMAIN_DESIGNS[idx]
        add_block(lines, f"## 4.{idx} F-{idx:02d} {design['name']}\n\n### 4.{idx}.1 需求理解\n\n{design['purpose']}")
        reqs = fr_rows.get(idx, [])
        if reqs:
            lines.append(md_table(
                ["需求", "★", "实现响应", "当前证据状态"],
                [[r["id"], r["star"], r["summary"], "设计覆盖；实施和测试证据待形成"] for r in reqs],
            ))
            lines.append("")
        parts = [
            ("components", "功能组成"), ("flow", "核心流程"), ("data", "数据与状态"),
            ("interfaces", "接口依赖"), ("exception", "异常与降级"),
            ("permission", "权限与审计"), ("verify", "验证方式"),
        ]
        for leaf, (key, title) in enumerate(parts, start=2):
            add_block(lines, f"### 4.{idx}.{leaf} {title}\n\n{design[key]}")
    add_block(lines, """
    ## 4.12 界面清单与导航结构

    Web一级入口包括应急工作台、预案、事件、指挥、资源、演练、值班、基础配置、统计视图、综合安防和系统运维。H5一级入口包括事件、任务、知识库、演练、打卡和盘点。关键操作从首页或工作台不超过3次页面跳转可达。事件详情作为跨模块上下文入口，避免用户在多个菜单重复查找。

    ## 4.13 异常处理与降级设计

    异常分为业务校验、权限拒绝、外部接口、数据质量、资源容量和基础环境六类。客户端展示用户可行动的提示，服务端记录结构化错误码、关联标识和上下文。外部能力失败时按可重试、不可重试和需人工处理分类；控制类指令失败不得静默重试造成重复动作。

    ## 4.14 浏览器、H5兼容与界面规范

    Web支持Chrome和Edge最新两个稳定版本；H5在甲方目标Android、iOS和宿主APP版本矩阵中验证。表单提供明确必填和错误定位；状态使用文字与颜色共同表达；时间统一采用北京时间并在接口中使用明确时区。最终宿主和终端版本【待人工确认】。

    ## 4.15 数据初始化与预置内容

    乙方提供人员、组织、预案、物资、站点、点位、字典和模板的导入模板，执行格式、必填、引用、重复和业务规则检查。甲方提供真实源数据并确认业务正确性。乙方保留映射、清洗、导入、失败和对账记录，最终成果双方确认。

    ## 4.16 角色—功能—数据权限矩阵
    """)
    lines.append(md_table(
        ["角色", "主要可操作功能", "受限操作", "数据范围"],
        [
            ["指挥长/调度员", "核实、启动、调度、关闭、调查", "门禁等高风险联动需授权确认", "授权事件和组织范围"],
            ["值班/安保", "上报、打卡、查看任务和态势", "不得发布预案或修改基础配置", "本人、班组及授权区域"],
            ["处置人员", "接收、反馈、上传和完成任务", "不得关闭事件", "本人任务与关联事件必要信息"],
            ["物资管理员", "台账、盘点和差异提交", "数量调整需复核", "授权站点"],
            ["演练管理员/执行人", "计划、执行、评估或反馈", "按职责分离", "授权演练"],
            ["系统管理员", "配置、接口、日志和运行检查", "不默认取得业务敏感内容", "系统配置及授权审计范围"],
        ],
    )); lines.append("")

    add_block(lines, """
    # 第五章 关键技术实现方案

    本章对关键技术统一按问题、设计、关键机制、异常场景和验证方法展开。具体产品选型留待M2评审，不影响机制与验收口径。
    """)
    tech_sections = [
        ("秒级预案启动", "预案启动需在3秒内形成通知与任务。采用本地事务写入事件、快照、任务和待发送消息，再由工作线程并行投递。以同一启动标识避免重复启动，以时间戳分解数据库、编排和通道耗时。数据库失败整体回滚；外部通道失败进入重试，不回滚已合法启动的事件。验证在目标数据规模和正常验收通道实测端到端时间。"),
        ("可靠消息推送", source_section(tech, "2.2 消息可靠推送")),
        ("任务编排与事件状态机", "事件、任务和消息分别建立受控状态机，跨对象通过领域事件协调。任务生成保存预案快照和处理人快照，催办与反馈不改变原始指令。消费者使用幂等键和去重表；定时扫描处理超时与重试。异常注入覆盖重复投递、乱序回执、进程重启和永久失败。验证以状态历史和业务数量守恒证明不重不漏。"),
        ("事件态势一张图", source_section(tech, "2.3 事件态势一张图")),
        ("定位链路", "定位源由甲方提供。适配层接收人员标识、坐标、楼层、精度和源时间，完成身份与楼层映射后推送前端。缓存只用于削峰并保留过期时间；展示必须标出最后更新时间。异常包括人员无法关联、坐标越界、楼层未知和源中断。验收对比源与展示时间戳和控制点误差，确认刷新不超过2秒且不降低源精度。"),
        ("视频GB/T 28181接入与回放", source_section(tech, "2.4 跨系统联动") + "\n\n视频适配覆盖设备目录、实时流、录像检索和回放控制。历史录像由既有系统保存不少于30天。本项目记录请求、流地址取得和首帧时间；验证跨30天边界检索、首帧不超过3秒、断流恢复和鉴权失败。"),
        ("门禁安全联动", "门禁开启采用授权确认、业务幂等、既有联锁和回执核验。系统不直接改变门禁安全逻辑。超时不等同于失败或成功，先查询结果；无法确认时告警并转人工。验证覆盖无权限、取消确认、联锁拒绝、重复请求、超时、明确失败和人工处置记录。"),
        ("统一中台适配", "身份、组织、权限、消息、工作流、文件、门户和数据汇聚均使用甲方中台。适配层将中台标识映射为应急业务标识，并缓存最小必要数据。接口升级通过契约版本和回归测试控制。中台不可用时的会话和写操作降级在安全方案中确定；不得自行建设另一套通用中台。"),
        ("H5宿主集成", source_section(tech, "2.6 H5 扫码与定位")),
        ("配置化流程", source_section(tech, "2.5 配置化流程")),
        ("接口容错与降级", "适配器为每类外部系统设置独立连接池、超时、有限重试、熔断、限流和健康检查。读接口可降级到最后成功数据并标时间；可补传写接口进入待处理队列；控制类写接口失败立即告警。验证通过故障注入检查故障不向核心事件事务扩散，恢复后能够对账。"),
        ("性能与容量", "性能设计依据不少于100人应急峰值、2万事件、20万任务、7.3万打卡记录每年等基线。读写分路径优化，常用查询建立组合索引，地图按视口加载，统计使用预聚合或受控缓存，附件与结构化数据分离。压测同时观察响应分位数、错误率、资源使用和队列积压，禁止只报告平均值。"),
        ("备份恢复与连续运行", source_section(tech, "2.7 可靠性与恢复")),
    ]
    for idx, (title, body) in enumerate(tech_sections, start=1):
        add_block(lines, f"## 5.{idx} {title}\n\n### 5.{idx}.1 问题与目标\n\n{body}\n\n### 5.{idx}.2 关键机制\n\n机制通过可配置参数、稳定业务标识、结构化日志和状态历史实现，关键动作形成可关联证据。设计评审冻结具体组件和参数，运行期变更进入配置与变更控制。\n\n### 5.{idx}.3 异常场景与验证\n\n测试覆盖正常路径、边界输入、重复请求、超时、依赖中断、恢复和权限拒绝。验证记录包含环境、数据、时间戳、请求与回执、结果和问题关闭证据，未取得现场条件时保持PENDING_EVIDENCE。")
        if idx == 2:
            add_block(lines, "[图5-1 可靠消息投递与人工降级]")


def build_bid_markdown() -> None:
    tech = (ROOT / "docs/work/technical_solution_v0.6.md").read_text(encoding="utf-8")
    lines: list[str] = []
    append_front_and_needs(lines)
    append_architecture(lines, tech)
    append_architecture_remainder(lines, tech)
    append_architecture_tail(lines, tech)
    append_functions_and_key_tech(lines, tech, parse_fr_rows())
    append_remaining_chapters(lines)
    BID_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {BID_MD.relative_to(ROOT)} lines={len(lines)} chars={len(BID_MD.read_text(encoding='utf-8'))}")


def _clear_template_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "请在 Word 中更新域"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def _clean_markdown(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^>\s*", "", text)
    text = text.replace("**", "").replace("`", "")
    return text


def _style_document(doc: Document) -> None:
    palette = {
        "Heading 1": (16, "1F4E79"),
        "Heading 2": (14, "2F75B5"),
        "Heading 3": (12, "365F91"),
    }
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Pt(21)
    for style_name, (size, color) in palette.items():
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)


def _write_headers_and_footers(doc: Document) -> None:
    title = "某自然博物馆智能运营中心建设项目——应急管理子系统"
    for section in doc.sections:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        header = section.header
        for paragraph in header.paragraphs:
            paragraph.clear()
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = hp.add_run(title + "  技术投标书")
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(89, 89, 89)
        footer = section.footer
        for paragraph in footer.paragraphs:
            paragraph.clear()
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("第 ")
        fr.font.size = Pt(8)
        _add_field(fp, "PAGE")
        fp.add_run(" 页")


def _extract_reused_figures() -> dict[str, Path]:
    out = ROOT / ".tmp/g1_08/reused_figures"
    out.mkdir(parents=True, exist_ok=True)
    mapping = {
        "图8-1": "word/media/image18.png",
        "图8-2": "word/media/image19.png",
        "图8-3": "word/media/image20.png",
        "图8-4": "word/media/image21.png",
    }
    result: dict[str, Path] = {}
    cached = all((out / Path(member).name).exists() for member in mapping.values())
    archive = None if cached else zipfile.ZipFile(EXISTING_BID)
    try:
        for label, member in mapping.items():
            target = out / Path(member).name
            if not target.exists():
                target.write_bytes(archive.read(member))
            result[label] = target
    finally:
        if archive is not None:
            archive.close()
    return result


def _figure_map() -> dict[str, Path]:
    figures = _extract_reused_figures()
    slides = ROOT / ".tmp/g1_08_ppt"
    figures.update({
        "图3-1": slides / "slide-1.png",
        "图3-2": slides / "slide-2.png",
        "图3-3": slides / "slide-3.png",
        "图3-4": slides / "slide-4.png",
        "图5-1": slides / "slide-5.png",
        "图9-1": slides / "slide-6.png",
    })
    missing = [str(path) for path in figures.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing figure previews: " + ", ".join(missing))
    return figures


def _add_figure(doc: Document, label_line: str, figures: dict[str, Path]) -> None:
    match = re.match(r"\[(图\d+-\d+)\s+(.+)\]", label_line)
    if not match:
        return
    label, caption = match.groups()
    path = figures[label]
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.1))
    cp = doc.add_paragraph(f"{label} {caption}")
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = False
    for run in cp.runs:
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(9)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Normal Table"
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "A6A6A6")
        borders.append(border)
    tbl_pr.append(borders)
    small = cols >= 6
    for row_index, values in enumerate(rows):
        for column_index in range(cols):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[column_index] if column_index < len(values) else ""
            cell.text = _clean_markdown(value).replace("<br>", "\n")
            if row_index == 0:
                _set_cell_shading(cell, "D9EAF7")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "宋体"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                    run.font.size = Pt(7 if small else 9)
                    run.font.bold = row_index == 0
    _set_repeat_table_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _set_landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.3)
    section.right_margin = Cm(1.3)


def _set_portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    if section.page_width > section.page_height:
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)


def build_docx() -> None:
    figures = _figure_map()
    doc = Document(REFERENCE)
    _clear_template_body(doc)
    _style_document(doc)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(90)
    r = cover.add_run("某自然博物馆智能运营中心建设项目\n——应急管理子系统")
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(24)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(38)
    sr = subtitle.add_run("投 标 文 件\n（技术标）")
    sr.bold = True
    sr.font.name = "Microsoft YaHei"
    sr._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    sr.font.size = Pt(30)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(80)
    mr = meta.add_run("项目编号：03\n投标人：严宇、何思源、任俊强\n法定代表人或授权代表：严宇\n日期：2026年9月10日\n用途：仅课程模拟")
    mr.font.name = "宋体"
    mr._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    mr.font.size = Pt(13)
    doc.add_page_break()
    toc_title = doc.add_paragraph("目 录", style="Title")
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc = doc.add_paragraph()
    _add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    doc.add_page_break()

    markdown_lines = BID_MD.read_text(encoding="utf-8").splitlines()
    index = 0
    started = False
    in_landscape = False
    while index < len(markdown_lines):
        raw = markdown_lines[index].rstrip()
        stripped = raw.strip()
        if not started:
            if re.match(r"^#{1,3}\s+编制说明", stripped):
                started = True
            else:
                index += 1
                continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(markdown_lines):
            table_lines = []
            while index < len(markdown_lines) and markdown_lines[index].strip().startswith("|"):
                table_lines.append(markdown_lines[index].strip())
                index += 1
            parsed = []
            for table_line in table_lines:
                values = [part.strip() for part in table_line.strip("|").split("|")]
                if all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    continue
                parsed.append(values)
            _add_table(doc, parsed)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = _clean_markdown(heading.group(2))
            if level == 1 and title.startswith("第十一章") and not in_landscape:
                _set_landscape(doc.add_section(WD_SECTION.NEW_PAGE))
                in_landscape = True
            elif level == 1 and title.startswith("第十二章") and in_landscape:
                _set_portrait(doc.add_section(WD_SECTION.NEW_PAGE))
                in_landscape = False
            elif level == 1:
                doc.add_page_break()
            doc.add_paragraph(title, style=f"Heading {level}")
            index += 1
            continue
        if re.match(r"^\[图\d+-\d+\s+.+\]$", stripped):
            _add_figure(doc, stripped, figures)
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(_clean_markdown(stripped[2:]), style="List Paragraph")
            paragraph.style = doc.styles["List Paragraph"]
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.37)
        else:
            paragraph = doc.add_paragraph(_clean_markdown(stripped))
            if len(stripped) <= 360:
                paragraph.paragraph_format.keep_together = True
            if raw.lstrip().startswith(">"):
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.right_indent = Cm(0.37)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(31, 78, 121)
        index += 1

    _write_headers_and_footers(doc)
    # The retained template starts the first section at page 1.  Later
    # landscape/portrait sections must continue that sequence rather than
    # restart page numbering at 1.
    for section in list(doc.sections)[1:]:
        pg = section._sectPr.find(qn("w:pgNumType"))
        if pg is not None:
            section._sectPr.remove(pg)
    doc.core_properties.title = "某自然博物馆智能运营中心建设项目——应急管理子系统 投标文件技术标"
    doc.core_properties.subject = "G1-08 技术投标书统一修订候选稿"
    doc.core_properties.author = "成员A（项目经理/技术标总编）"
    doc.core_properties.comments = "项目编号03；仅用于课程模拟；人员方案按用户确认的三名成员编制，本次不主张业绩。"
    EXISTING_BID.parent.mkdir(parents=True, exist_ok=True)
    doc.save(EXISTING_BID)
    check = Document(EXISTING_BID)
    headings = [p for p in check.paragraphs if p.style.name.startswith("Heading")]
    print(
        f"Wrote {EXISTING_BID.relative_to(ROOT)} paragraphs={len(check.paragraphs)} "
        f"tables={len(check.tables)} images={len(check.inline_shapes)} headings={len(headings)} sections={len(check.sections)}"
    )


def append_architecture_remainder(lines: list[str], tech: str) -> None:
    add_block(lines, source_section(tech, "1.2 总体逻辑架构"))
    add_block(lines, """
    ## 3.4 业务架构

    业务架构按准备、监测、处置和改进组织。准备阶段维护预案、人员、物资、值班和演练；监测阶段接收Web/H5上报与既有告警；处置阶段完成核实、启动、任务和联动；改进阶段形成评估、调查、知识和预案修订。事件标识贯穿全部阶段。

    [图3-2 应急管理业务闭环]

    ## 3.5 技术架构与选型

    投标阶段冻结能力边界，不虚构具体产品。展现层需支持主流Chrome/Edge最新两个稳定版本及目标Android、iOS宿主；服务层采用开放、成熟、可在隔离环境部署的模块化框架；数据层满足事务、空间和审计需求；异步能力支持可靠投递、定时任务与重试；接口层支持REST/OpenAPI 3.0并适配GB/T 28181。

    ### 3.5.1 关键技术选型及理由

    具体框架、数据库、消息组件、GIS引擎和监控组件在M2设计评审前通过候选对比确定。对比维度包括业务适配、离线安装、许可证、漏洞、国产环境兼容、运维复杂度、恢复能力和团队掌握程度。

    ### 3.5.2 选型合规性说明

    候选组件必须进入SBOM，记录名称、版本、许可证、来源、漏洞和替换方案。包含运行时在线激活、不可离线安装、许可证冲突或高危漏洞无法清零的组件不得准入。

    ## 3.6 数据架构
    """)


def append_remaining_chapters(lines: list[str]) -> None:
    plan = (ROOT / "docs/work/A_PM/project_plan_v1.md").read_text(encoding="utf-8")
    risk = (ROOT / "docs/work/A_PM/risk_register_v1.md").read_text(encoding="utf-8")
    add_block(lines, """
    # 第六章 开源组件与许可证合规

    ## 6.1 合规管理策略

    投标阶段不虚构尚未冻结的具体技术产品和版本。M2设计评审前建立候选组件清单和软件物料清单SBOM，记录组件名称、版本、来源、许可证、直接或间接依赖、已知漏洞、使用方式、修改情况、替代方案和批准人。开发、测试、镜像和交付环境使用同一锁定清单。

    ## 6.2 开源组件清单

    组件清单按Web前端、H5、服务框架、数据访问、空间处理、任务与消息、视频适配、接口文档、日志监控、安全扫描、测试与构建分类维护。具体名称和版本【待M2技术选型确认】。未确认项不得被写成正式承诺，确认后须通过许可证与漏洞门禁。
    """)
    lines.append(md_table(
        ["类别", "所需能力", "准入证据", "当前状态"],
        [
            ["Web/H5框架", "组件化界面、浏览器与宿主适配", "版本、许可证、构建锁文件、兼容记录", "【待M2确认】"],
            ["服务与数据", "业务模块、事务、查询、空间索引", "选型对比、许可证、容量与恢复验证", "【待M2确认】"],
            ["异步与调度", "可靠投递、周期任务、失败重试", "消息语义、持久化、故障演练", "【待M2确认】"],
            ["视频与地图", "GB/T 28181、地图加载与点位", "协议和授权边界、接口验证", "【待M2确认】"],
            ["安全与测试", "SCA、扫描、测试与报告", "工具版本、规则库、结果和复测", "【待M2确认】"],
        ],
    )); lines.append("")
    add_block(lines, """
    ## 6.3 重点合规问题分析

    许可证评审同时检查源代码修改义务、再分发义务、版权与NOTICE保留、专利条款、商标限制、网络服务条款和商业插件授权。强传染性或与交付方式冲突的许可证需替换或获得书面法律评估。第三方SDK须确认甲方已有合法授权和离线运行条件。

    ## 6.4 知识产权承诺

    交付包只包含经批准的自研成果、甲方授权资料和合规第三方组件。甲方地图、视频SDK、宿主APP和业务数据的权利边界由甲方确认；乙方不复制或扩散超出项目授权范围的资料。发现许可证或权利风险时停止纳入交付并登记问题。

    ## 6.5 开源许可证识别与项目流程

    流程为需求提出、候选检索、许可证识别、SCA与漏洞检查、架构和合规复核、准入登记、版本锁定、持续扫描、交付审计。组件升级必须重新执行差异审查。交付时输出SBOM、许可证文本、NOTICE、扫描结果、例外批准和替换说明。

    # 第七章 非功能设计

    ## 7.1 性能效率设计

    性能设计采用端到端时间和分位数口径。预案启动、告警接入、任务下达、消息、定位、视频、页面、地图、打卡、视图和报表分别设置时间戳采集点。缓存只用于可接受陈旧度的读场景，写入和高风险指令以一致性与可追踪为先。峰值不少于100人时同时观察响应、错误、数据库连接、线程、队列和外部通道。
    """)
    performance = [
        ["预案启动", "通知与任务下发≤3秒", "预生成任务模板、短事务、异步并行投递", "服务端时间戳、任务和消息记录"],
        ["告警接入", "≤2秒", "轻量校验、异步后处理、来源幂等", "告警源与事件时间戳"],
        ["任务下达", "确认后≤3分钟", "审批催办、启动编排和异常告警", "审批与任务时间线"],
        ["消息", "≥20路、≥99%", "并发窗口、重试、回执和人工降级", "正常验收通道发送与回执"],
        ["定位", "≤2秒、亚米级源精度不降低", "增量刷新、楼层映射、过期标识", "源与展示坐标和时间"],
        ["视频", "首帧≤3秒", "目录缓存、按需取流、播放器超时", "请求、取流和首帧时间"],
        ["Web/地图/打卡/报表", "95%≤3秒、≤2秒、≤1秒、≤5秒", "索引、视口加载、短写事务、统计预聚合", "性能报告和原始采样"],
    ]
    lines.append(md_table(["对象", "目标", "保障设计", "验证证据"], performance)); lines.append("")
    add_block(lines, """
    ## 7.2 信息安全性设计

    方案按不低于等保二级中与应用软件直接相关的要求建设。身份优先复用中台，权限采用角色、组织、事件和数据范围组合控制；服务接口校验令牌、权限、参数和幂等；敏感配置与代码分离；传输按甲方证书和网络方案加密；日志记录主体、动作、对象、结果和关联标识并防止普通用户修改。

    开发与交付执行SCA、代码与依赖漏洞扫描、Web渗透测试、越权和注入专项测试。高危漏洞必须清零并提供复测结果。最终定级、备案、基础环境和正式第三方测评由甲方负责；本项目配合整改应用侧问题。

    ## 7.3 可靠性设计

    核心事件事务与外部调用解耦，消息和接口使用持久化台账。服务重启后从持久状态恢复待办；外部依赖按接口独立熔断，避免级联故障。后备电源由甲方提供，本项目给出关键服务清单和启动顺序。每日增量、每周全量备份，恢复演练验证事件、任务、消息和值班数据一致性。

    ## 7.4 可移植性与兼容性设计

    系统以容器化制品、离线依赖、环境检查、初始化脚本、升级和回滚脚本交付，配置和代码分离。从干净操作系统到可访问不超过2小时，并由非乙方人员一次独立部署成功。支持主流Linux或甲方认可国产操作系统，最终兼容组合在甲方环境确认后冻结。

    ## 7.5 易用性与维护性设计

    关键操作不超过3次页面跳转；错误提示指出原因和处理动作；高风险操作二次确认；状态同时用文字和视觉标识。模块边界、接口契约、结构化日志和健康检查降低维护难度。核心模块单元测试覆盖率不低于70%，需求、设计、代码、测试和验收保持双向追踪。

    ## 7.6 技术指标响应总表
    """)
    lines.append(md_table(["类别", "关键目标", "设计位置", "证据状态"], [
        ["功能", "39条FR全覆盖，验收用例通过率100%", "第四章、第九章、第十一章", "设计覆盖，待实施验收"],
        ["性能", "PE-01至PE-12全部达标", "第5.12节、第7.1节", "待目标环境实测"],
        ["可靠性", "7×24、可用率≥99.5%、MTTR≤2小时", "第5.13节、第7.3节", "待试运行与演练"],
        ["安全", "应用侧等保二级相关要求、高危0", "第7.2、7.7节", "待扫描、渗透与复测"],
        ["部署", "干净环境≤2小时、一次独立成功", "第3.7节、第7.4节", "待甲方环境验证"],
    ])); lines.append("")
    add_block(lines, """
    ## 7.7 安全合规控制措施映射
    """)
    lines.append(md_table(["控制域", "应用侧措施", "验证方法", "责任边界"], [
        ["身份鉴别", "中台统一身份、会话控制、口令策略", "认证、失效、锁定和绕过测试", "甲方提供身份能力，乙方适配"],
        ["访问控制", "角色、组织、事件、数据范围和服务账号最小权限", "越权矩阵和接口测试", "乙方应用侧"],
        ["安全审计", "操作、接口、配置、导出和高风险指令留痕≥180天", "日志完整性和查询抽样", "乙方应用，甲方提供环境"],
        ["数据保护", "敏感字段控制、导出审计、备份和传输保护", "权限、恢复和传输检查", "双方按澄清分工"],
        ["漏洞管理", "SCA、扫描、渗透、修复复测、高危0", "报告与复测证据", "乙方应用侧"],
    ])); lines.append("")

    add_block(lines, "# 第八章 项目实施方案\n\n## 8.1 项目组织")
    add_block(lines, source_section(plan, "4 三人成员职责与交付责任"))
    add_block(lines, "## 8.2 进度计划\n\n### 8.2.1 工作分解结构")
    add_block(lines, source_section(plan, "2.3 WBS 工作包"))
    add_block(lines, "[图8-1 项目工作分解结构WBS]")
    add_block(lines, "### 8.2.2 总体进度与关键路径")
    add_block(lines, source_section(plan, "3.1 阶段安排"))
    add_block(lines, source_section(plan, "3.2 任务依赖与关键路径"))
    add_block(lines, "[图8-2 总体进度甘特图]\n\n[图8-3 项目阶段与正式里程碑]")
    add_block(lines, "### 8.2.3 里程碑交付与准出")
    add_block(lines, source_section(plan, "3.4 正式里程碑与准出条件"))
    add_block(lines, "## 8.3 风险管理")
    add_block(lines, source_section(risk, "1.2 评估口径"))
    add_block(lines, source_section(risk, "1.3 风险登记与评估表"))
    add_block(lines, "[图8-4 项目风险矩阵]")
    add_block(lines, source_section(risk, "2 跟踪机制"))
    add_block(lines, "## 8.4 配置管理")
    add_block(lines, source_section(plan, "5.2 配置与文档管理"))
    add_block(lines, source_section(plan, "5.3 Git 工作安排"))
    add_block(lines, source_section(plan, "5.4 变更控制"))
    add_block(lines, """
    配置项包含需求、设计、源代码、构建脚本、容器镜像、数据库脚本、接口契约、测试资产、部署配置、用户文档和验收证据。每个里程碑执行配置审计，确认版本、校验值、追踪关系和签认状态一致。生产配置不进入公开日志，密钥由甲方受控设施管理。

    ## 8.5 沟通机制
    """)
    add_block(lines, source_section(plan, "5.1 沟通与报告"))
    lines.append(md_table(["沟通活动", "频率/触发", "参与者", "输出"], [
        ["项目例会", "每周", "甲乙双方项目负责人及相关成员", "周报、决定、问题与行动项"],
        ["接口专题会", "接口资料或联调窗口变化时", "甲方协调人、B及接口方", "契约、前置条件、问题单"],
        ["风险升级", "重大风险识别后24小时内", "A、甲方负责人、相关责任人", "书面风险报告与处置决定"],
        ["里程碑评审", "M1至M5", "甲方、A/B/C及评审人员", "评审记录、问题与准出结论"],
        ["变更评审", "范围、数字、责任或验收变化时", "双方授权人员", "变更影响与书面确认"],
    ])); lines.append("")
    add_block(lines, """
    ## 8.6 交付物响应清单

    交付包括源代码及构建文件、容器镜像与脚本、数据库初始化与升级脚本、H5包、计划、需求、设计、数据库、接口、测试、用户、部署运维、培训、开源合规、追踪矩阵、试运行、总结文档，另含基础数据初始化成果、演练和培训记录、质保承诺。完整清单见附录G。

    ## 8.7 开发方法与工程实践

    项目按里程碑控制与短周期增量结合执行。每个增量从需求条款和验收条件开始，形成设计、实现、自动测试和演示证据。高风险接口和性能原型前置，缺少甲方环境时使用模拟服务只验证适配逻辑，不能写成真实连通结论。

    ## 8.8 项目管理制度与报告模板

    项目使用任务看板、周报、风险登记册、Issue、变更记录、评审记录和配置审计。周报至少包含完成项、下周计划、偏差、风险、接口前置和需甲方决策事项。状态必须有证据，不以主观百分比替代。

    ## 8.9 试运行保障方案

    M3初验后形成试运行启动记录并立即连续监测，W17至W20保留运行、工单、缺陷和可用率证据。试运行至少支持一次真实应急演练。问题按缺陷和服务级别处理；影响验收的缺陷未关闭时不得准出M4和M5。

    # 第九章 质量保障方案

    ## 9.1 质量目标

    质量目标包括全部39条FR有测试用例、验收用例通过率100%、核心模块单元测试覆盖率不低于70%、高危漏洞为0、性能指标全部达标、文档与系统和代码一致、非乙方人员一次独立部署成功。每项目标均指定证据和责任人。

    [图9-1 需求到验收的证据链]

    ## 9.2 测试策略

    单元测试验证状态机、规则、映射和异常分支；集成测试验证数据库、消息、文件、地图和适配器；系统测试验证完整业务、性能、兼容、安全、恢复和可用性；验收测试由甲方组织，乙方配合。测试从需求和风险派生，正常路径与失败降级同等重要。

    ## 9.3 质量保障措施

    需求基线、设计评审、代码评审、持续集成、静态检查、SCA、自动测试、手工探索、性能测试、安全测试、部署演练和文档审计构成质量门禁。每次发现记录环境、步骤、期望、实际和证据，修复后由非修复人复验。

    ## 9.4 测试用例样例
    """)
    lines.append(md_table(["用例ID", "目标", "前置", "步骤摘要", "预期与证据"], [
        ["TC-PLAN-START-001", "验证预案启动≤3秒", "已发布预案、正常消息通道", "确认事件并启动，记录各时间戳", "任务和通知入队≤3秒；事件、任务、消息记录"],
        ["TC-MSG-020", "验证≥20路与≥99%", "正常验收通道、接收终端", "并发发送并收集平台、通道和终端回执", "并行数与到达率达标；原始回执"],
        ["TC-LOC-002", "验证定位刷新和精度", "甲方亚米级源与控制点", "对比源和展示坐标、时间", "≤2秒且不降低源精度"],
        ["TC-VIDEO-003", "验证实时和历史回放", "GB/T 28181账号、≥30天录像", "实时调阅、检索边界录像、断流恢复", "首帧≤3秒，回放可用，异常有提示"],
        ["TC-DOOR-004", "验证安全联动", "授权账号、门禁联锁和回执", "确认下发、拒绝、超时、重复和失败", "服从联锁，失败告警并人工降级"],
        ["TC-DEPLOY-005", "验证独立部署", "干净环境和离线制品", "非乙方人员按手册部署", "≤2小时且一次成功"],
    ])); lines.append("")
    add_block(lines, """
    ## 9.5 缺陷分级与测试准入准出

    一级缺陷导致核心业务不可用、数据严重错误或安全高危；二级缺陷导致主要功能受阻但有临时方案；三级缺陷为局部功能或易用性问题；四级为文档和轻微显示问题。进入系统测试前需通过单元和集成门禁；进入初验前不得存在未处置一级缺陷；竣工验收前高危漏洞和阻断验收缺陷必须清零。

    ## 9.6 文档质量保障

    文档与代码同库同版本，包含封面、版本、修订、编号、来源、图表和签认状态。需求、设计、接口、测试和验收建立双向追踪。电子交付包含可编辑源文件与PDF。部署运维手册必须由非编制人员按干净环境实测。

    ## 9.7 评审检查单样例
    """)
    lines.append(md_table(["评审对象", "必查项", "不通过条件"], [
        ["需求", "39条FR、34条★、PE/NFR、责任和验收", "遗漏★、数字无来源、范围漂移"],
        ["设计", "状态、数据、接口、异常、权限和验证", "只有结论、无法验证、依赖不清"],
        ["代码/配置", "追踪、测试、日志、密钥和依赖", "未评审高风险变更、高危漏洞"],
        ["测试", "环境、数据、步骤、结果、原始证据", "只报结论、不能复现、指标口径错误"],
        ["交付", "版本、清单、校验、部署和文档一致", "缺项、版本不一致、无法独立部署"],
    ])); lines.append("")
    add_block(lines, """
    ## 9.8 测试数据与测试环境管理

    测试数据分合成数据、脱敏样例、甲方提供验收数据和运行数据。甲方数据只在授权私有环境使用，不上传公共服务。环境记录操作系统、容器、数据库、浏览器、宿主、网络、接口版本、账号权限、时间同步和数据集版本。性能、安全和验收结果必须绑定环境。

    ## 9.9 验收测试组织方案

    验收分初验和竣工验收，由甲方组织、乙方配合。顺序为文档审查、功能用例、性能、四系统端到端联动及一次模拟联合演练、安全复测、部署与可移植性。验收不合格时在15日内完成整改并申请复验。全部FR用例通过率100%，PE全部达标，高危0，第三方独立部署一次成功。

    ## 9.10 质量文化与持续改进

    每周质量数据进入周报，重大偏差进入Issue或风险。重复缺陷执行根因分析并改进检查单、测试或设计规则。人工复核负责确认AI辅助内容、数字、★条款和图表，工具结果不能替代责任人结论。

    # 第十章 培训与售后服务方案

    ## 10.1 培训方案
    """)
    lines.append(md_table(["对象", "时长", "内容", "考核与证据"], [
        ["指挥人员", "1天", "核实、启动、指挥、联动、关闭与调查", "桌面推演、签到、考核和录屏"],
        ["值班/安保", "半天", "打卡、上报、告警、态势和先期处置", "实操记录与考核"],
        ["处置人员", "半天", "任务接收、反馈、附件和完成", "H5实操"],
        ["物资管理员", "半天", "台账、盘点、差异和有效期", "盘点演练"],
        ["系统管理员", "1天", "配置、接口、日志、备份恢复和部署", "独立操作考核"],
    ])); lines.append("")
    add_block(lines, """
    培训在甲方地点实施，结合馆内预案开展桌面推演。交付课件、录屏、签到和考核记录。新用户完成完整事件上报与处置的学习目标不超过2小时。

    ## 10.2 售后服务方案

    质保期为12个月，提供缺陷修复和适应性升级。建立7×24申告渠道和7×12远程支持，重大活动按甲方计划提供保障。工单记录受理、分级、诊断、恢复、根因、修复、复测和关闭。

    ## 10.3 应急预案

    一级故障30分钟响应，4小时内恢复或提供临时方案；二级故障30分钟响应、2小时内解决，复杂问题最迟1个工作日；三级故障8小时响应、2小时内给出意见、5个工作日内解决。超时按项目升级链路报告，并保护事件、任务、消息和数据。

    ## 10.4 知识转移计划

    知识转移包括架构、业务规则、数据模型、接口、部署、监控、备份恢复、故障排查和版本发布。甲方人员在乙方指导下完成一次安装、一次备份恢复、一次接口问题定位和一次版本回滚，结果进入移交记录。

    ## 10.5 培训教材与手册目录

    教材包括角色快速入门、指挥操作、H5现场操作、资源和值班、演练、系统配置、接口监控、部署运维、备份恢复和常见故障。每份手册注明适用版本、读者、前置条件和操作证据。

    ## 10.6 服务级别承诺

    SLA以用户需求书为准，不因本方案描述而扩大或缩小。服务统计区分响应、临时恢复和最终解决时间；等待甲方环境、接口方或业务确认的时间按双方工单记录处理。质保期满提交服务总结、遗留问题和资产移交清单。
    """)

    append_compliance_chapter(lines)
    append_people_achievements_and_appendices(lines)


def bid_location(clause_id: str) -> str:
    fr = re.match(r"FR-(\d{2})", clause_id)
    if fr:
        return f"第4.{int(fr.group(1))}节；第11章"
    prefix = clause_id.split("-")[0]
    mappings = {
        "STD": "第3.2节、第6章、第7章", "REQ": "第2章、第3章",
        "PE": "第5.12节、第7.1节", "NFR": "第7章",
        "INT": "第3.10节、第4.11节、第5.11节", "DATA": "第3.6节、第4.15节",
        "IMPL": "第8.2节", "PM": "第8.4至8.8节", "TRAIN": "第10.1节",
        "SLA": "第10.2、10.3、10.6节", "RISK": "第8.3节",
        "ACC": "第9.9节", "DEL": "第8.6节、附录G",
        "DOC": "第9.6节、附录G", "TB": "附录F",
    }
    return mappings.get(prefix, "第2至10章相关章节；第11章")


def append_response_table(lines: list[str], selected: list[dict[str, str]]) -> None:
    table_rows = []
    for row in selected:
        response = "完全响应。" + (row["vendor"] or "按本技术标对应章节落实并提交证据")
        evidence = row["evidence"] or row["verification"]
        deviation = "无偏离" if row["deviation_status"] in {"FULLY_COMPLIANT", "PENDING_EVIDENCE"} else row["deviation_status"]
        table_rows.append([
            row["matrix_id"] + "／" + row["clause_id"], row["requirement"], response,
            bid_location(row["clause_id"]),
            evidence + "；成熟度：" + row["maturity"] + "；证据状态：" + row["deviation_status"],
            deviation,
        ])
    lines.append(md_table(["矩阵/条款", "要求", "响应", "方案位置", "验证证据", "偏离结论"], table_rows))
    lines.append("")


def append_compliance_chapter(lines: list[str]) -> None:
    rows = parse_compliance_rows()
    add_block(lines, """
    # 第十一章 技术规格响应与偏离表

    本章以control/compliance_matrix.md的168行受控记录为唯一来源。表中设计响应说明本方案如何承接要求；验证证据是实施后应形成的证据，不表示现场已经完成。当前矩阵18行为FULLY_COMPLIANT，150行为PENDING_EVIDENCE；正偏离0、负偏离0、不适用0。

    ## 11.1 实质性条款响应表
    """)
    append_response_table(lines, [r for r in rows if r["star"] == "是"])
    add_block(lines, "## 11.2 功能需求响应表")
    append_response_table(lines, [r for r in rows if r["clause_id"].startswith("FR-")])
    add_block(lines, "## 11.3 性能与非功能需求响应表")
    append_response_table(lines, [r for r in rows if r["clause_id"].startswith(("PE-", "NFR-"))])
    add_block(lines, "## 11.4 标准、架构、数据与接口响应表")
    append_response_table(lines, [r for r in rows if r["clause_id"].startswith(("STD-", "REQ-", "DATA-", "INT-"))])
    add_block(lines, "## 11.5 实施、服务、部署、验收与交付响应表")
    append_response_table(lines, [r for r in rows if r["clause_id"].startswith(("IMPL-", "PM-", "TRAIN-", "SLA-", "RISK-", "ACC-", "DEL-", "DOC-"))])
    add_block(lines, "## 11.6 课程门禁与评分响应表")
    append_response_table(lines, [r for r in rows if r["clause_id"].startswith("TB-")])
    add_block(lines, """
    ## 11.7 偏离结论与证据状态

    全部条款当前为无偏离响应，未提出优于需求的无依据承诺。PENDING_EVIDENCE表示设计或管理安排已经建立，但接口、环境、测试、试运行、签认或验收证据需要在后续里程碑形成。任何★条款出现负偏离、关键数字无来源或高影响问题未关闭时，项目不得冻结或提交竣工验收。
    """)


def append_people_achievements_and_appendices(lines: list[str]) -> None:
    add_block(lines, """
    # 第十二章 项目团队

    ## 12.1 岗位配置与职责
    """)
    lines.append(md_table(["岗位", "主要职责", "阶段投入", "人员信息"], [
        ["项目经理/技术标总编", "范围、计划、风险、沟通、交付与跨文档一致性", "全周期", "【待人工确认】"],
        ["技术负责人/架构师", "架构、数据、接口、NFR、部署和技术验证", "M1至M5", "【待人工确认】"],
        ["需求与合规负责人", "需求、RTM、★、标准、偏离和验收追踪", "M1至M5", "【待人工确认】"],
        ["Web/H5工程师", "管理端、指挥端、移动端和宿主适配", "设计至试运行", "【待人工确认】"],
        ["后端/数据工程师", "业务服务、状态机、数据、任务和消息", "设计至试运行", "【待人工确认】"],
        ["接口/GIS工程师", "地图、定位、视频及外部系统适配", "M2至M4", "【待人工确认】"],
        ["测试/安全工程师", "四级测试、性能、安全、兼容和恢复", "设计至M5", "【待人工确认】"],
        ["实施与运维工程师", "部署、迁移、培训、试运行和质保", "M3至质保期", "【待人工确认】"],
    ])); lines.append("")
    add_block(lines, """
    ## 12.2 拟投入关键人员简历

    ### 12.2.1 项目经理

    姓名、学历、资格、相关经验、在本项目职责和可投入时间均为【待人工确认】。提交前须附可核验材料并取得本人授权。

    ### 12.2.2 技术负责人/架构师

    姓名、架构与集成经验、证书、项目证明和投入承诺均为【待人工确认】。

    ### 12.2.3 开发与接口人员

    Web、H5、后端、数据、GIS和接口人员姓名、技能证据及投入均为【待人工确认】。

    ### 12.2.4 测试、安全、实施与运维人员

    人员姓名、测试或安全能力证明、部署与培训经验均为【待人工确认】。

    ## 12.3 人员投入计划

    具体人数、人月和到岗日期没有真实依据，统一标记【待人工确认】。项目计划仅冻结A/B/C在课程第一关的职责，不应被解释为合同实施团队规模。

    # 第十三章 类似项目业绩

    ## 13.1 业绩填报说明

    当前资料未提供投标人真实名称、合同、验收报告或客户证明，本章不得引用教学案例中的虚构业绩。以下字段保留给人工补充，未补充时不作业绩主张。

    ### 13.1.1 类似应急管理系统业绩

    项目名称、客户、合同时间、范围、金额、验收情况、联系人及证明材料：【待人工确认】。

    ### 13.1.2 博物馆或人员密集场所集成业绩

    项目名称、视频、门禁、消防、消息和中台集成范围及验收证明：【待人工确认】。

    ### 13.1.3 私有化部署与安全交付业绩

    项目名称、容器化、隔离环境、安全测试和运维移交证明：【待人工确认】。

    # 第十四章 合理化建议

    建议甲方在M1同步确认接口清单、业务责任人和验收数据；在M2前优先开放视频与消息真实连通窗口；统一楼栋、楼层、点位、组织和人员编码；将门禁联动授权和人工降级纳入现场规程；以真实演练检验预案、任务、消息和资源数据；在试运行期间按周复盘未送达、接口失败、缺卡、盘点差异和演练问题。

    这些建议不改变项目范围、指标、责任或里程碑。需要增加功能、硬件、第三方服务或正式测评时，双方应先进行书面变更评估。

    # 附录A 术语和缩略语
    """)
    terms = [
        ["预案快照", "事件启动时固化的预案版本、流程、任务和人员引用"],
        ["事件状态机", "控制事件合法状态迁移和历史记录的规则"],
        ["幂等键", "识别重复请求并保证同一业务动作只生效一次的标识"],
        ["回执", "平台受理、通道发送、终端到达或业务确认的状态证据"],
        ["H5", "嵌入甲方现有智慧管理APP的移动Web应用"],
        ["GIS", "地理信息系统能力，用于楼层、点位和业务图层"],
        ["GB/T 28181", "公共安全视频监控联网的标准协议要求"],
        ["REST", "基于HTTP资源和方法组织的接口风格"],
        ["OpenAPI 3.0", "用于描述和生成接口文档的规范"],
        ["RBAC", "基于角色的访问控制"],
        ["SCA", "软件成分分析，用于依赖漏洞与许可证识别"],
        ["SBOM", "软件物料清单，记录直接和间接组件"],
        ["RTO", "恢复时间目标"],
        ["RPO", "恢复点目标"],
        ["MTTR", "平均修复或恢复时间，本项目单点故障目标不超过2小时"],
        ["SLA", "服务级别约定"],
        ["RTM", "需求追踪矩阵"],
        ["PENDING_EVIDENCE", "设计已覆盖但实施、测试或验收证据尚未形成"],
    ]
    lines.append(md_table(["术语", "说明"], terms)); lines.append("")
    add_block(lines, """
    # 附录B 接口与报文设计节选

    报文仅展示语义，字段和地址在M2接口契约评审时冻结。示例中的标识为虚构测试值，不代表真实数据。

    ## B.1 申请访问令牌

    请求字段包括client_id、授权凭据和scope；响应包括access_token、token_type、expires_in和trace_id。凭据不得写入日志，失败返回统一认证错误码。

    ## B.2 提交事件与启动预案
    """)
    lines.append(md_table(["接口", "关键请求字段", "响应", "幂等与审计"], [
        ["POST /api/v1/events", "source、type、occurred_at、location、description、attachments", "event_id、status、trace_id", "client_request_id去重，保留上报来源"],
        ["POST /api/v1/events/{id}/confirm", "decision、comment、workflow_task_id", "事件状态与时间", "校验当前状态和审批权限"],
        ["POST /api/v1/events/{id}/plans/{plan}/start", "plan_version、confirm_token、client_request_id", "operation_id、task_count、message_count", "启动标识唯一，记录授权人"],
    ])); lines.append("")
    add_block(lines, """
    ## B.3 查询任务状态

    GET /api/v1/tasks/{task_id}返回任务状态、处理人、进度、最后更新时间、反馈摘要和关联事件。客户端根据状态版本避免旧响应覆盖新状态。

    ## B.4 消息回执、定位与视频请求
    """)
    lines.append(md_table(["接口", "主要字段", "状态语义"], [
        ["POST /api/v1/message-receipts", "message_id、channel_status、terminal_status、occurred_at", "回执可乱序到达，服务端按状态规则合并"],
        ["GET /api/v1/events/{id}/locations", "person_id、x、y、floor、accuracy、source_time", "返回源时间和过期标识"],
        ["POST /api/v1/video/live", "device_id、channel_id、event_id", "返回播放会话和有效期，不存储录像"],
        ["POST /api/v1/video/records/search", "channel_id、start、end", "返回既有系统录像索引"],
        ["POST /api/v1/access-doors/{id}/open", "event_id、confirm_token、client_request_id", "返回受理和联锁回执；失败转人工"],
    ])); lines.append("")
    add_block(lines, "## B.5 错误码与重试语义")
    lines.append(md_table(["错误码", "含义", "是否重试", "处理"], [
        ["AUTH-001", "令牌无效或过期", "否", "重新认证"],
        ["PERM-003", "权限不足", "否", "阻断并审计"],
        ["STATE-009", "非法状态迁移", "否", "刷新状态后按流程处理"],
        ["EXT-TIMEOUT", "外部接口超时", "按接口策略", "有限重试、查询结果或降级"],
        ["EXT-CIRCUIT", "外部接口熔断", "否", "展示故障并转人工或等待恢复"],
        ["DUP-001", "重复业务请求", "否", "返回原处理结果"],
        ["DATA-VALID", "数据校验失败", "修正后", "返回字段级错误"],
    ])); lines.append("")

    add_block(lines, "# 附录C 条款—投标方案索引\n\n## C.1 功能需求索引")
    index_rows = []
    for domain, values in parse_fr_rows().items():
        for row in values:
            index_rows.append([row["id"], row["star"], DOMAIN_DESIGNS[domain]["name"], f"第4.{domain}节；第11.2节", f"TC-F{domain:02d}用例组"])
    lines.append(md_table(["需求", "★", "功能域", "方案位置", "验证索引"], index_rows)); lines.append("")
    add_block(lines, "## C.2 非功能、接口、过程与交付索引")
    lines.append(md_table(["条款类别", "方案位置", "验证与证据"], [
        ["20项法规标准", "第3.2节、第6章、第7章、第11.4节", "适用性矩阵、评审和测试证据"],
        ["PE/NFR", "第5.12至5.13节、第7章、第11.3节", "性能、安全、恢复和兼容报告"],
        ["接口", "第3.10、3.14、4.11、5.5至5.11节", "契约、联调、异常、降级和对账"],
        ["实施与项目管理", "第8章", "计划、周报、里程碑、配置和风险"],
        ["质量与验收", "第9章、附录E", "用例、缺陷、评审、验收和签认"],
        ["培训与服务", "第10章", "课件、签到、考核、工单和总结"],
        ["交付", "第8.6节、附录G", "清单、版本、校验和签收"],
    ])); lines.append("")
    append_data_quality_appendices(lines)


def append_data_quality_appendices(lines: list[str]) -> None:
    add_block(lines, "# 附录D 核心数据字典\n\n## D.1 emergency_event（应急事件）")
    lines.append(md_table(["字段", "类型/约束", "说明"], [
        ["id", "字符串，主键", "事件标识"],
        ["source_type/source_id", "枚举/字符串", "Web、H5或外部告警来源及原标识"],
        ["event_type_id", "外键", "事件类型版本"],
        ["status/status_version", "枚举/整数", "当前状态及并发控制版本"],
        ["occurred_at/reported_at", "带时区时间", "发生与接报时间"],
        ["building/floor/location", "字符串/空间", "位置和楼层引用"],
        ["plan_snapshot_id", "外键，可空", "启动后关联的预案快照"],
        ["created_by/updated_at", "审计字段", "创建主体和更新时间"],
    ])); lines.append("")
    add_block(lines, "## D.2 plan_snapshot与emergency_task")
    lines.append(md_table(["字段", "对象", "说明"], [
        ["snapshot_id/plan_version", "预案快照", "快照标识与来源版本"],
        ["content_hash", "预案快照", "内容校验值"],
        ["task_id/event_id", "任务", "任务与事件标识"],
        ["template_id/template_snapshot", "任务", "模板引用与快照"],
        ["assignee_id/role_snapshot", "任务", "处理人和职责快照"],
        ["status/status_version", "任务", "状态和并发版本"],
        ["accepted_at/completed_at", "任务", "接收和完成时间"],
        ["feedback/attachment_ids", "任务", "反馈和附件"],
    ])); lines.append("")
    add_block(lines, "## D.3 message_delivery与interface_call")
    lines.append(md_table(["字段", "对象", "说明"], [
        ["message_id/business_key", "消息", "消息标识和业务幂等键"],
        ["recipient/channel", "消息", "接收人和通道"],
        ["status/receipt_at", "消息", "待发送至确认的状态和回执时间"],
        ["retry_count/next_retry_at", "消息", "重试次数和时间"],
        ["trace_id/interface_code", "接口调用", "关联标识和接口代码"],
        ["request_digest/response_code", "接口调用", "脱敏摘要和响应码"],
        ["duration_ms/result", "接口调用", "耗时和结果"],
    ])); lines.append("")
    add_block(lines, "## D.4 resource、checkin与audit_log")
    lines.append(md_table(["字段", "对象", "说明"], [
        ["resource_id/site_id/batch_no", "物资", "物资、站点和批次"],
        ["quantity/expiry_date", "物资", "现有数量和有效期"],
        ["checkin_id/task_id/person_id", "打卡", "记录、任务和人员"],
        ["point_id/source_location/distance", "打卡", "点位、源坐标和距离"],
        ["server_time/validation_result", "打卡", "服务端时间和校验结论"],
        ["audit_id/actor/action/target", "审计", "主体、动作和对象"],
        ["before_after/trace_id/created_at", "审计", "变更摘要、关联标识和时间"],
    ])); lines.append("")
    add_block(lines, "# 附录E 质量度量与验收用例框架\n\n## E.1 单元测试覆盖率目标分解")
    lines.append(md_table(["模块", "覆盖目标", "重点"], [
        ["事件与状态机", "核心模块≥70%", "合法迁移、并发和关闭门禁"],
        ["预案启动与任务生成", "核心模块≥70%", "快照、事务和幂等"],
        ["消息与接口适配", "核心模块≥70%", "重试、回执、熔断和对账"],
        ["权限与审计", "核心模块≥70%", "越权、数据范围和日志"],
        ["资源、值班与演练", "核心模块≥70%", "盘点、打卡、周期调度"],
    ])); lines.append("")
    add_block(lines, "## E.2 验收测试用例框架清单")
    lines.append(md_table(["用例组", "覆盖", "重点证据"], [
        ["TC-F01至TC-F11", "39条FR", "每条需求至少一条验收用例及结果"],
        ["TC-PE", "PE-01至PE-12", "原始时间戳、回执、分位数和资源数据"],
        ["TC-SEC", "应用安全", "SCA、扫描、渗透、越权、注入和高危0"],
        ["TC-INT", "四系统及其他接口", "真实契约、联调、断连恢复和联合演练"],
        ["TC-PORT", "容器化与迁移", "干净环境≤2小时且一次成功"],
        ["TC-REC", "备份恢复", "备份校验、恢复时间和数据一致性"],
    ])); lines.append("")
    add_block(lines, """
    ## E.3 质量度量月报指标

    月报记录需求覆盖、用例设计与执行、通过率、缺陷新增关闭与遗留、核心模块覆盖率、静态检查、SCA与漏洞、构建成功率、接口连通、性能基线、文档评审、配置审计、风险和变更。数据来自工具与受控记录，不能用手工修饰的结论替代。

    # 附录F 评分要点响应对照
    """)
    lines.append(md_table(["评分要点", "本方案响应", "位置与证据"], [
        ["需求理解", "角色、11域、难点、场景、优先级、12项澄清和价值度量", "第二章、粒度审计"],
        ["方案与验证", "16项总体设计、11域八要素、13项关键技术、NFR", "第三至七章、G1-04"],
        ["WBS与风险", "工作包、甘特、里程碑、16项风险和跟踪机制", "第八章、G1-06、G1-07"],
        ["文档质量", "14章和附录A至G、逐条响应、模板继承与视觉复核", "全文、三轮复核记录"],
        ["招标解析", "168行矩阵、39/39 FR、34/34 ★FR、5项其他★", "第十一章、G1-05"],
        ["AI策略", "原始Prompt、人工复核、Git证据和可解释边界", "governance与logs/prompts"],
    ])); lines.append("")
    add_block(lines, "# 附录G 交付物清单")
    lines.append(md_table(["类别", "交付物", "形式", "节点/证据"], [
        ["软件", "源代码、构建文件、容器镜像、数据库脚本、H5包", "受控仓库与离线制品", "M3至M5，版本与校验"],
        ["计划与需求", "开发计划、SRS、RTM", "可编辑源文件与PDF", "M1及持续更新"],
        ["设计", "概要、详细、数据库、接口设计", "可编辑源文件与PDF/OpenAPI", "M2"],
        ["测试", "测试计划、设计、用例、报告", "源文件、结果和原始证据", "M2至M5"],
        ["用户与运维", "用户手册、部署运维手册", "可编辑源文件与PDF", "试运行前"],
        ["合规", "开源组件清单、许可证说明、SBOM、安全报告", "清单与报告", "随版本更新，M5归档"],
        ["运行与总结", "试运行报告、开发总结、验收材料", "报告与签认", "M4至M5"],
        ["数据", "模板、映射、清洗、导入、对账和初始化成果", "数据包与记录", "双方确认"],
        ["培训与服务", "课件、录屏、签到、考核、演练记录、质保承诺", "文件与记录", "培训、试运行和验收"],
    ])); lines.append("")
    add_block(lines, "本技术标全文完。G1-08状态为REVIEW，等待B技术复核和C合规复核，不在本次自行标记DONE。")


def append_architecture_tail(lines: list[str], tech: str) -> None:
    add_block(lines, source_section(tech, "1.6 数据层"))
    add_block(lines, """
    数据按主数据、交易数据、状态历史、接口台账、附件与审计分区管理。事件、任务和消息使用稳定业务标识；外部系统标识通过映射表保存。过程状态使用追加历史避免覆盖。空间对象记录坐标系、楼栋、楼层、来源和更新时间。

    ## 3.7 部署架构
    """)
    add_block(lines, source_section(tech, "1.11 部署架构与最低配置交付"))
    add_block(lines, """
    ### 3.7.1 系统配置建议

    甲方提供服务器、网络、域名证书、备份介质和后备电源。乙方提交最低配置、端口、域名、证书、时间同步、存储增长和备份要求。最终CPU、内存和存储数值根据甲方环境清单与容量测试固化，当前标记【待人工确认】。

    [图3-3 私有化部署拓扑]

    ## 3.8 逻辑架构与模块划分
    """)
    add_block(lines, source_section(tech, "1.5 应用服务"))
    add_block(lines, "## 3.9 核心数据模型设计")
    lines.append(md_table(
        ["聚合根", "关键子对象", "一致性规则", "主要证据"],
        [
            ["预案", "版本、流程、任务模板、附件、通知对象", "发布后不原位修改，启动形成快照", "版本差异与启动快照"],
            ["事件", "上报、告警、审批、续报、评估、调查", "状态迁移受控，历史追加", "状态历史与审计"],
            ["任务", "处理人、反馈、附件、催办、消息", "任务快照与幂等下发", "任务日志与回执"],
            ["资源", "人员、小组、站点、物资批次、盘点", "数量变更有流水，过期物资不可调派", "盘点对账"],
            ["值班", "班次、点位、二维码、打卡、预警", "服务端校验身份、时段和范围", "打卡与定位证据"],
            ["接口", "契约、调用、回执、重试、对账", "关联标识贯穿请求和业务", "接口与联调报告"],
        ],
    )); lines.append("")
    add_block(lines, "## 3.10 接口设计概览")
    add_block(lines, source_section(tech, "1.10 接口层与外部系统"))
    add_block(lines, """
    ## 3.11 核心业务流程

    ### 3.11.1 预案编制与发布流程

    草稿配置完成后执行引用完整性和流程可达性校验，再经工作流审核发布。发布动作固化版本、依赖对象和附件校验值。停用只影响新实例，不改变历史事件。

    ### 3.11.2 事件接报与启动流程

    上报或告警形成待核实事件，审批通过后转已确认。授权人员选择预案并启动，本地事务写入事件状态、预案快照、任务和待发送消息；异步投递处理通道波动。

    ### 3.11.3 处置、关闭与改进流程

    处置人员反馈任务，指挥端按时间线聚合。关闭前检查强制任务和异常。关闭后按模板评估，需要时发起事故调查，问题进入改进措施并回到预案修订。

    ### 3.11.4 演练和值班流程

    周期调度生成演练或打卡任务，H5执行并回传。系统对频次、应到实到、差异和改进进行统计，确保日常准备形成可验收记录。

    ## 3.12 验证计划
    """)
    add_block(lines, source_section(tech, "3.1 验证原则"))
    add_block(lines, source_section(tech, "3.2 关键指标验证矩阵"))
    add_block(lines, """
    ## 3.13 部署拓扑与环境配置建议

    建议按接入区、应用区、数据区和运维区划分。API入口终止受控访问；应用与任务服务部署在容器网络；数据库和文件位于数据区；外部适配器按甲方网络边界访问既有系统。日志、监控和备份单独受控。网络区划最终以甲方安全方案为准。

    ## 3.14 接口设计节选

    所有接口包含版本、认证、关联标识、幂等键、超时、错误码和审计约定。写接口返回受理结果与业务状态查询地址，不能把受理成功等同于最终业务成功。详细报文见附录B。

    ## 3.15 核心状态流转设计

    事件状态为待核实、已确认、处置中、待关闭、已关闭和调查中；任务状态为待接收、已接收、处理中、待复核、已完成、已取消和失败；消息状态为待发送、平台受理、通道发送、到达、确认、失败和人工处理。状态迁移由服务端校验，非法跳转返回明确错误并记录。

    [图3-4 事件、任务与消息状态关系]

    ## 3.16 集成边界与约束条件

    甲方负责中台、地图、定位、视频、消息、安防消防、信息发布、门禁等既有能力及资料、账号、环境和窗口。乙方负责应急业务适配、字段映射、联调、容错、降级、验证与留痕。既有硬件和系统自身改造、定位基站采购、地图重测绘或三维建模、原生APP、基础硬件采购、政府专线和正式第三方等保测评不默认属于乙方范围。
    """)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["analyze", "build-audit", "build-md", "build-docx"])
    args = parser.parse_args()
    if args.command == "analyze":
        analyze()
    elif args.command == "build-audit":
        build_audit()
    elif args.command == "build-md":
        build_bid_markdown()
    elif args.command == "build-docx":
        build_docx()


if __name__ == "__main__":
    main()
