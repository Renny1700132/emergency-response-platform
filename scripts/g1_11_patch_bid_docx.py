"""Targeted G1-11 consistency repair for the published technical bid."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / ".tmp" / "g1_11_python_libs"))
from docx import Document

TARGET = ROOT / "docs" / "deliverables" / "00-投标文件技术标.docx"


def paragraphs(parent):
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from paragraphs(cell)


def replace_text(paragraph, old, new):
    if old not in paragraph.text:
        return False
    old_run = paragraph.runs[0] if paragraph.runs else None
    text = paragraph.text.replace(old, new)
    paragraph.clear()
    run = paragraph.add_run(text)
    if old_run:
        run.bold = old_run.bold
        run.italic = old_run.italic
        run.underline = old_run.underline
        run.font.name = old_run.font.name
        run.font.size = old_run.font.size
    return True


doc = Document(TARGET)
replacements = {
    "我方已按冻结基线理解本项目的范围、39条功能需求、34条带★功能条款、5项非功能或部署类★条款、20项法规标准和全部实施、服务、验收要求。本方案对上述要求均作响应，未提出正偏离或负偏离。历史录像回放、人员定位、疏散门禁、统一消息、中台、地图、H5、部署环境、基础数据、外部接口和应用安全均按12项课程模拟甲方书面澄清执行。": "我方已按冻结基线理解本项目的范围、39条功能需求、34条带★功能条款、5项非功能或部署类★条款、20项法规标准和全部实施、服务、验收要求。本方案对上述要求均作响应，未提出正偏离或负偏离。历史录像回放、人员定位、疏散门禁、统一消息、中台、地图、H5、部署环境、基础数据、外部接口和应用安全均参考12项课程模拟甲方/招标人书面澄清组织方案；原始需求书的明确强制要求始终优先。",
    "十二项澄清是本课程模拟中已经确定的甲方解释，用于解决真实需求书中的歧义并冻结当前编标口径。它们不代表现实单位签章或法律文件，也不证明接口和指标已现场验收。": "十二项澄清是本课程模拟中已经确定的甲方/招标人解释，用于解决真实需求书中的歧义并冻结当前编标口径。它们不代表现实单位签章或法律文件，也不证明接口和指标已现场验收。澄清不得覆盖原始需求书的明确强制要求；凡涉及既有系统存储、基础环境、地图生产、宿主发布或第三方测评等范围及责任分配，均为课程模拟实施边界，须在现实采购或实施前取得教师/模拟甲方的确认后方可作为实际责任划分依据。",
    "历史回放强制；既有视频系统负责录像与不少于30天保存，本项目负责调阅、回放和验证": "课程模拟澄清：既有视频系统负责录像与不少于30天保存；本项目负责调阅、回放和验证。该责任分配不覆盖原始需求书明确强制要求，现实采购或实施前待教师/模拟甲方确认。",
    "实时视频一键调阅、GB/T 28181；结合 P1 §3.3、§8.1及模拟裁决，历史录像回放按强制要求处理，既有视频系统负责存储及≥30天保存，乙方负责接口调阅、回放和验证": "本条整体为★：实时视频一键调阅、GB/T 28181；其中“历史回放”子句为“宜支持”。但原始需求书 §8.1(1) 接口要求及 FR-11.3 对录像检索/历史回放提出强制接口响应，故本方案实施调阅、回放和验证；既有视频系统存储及≥30天保存的课程模拟责任边界待教师/模拟甲方确认",
    "指标保障： 在甲方提供的正常可用验收通道内支持不少于 20 路并行，终端到达率不低于 99%（KN-006、KN-007）。": "指标保障： PE-04 消息推送不少于 20 路并行、终端到达率不低于 99%（内部台账 KN-006、KN-007）；在甲方提供的正常可用验收通道内验证。",
    "指标保障： 地图常规操作不超过 2 秒（KN-035），人员位置刷新不超过 2 秒/次且不降低甲方亚米级源精度（KN-008、KN-009），综合安防视图刷新不超过 30 秒（KN-016），应急信息视图默认不超过 60 秒（KN-017）。": "指标保障： PE-08 地图常规操作不超过 2 秒（内部台账 KN-035）；PE-05 人员位置刷新不超过 2 秒/次且不降低甲方亚米级源精度（内部台账 KN-008、KN-009）；PE-10 综合安防视图刷新不超过 30 秒、应急信息视图默认不超过 60 秒（内部台账 KN-016、KN-017）。",
    "指标保障： 报警/告警输入响应不超过 2 秒（KN-011）；视频首帧不超过 3 秒（KN-010）；竣工验收完成视频、信息发布、物联网和中台四系统端到端联动（KN-064）。": "指标保障： PE-02 报警/告警输入响应不超过 2 秒（内部台账 KN-011）；PE-06 视频首帧不超过 3 秒（内部台账 KN-010）；竣工验收完成视频、信息发布、物联网和中台四系统端到端联动（内部台账 KN-064）。",
    "指标保障： 扫码打卡写入/回传不超过 1 秒（KN-036）；位置刷新不超过 2 秒/次且不降低源精度（KN-008、KN-009）。": "指标保障： PE-09 扫码打卡写入/回传不超过 1 秒（内部台账 KN-036）；PE-05 位置刷新不超过 2 秒/次且不降低源精度（内部台账 KN-008、KN-009）。",
}

count = 0
for paragraph in paragraphs(doc):
    for old, new in replacements.items():
        if replace_text(paragraph, old, new):
            count += 1

anchor = next(p for p in doc.paragraphs if p.text.strip() == "5.12 性能与容量")
if not any("PE-01 预案启动通知/任务下发" in p.text for p in doc.paragraphs):
    intro = doc.add_paragraph("性能条款在正式响应中以原始需求书的 PE 编号为主，KN 仅用于内部控制台账。PE-01—PE-12 的逐项阈值、测试方法和证据索引如下；正文引用单项指标时同步标注对应 PE 编号。")
    table = doc.add_table(rows=1, cols=4)
    table.style = doc.tables[27].style
    for cell, value in zip(table.rows[0].cells, ["PE 编号", "对外响应指标", "验证要点", "内部台账"]):
        cell.text = value
    rows = [
        ("PE-01", "预案启动通知/任务下发 ≤3秒", "服务端启动、任务与消息时间戳", "KN-003"),
        ("PE-02", "报警/告警输入响应 ≤2秒", "告警源与事件时间戳", "KN-011"),
        ("PE-03", "事件确认至处置任务下达 ≤3分钟", "审批与任务时间线", "KN-012"),
        ("PE-04", "消息 ≥20路并行、到达率 ≥99%", "正常验收通道发送与回执", "KN-006、KN-007"),
        ("PE-05", "人员位置刷新 ≤2秒/次、源精度亚米级且不降低", "源与展示坐标、时间戳", "KN-008、KN-009"),
        ("PE-06", "视频调阅首帧 ≤3秒", "请求、取流与首帧时间", "KN-010"),
        ("PE-07", "Web 普通业务页面 95%请求 ≤3秒", "分位数、采样与资源水位", "KN-034"),
        ("PE-08", "地图常规操作 ≤2秒", "视口操作、时间戳与资源水位", "KN-035"),
        ("PE-09", "移动端扫码打卡写入/回传 ≤1秒", "客户端、服务端与回执时间戳", "KN-036"),
        ("PE-10", "综合安防视图 ≤30秒、应急信息视图默认 ≤60秒", "视图刷新与来源时间", "KN-016、KN-017"),
        ("PE-11", "应急峰值并发在线用户 ≥100人", "并发负荷、错误率和资源水位", "KN-028"),
        ("PE-12", "一年期统计报表生成 ≤5秒", "统计请求、结果与资源水位", "KN-037"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    anchor._p.addnext(intro._p)
    intro._p.addnext(table._tbl)
    count += 1

doc.save(TARGET)
print(f"saved {TARGET}; replacements={count}")
