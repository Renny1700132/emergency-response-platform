"""G5-03: independently aggregate committed evidence and clone formal Reference.

No test executor or business implementation is changed. Rendering is a separate gate.
"""
from pathlib import Path
from collections import Counter
from copy import deepcopy
import csv, hashlib, json, re, shutil, subprocess
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

ROOT = Path(__file__).resolve().parents[2]
EV = ROOT / 'evidence/g5/G5-03'
WORK = ROOT / 'docs/work/A_PM/g5_final'
OUT = ROOT / 'docs/deliverables/G5'
REF = ROOT / 'docs/reference/36-测试报告（教学样例）.docx'
for p in (EV, WORK, OUT): p.mkdir(parents=True, exist_ok=True)
def read(p): return (ROOT / p).read_text(encoding='utf-8-sig')
def js(p): return json.loads(read(p))
def save(p, x): p.write_text(json.dumps(x, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
def git(*args): return subprocess.check_output(['git', *args], cwd=ROOT, text=True, encoding='utf-8').strip()
HEAD = git('rev-parse', 'HEAD')
f = js('evidence/g5/G5-01/functional-gate-raw.json')
p = js('evidence/g5/G5-02/performance-raw.json')
q = js('evidence/g5/G5-02/quality-summary.json')
ac = f['acResults']
counts = Counter(a['result'] for a in ac)
assert counts == {'PASS':11, 'NOT_RUN':59, 'BLOCKED':47}
assert len(ac)==117 and len({a['acId'] for a in ac})==117 and len({a['frId'] for a in ac})==39
spec = read('docs/work/B_TECH/spec.md')
star = {int(m[0]) for m in re.findall(r'^### G2-FR-(\d{3})｜([^\n]+)', spec, re.M) if '｜★｜' in '｜'+m[1]}
assert len(star)==34
star_pass = sum(all(a['result']=='PASS' for a in ac if a['frNumber']==n) for n in star)
assert star_pass==1
csvrows = list(csv.DictReader(read('evidence/g5/G5-02/performance-samples.csv').splitlines()))
metrics=[]
for m in p['measurements']:
    samples=m['samples']; durations=sorted(s['durationMs'] for s in samples)
    import math
    vals={k:durations[math.ceil(len(durations)*r)-1] for k,r in [('p50Ms',.5),('p95Ms',.95),('p99Ms',.99)]}
    for k,v in vals.items(): assert round(v,3)==m[k]
    rows=[r for r in csvrows if r['metric']==m['id']]
    assert len(rows)==len(samples)
    for i,(r,s) in enumerate(zip(rows,samples)):
        assert int(r['index'])==i+1 and int(r['status'])==s['status'] and float(r['duration_ms'])==s['durationMs'] and r['trace_id']==s['traceId'] and r['path']==s['path'] and r['method']==s['method']
    fails=sum(not 200<=s['status']<300 for s in samples)
    assert fails==m['failures']
    metrics.append({'id':m['id'],'samples':len(samples),'concurrency':m['concurrency'], **vals,'failures':fails,'failureRate':fails/len(samples)})
assert len(csvrows)==1160

# Read all supplied text evidence; index physical bytes without rewriting history.
inputs = [ROOT / x for x in ['AGENTS.md','constitution.md','tasks.md','governance/g5_quality_workflow.md','governance/ai_logging.md','governance/source_priority.md','control/facts.md','control/key_numbers.md','control/issues.md','control/overrides.md','docs/work/A_PM/g5_quality_gate_plan.md','docs/work/A_PM/software_requirements_specification_v0.1.md','docs/work/B_TECH/spec.md','docs/work/C_REQ/test_plan.md','docs/work/C_REQ/rtm_v4.md','docs/work/C_REQ/rtm_v1.md','docs/work/C_REQ/rtm_g5_increment.md','docs/work/B_TECH/g5_technical_validation.md']]
inputs += list((ROOT/'evidence/g5/G5-01').rglob('*'))+list((ROOT/'evidence/g5/G5-02').rglob('*'))
inputs += list((ROOT/'logs/reviews').glob('*G5-01*'))+list((ROOT/'logs/reviews').glob('*G5-02*'))
inputs += list((ROOT/'logs/prompts').glob('2026-10-0*-*.md'))
idx=[]
for path in sorted(set(x for x in inputs if x.is_file())):
    data=path.read_bytes()
    if path.suffix in ('.md','.json','.log','.csv','.mjs'): data.decode('utf-8-sig')
    idx.append({'file':path.relative_to(ROOT).as_posix(),'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest()})
save(EV/'input-index.json', {'head':HEAD,'files':idx,'hashBasis':'physical workspace bytes; G5-02 canonical/Git-blob manifest separately validated'})

design_lines={int(m.group(1)):m.group(0) for m in re.finditer(r'^\| G2-FR-(\d{3}) /[^\n]+',read('docs/work/C_REQ/rtm_v1.md'),re.M)}
assert len(design_lines)==39
chains=[]
for a in ac:
    n=a['frNumber']; backlog=n>=30
    deferred=[] if a['result']=='PASS' else ['IMPLEMENTATION_DEFERRED' if backlog else 'EVIDENCE_DEFERRED']
    if a['result']!='PASS' and any(s in (a['expected']+a['note']) for s in ['验收通道','目标设备','定位','视频','地图','消息','安防','宿主','用户学习']): deferred.append('REAL_WORLD_DEFERRED')
    impl=[] if backlog else (['backend/src/event-workflow.mjs','backend/src/event-persistence.mjs','frontend/src/features/response/workflow.ts'] if 13<=n<=16 or n==22 else ['backend/src/sprint2-service.mjs','backend/src/sprint2-persistence.mjs','frontend/src/features/formal/module-workbench.ts'])
    for path in impl: assert (ROOT/path).is_file()
    chains.append({**a,'star':n in star,'designIds':[f'DLD-TR-{n:03}',f'DBD-TR-{n:03}',f'API-TR-{n:03}'], 'designMapping':design_lines[n], 'designSource':'docs/work/C_REQ/rtm_v1.md', 'implementationFiles':impl,'implementationStatus':'UNIMPLEMENTED_BACKLOG' if backlog else 'EXISTS_NOT_FULLY_VERIFIED', 'plannedTC':a['acId'].replace('AC-','TC-',1), 'testExecutionVersion':f['testedSource'], 'rawLocator':f"evidence/g5/G5-01/functional-gate-raw.json#acResults/{ac.index(a)}",'courseStatus':'COURSE_ACCEPTED_EVIDENCE' if a['result']=='PASS' else 'COURSE_DEFERRED', 'deferredStatuses':deferred, 'issue':'ISSUE-G5-00-001' if backlog else ('ISSUE-G5-01-002' if deferred else 'NONE'), 'followup':'实现并执行完整AC；独立复核后才能修改真实结果' if backlog else ('补齐完整AC断言与环境，重跑并交指定复核人' if deferred else '保留真实PASS及模拟身份/消息/联锁边界')})
save(EV/'rtm-v5.json', {'version':'V5.0-review','status':'PENDING_C_FINAL_REVIEW','scope':f['scope'],'results':dict(counts),'starResults':{'PASS':1,'BLOCKED':33},'identifierTraceability':'117/117 linked; missing execution is explicitly deferred, not evidence PASS','rows':chains})

def h(level,t): return ('heading',level,t)
def para(t): return ('paragraph',t)
def table(caption,headers,rows): return ('table',caption,headers,rows)
boundary='依据 OVR-034 / OVR-035，G5-01 的 A Review 与 G5-02 的 C Review 已接受课程条件验收。本报告为 G5-03 A 候选，COURSE_ACCEPTED 为建议结论，等待 C 最终独立复核；真实实现与目标环境保持 IMPLEMENTATION_AND_REAL_WORLD_DEFERRED。课程收口不表示甲方生产验收完成。'
common=[h(1,'1 范围和版本'),para('项目为某自然博物馆智能运营中心建设项目——应急管理子系统；主责何思源（A / @WhiteApricot），复核任俊强（C，待独立复核）。编制日期 2026-10-11，版本 V0.1，状态 REVIEW 候选。'),para(boundary),para(f'受控设计输入 BASELINE-G3-M3-R1.0；报告汇总 HEAD {HEAD}。功能执行来源 {f["testedSource"]}；性能来源 e6e12b59fc5749cdf0475ed9631bac7d21d4bb1a，原始时间 {p["generatedAt"]}。不同轮次分别披露，不拼接最佳数字。')]
test=common+[h(1,'2 环境和三级测试'),para('G5-01：Windows x64，Node v24.21.0，Edge 154.0.4258.62；正式页面在 Vitest/jsdom 挂载，通过真实 Node HTTP 与隔离 PostgreSQL 链路执行。身份、消息、门禁联锁仍为注入测试缝或 SIMULATED_EVIDENCE。系统包 2/2 为 C 归档实跑，A 本轮未重复清表。'),table('表 2-1 三级执行结果',['级别','真实归档结果','边界'],[['单元及护栏','13/13 PASS','领域规则与工程护栏，不是13条AC'],['模块和HTTP集成','18/18 PASS','内存/模拟依赖'],['PostgreSQL后端集成','2/2 PASS','独立数据库集成，不等于UI系统包'],['前端混合套件','45 PASS / 2 skipped','函数、组件、契约及1条内存HTTP E2E；不得统称系统测试'],['系统专用包','2/2 PASS','Web/H5→HTTP→PostgreSQL；非全部117AC或真实浏览器/宿主验收']]),h(1,'3 用例和功能统计'),table('表 3-1 完整分母和实际状态',['对象','PASS','NOT_RUN','BLOCKED'],[['117 AC',11,59,47],['39 FR（全部AC通过）',1,0,38],['34 ★FR（全部AC通过）',1,0,33],['102 ★AC',sum(a['result']=='PASS' and a['frNumber'] in star for a in ac),sum(a['result']=='NOT_RUN' and a['frNumber'] in star for a in ac),sum(a['result']=='BLOCKED' and a['frNumber'] in star for a in ac)]]),para('计划 117 TC 与 117 AC 一一编号。AC 全量通过比例为 11/117=9.40%；NOT_RUN 与 BLOCKED 不计 PASS。只有 G2-FR-013 的三条 AC 全部通过；有至少一条 PASS 的 FR 数量见 RTM v5，不以套件总绿外推全量需求。具体预期、实际、断言、用例和证据见 rtm-v5.json 117 行及正式 RTM 附录。'),h(1,'4 覆盖率和工程检查'),table('表 4-1 分轮次覆盖率',['来源','语句','分支','函数','行'],[['G5-01前端',95.57,77.73,97.36,100],['G5-01后端',87.69,77.74,85.24,87.69],['G5-02前端',95.57,77.73,97.36,100],['G5-02后端',87.75,77.77,85.36,87.75],['G5-02工程护栏',96.75,81.63,90,96.75]]),para('单位为百分比。前端按既有 coverage include、后端 backend/src/*.mjs、护栏 scripts/g4/lib/*.mjs；覆盖率不是全部业务或117AC验收覆盖。G5-02较后轮增加容器入口回归，后端21项，不将G5-01的20项后端coverage轮次改写。最终 quality 实跑见 G5-03/final-quality.log，OpenAPI 0 error / 14 warning；历史安全扫描67文件、依赖根/前端漏洞0，高危0；这不是远程CI、渗透测试或正式等保测评。'),h(1,'5 兼容性和易用性'),para(read('evidence/g5/G5-01/compatibility-matrix.md').split('## ISSUE')[0].replace('# G5-01 兼容性最小矩阵（A Review 整改版）','').strip()),para('错误语义有加载、空态、403、网络失败、traceId、上传失败保留/重试、幂等/版本冲突断言；超时人工降级为 SIMULATED_EVIDENCE。Chrome双版本、Edge第二版本、Android/iOS真实宿主及KN-039新用户学习计时均为 REAL_WORLD_DEFERRED。'),h(1,'6 缺陷和遗留'),para('本轮产品缺陷一般1项，已修复且由A既有Review接受；致命0、严重0、轻微0，限实际执行范围。未执行/阻断范围不能宣称无缺陷。技术/证据整改与008开放改进分别见《缺陷看板与修复记录》。旧87PASS/28★PASS及24PASS已作废，历史失败和首次执行环境错误全部保留。'),para('遗留逐AC清单见正式RTM附录及rtm-v5.json：030—039的30AC为IMPLEMENTATION_DEFERRED；59 NOT_RUN及其余17 BLOCKED为EVIDENCE_DEFERRED；兼容、用户和目标系统依赖叠加REAL_WORLD_DEFERRED，不互斥求和。六类真实甲方系统0/6、真实非乙方部署操作者0，长期可用率、目标容量、视频首帧/30天、真实消息到达率和定位/地图端到端均须补测。'),h(1,'7 双结论'),para(boundary),para('IMPLEMENTATION / REAL_WORLD STATUS：11 PASS / 59 NOT_RUN / 47 BLOCKED；34★FR为1 PASS /33 BLOCKED。全量真实门禁仍有阻断，不能宣布117AC或34★全部通过，不允许据A自检直接关闭第五关。'),h(1,'8 原始证据和附件'),para('G5-01 functional-gate-raw.json、ac-117-matrix.md、execution-summary.md、compatibility-matrix.md、error-semantics-review.md、defect-register.md；G5-02 quality-summary.json与quality-gate.log；完整路径和hash见G5-03/input-index.json。四份正式成果同属于本次V0.1候选；C Review请求见G5-03/review-request.md。')]

perf=common+[h(1,'2 方法 环境和数据规模'),para('原始执行 Windows 10.0.22631 x64，Node v24.14.0，AMD Ryzen 7 7735H，16逻辑核，总内存16,366,567,424 bytes。真实本地HTTP监听器+内存持久层+SIMULATED_EVIDENCE外部适配器。工具为scripts/g5/run-technical-gate.mjs，Node性能计时；20次查询预热后执行有界并发请求。'),para('规模是1160个实测样本，非1160名用户或目标数据库容量。普通查询为空事件库阶段；后续20次事件核实/启动、模拟告警/签到/位置与小型资源/统计路径依次执行。未覆盖一年统计规模、2万事件/20万任务、甲方生产PostgreSQL与网络。只归档一轮，短窗口没有稳定时段和多轮方差结论。'),h(1,'3 真实分位数和失败率'),para('nearest-rank：对所有样本耗时升序排列，取ceil(N×q)位置；不剔除失败请求。P50/P95/P99均为ms，由JSON逐样本独立复算且与CSV逐条对照，1160/1160一致。小样本的P99等于最大值是算法结果，不是由P95或max估算。'),table('表 3-1 分组实际样本',['指标','N/并发','P50','P95','P99','失败率'],[[m['id'],f"{m['samples']}/{m['concurrency']}",f"{m['p50Ms']:.3f}",f"{m['p95Ms']:.3f}",f"{m['p99Ms']:.3f}",f"{m['failureRate']:.1%}"] for m in metrics]),para('PE-07的冻结判据为95%请求≤3s，P99作为补充披露，不能替换判据。PE-11只证明本地100并发技术路径。PE-04模拟20路20/20接受，durationMs=0.065；不能外推真实到达率。PE-01/03确认至完成最大3.005ms，消息仍为模拟。PE-05输入0.5m源精度保留，不证明真实定位基础设施或持续刷新。PE-06首帧及≥30天保存、PE-10真实刷新均NOT_RUN。'),h(1,'4 资源数据和测量限制'),para('原始数据有CPU型号/逻辑核/总内存静态配置；没有进程CPU利用率、RSS、I/O、网络吞吐和连续监控时序。资源运行数据为EVIDENCE_DEFERRED，禁止补造。后续目标规模补测须同步收集资源时序、重复轮次、稳定窗口及容量，当前数据仅支持本地短时路径结论。'),h(1,'5 故障和恢复'),para('fault-drill-raw.json：无权403、MESSAGE_TIMEOUT、门禁MANUAL_DEGRADATION且automaticReplay=false；外部故障为SIMULATED_EVIDENCE。该早轮JSON内databaseRecovery=BLOCKED是当时快照，后续独立隔离库演练另有原始文件，不改写早轮。'),para('postgres-recovery-drill.json：Windows/Node24.14.0/PostgreSQL15.14隔离一次性集群；迁移up/down/up、服务stop/start、custom-format备份、销毁重建、pg_restore；37表及1条哨兵前后一致。备份48,434bytes；原始事件计时可回查。本轮不重复破坏性测试，不把短演练写成全年7×24或99.5%可用率证明。'),para('docker-clean-deploy.json及log：WSL2/Linux，Docker28.1.1/Compose2.35.1，Node24-alpine，空卷和无缓存应用构建；2026-10-09 22:50:48—22:51:05 +08:00共17秒，health/ready/重启ready=200。B操作者，非真实非乙方独立执行；清理审计强度改进008仍OPEN NON_BLOCKING，不写新脚本已正向重跑。'),h(1,'6 原始索引与版本适用性'),para('原始JSON：evidence/g5/G5-02/performance-raw.json、fault-drill-raw.json、postgres-recovery-drill.json、docker-clean-deploy.json；CSV：performance-samples.csv；logs：quality-gate.log、docker-clean-deploy.log及attempt-1/2.log、postgres-migration-attempt.log；manifest23项已按canonical LF及HEAD Git blob验证。环境、恢复与部署是各自轮次，不能合成“同一生产环境”。'),para('报告HEAD相对性能基线的server/event-workflow/sprint2-service/external-adapters/message-port无修改。后续main入口和Dockerfile修复不改变性能被测内存HTTP路径；本次仅报告与证据整理，历史性能仍代表该路径，未重跑性能。最终quality通过，详见版本影响核验JSON。'),h(1,'7 双结论'),para(boundary),para('局部本地性能均达到相应技术路径阈值，8组1160样本0失败；目标环境/容量/端到端、长期资源与真实接口继续REAL_WORLD_DEFERRED或EVIDENCE_DEFERRED。不得把本地P99写成生产SLA。')]

# Defects: separate product, evidence/review and environment events; never infer severity counts.
compat_rows=[]
for line in read('evidence/g5/G5-01/compatibility-matrix.md').split('## ISSUE')[0].splitlines():
    if line.startswith('|') and not line.startswith('|---'):
        cells=[x.strip() for x in line.strip('|').split('|')]
        if cells[0]!='对象': compat_rows.append([cells[0],cells[2],cells[4],cells[5]])
test=[table('表 5-1 实际兼容矩阵',['对象','实际环境','结果','限制'],compat_rows) if b[0]=='paragraph' and b[1].startswith('| 对象') else b for b in test]

defects=[
('DEF-G5-01-001','一般产品缺陷','H5窄屏卡片和状态裁切','390px应完整显示 / 旧截图右侧裁切','124419f','H5布局回归及Edge390/scrollWidth390；A复审005','CLOSED / VERIFIED_BY_A','相关移动FR；全部宿主验收仍延期'),
('ISSUE-G5-01-003','P1证据阻断','套件总绿被批量派生AC PASS','逐AC完整条件 / 87及24PASS被否决','124419f / 930894c / b3db2d5','5事件组合排除断言；系统2/2；A第三轮复验','CLOSED / VERIFIED_BY_A','117AC；最后缺口013-02'),
('ISSUE-G5-01-004','P1证据阻断','45项混合前端被称完整系统测试','同链路系统包 / Mock与内存E2E不能覆盖全部AC','930894c','Web/H5→HTTP→PostgreSQL2/2；A第二轮复验','CLOSED / VERIFIED_BY_A','MVP87AC系统层级'),
('ISSUE-G5-02-003','MAJOR证据阻断','manifest换行与Git字节不一致','hash逐文件可复核 / 旧8/13或12/13不一致','279976c / 6946645','canonical+Git blob23/23；C第三轮复验','CLOSED / VERIFIED_BY_C','全技术证据包'),
('ISSUE-G5-02-004','MINOR证据问题','模拟消息耗时报告0.037与原始0.065不一致','同轮次回指 / 原报告值漂移','279976c','原JSON0.065；C第二轮复验','CLOSED / VERIFIED_BY_C','PE-04 / AC-G2-FR-022-03局部'),
('ISSUE-G5-02-005','MINOR审计问题','日志完整hash错误','有效对象 / 原对象不存在','6946645','保留错误追加更正；C第三轮对象检查','CLOSED / VERIFIED_BY_C','AI审计；非功能AC N/A'),
('ISSUE-G5-02-006','MAJOR部署缺陷','相对容器入口未启动服务','服务监听 / 进程0退出但未监听','5fc359a','入口回归、Docker health/ready/重启；C第五轮','CLOSED / VERIFIED_BY_C','KN-005 / ENG部署'),
('ISSUE-G5-02-007','MAJOR部署问题','Node20镜像低于engines','Node24 / Node20 EBADENGINE','5fc359a','Node24无缓存构建和quality；C第五轮','CLOSED / VERIFIED_BY_C','KN-005 / runtime'),
('ISSUE-G5-02-008','MINOR非阻断改进','未来部署门禁写死时限且退出清理审计不足','真实时限/显式清理 / 原脚本不足','7b86b8c','静态整改；C第六轮接受边界；新Bash正向NOT_RUN','OPEN / NON_BLOCKING_IMPROVEMENT','KN-005/065；下次真实部署随验'),
('ISSUE-G5-00-001','真实全量门禁阻断','030—039缺实现','全需求验收 / 30AC无实现','N/A未实现','无复测；OVR-035课程关闭接受','CLOSED_FOR_COURSE_STAGE / IMPLEMENTATION_DEFERRED','30AC / 6★FR'),
('ISSUE-G5-01-002','真实资源/证据阻断','兼容宿主/新用户/逐AC证据不足','完整验收 / 59NOT_RUN及17非backlog BLOCKED','N/A延期','无新增实跑；A最终接受OVR-035','CLOSED_FOR_COURSE_STAGE / EVIDENCE_AND_REAL_WORLD_DEFERRED','未执行AC及KN-038/039'),
('ISSUE-G5-02-001','真实环境阻断','甲方六类系统未取得','真实联调 / 0/6实测','1ff5464（仅课程记录）','SIMULATED_EVIDENCE6/6；C最终接受OVR-034','CLOSED_FOR_COURSE_STAGE / REAL_WORLD_DEFERRED','PE-02/04/05/06/08/10'),
('ISSUE-G5-02-002','真实独立部署阻断','非乙方真实操作者缺失','独立执行一次 / 真实操作者0','1ff5464（仅课程见证）','B17秒部署+SIMULATED_COURSE_ROLE；C最终接受','CLOSED_FOR_COURSE_STAGE / REAL_WORLD_DEFERRED','KN-065；数据库/Docker局部已接受')]
defect=common+[h(1,'2 分类统计和关闭责任'),para('G5-01产品看板：一般1项已关闭，致命0、严重0、轻微0，适用已执行范围。G5-02容器入口/运行时、证据链和日志问题按原Issue定级逐项列出，不能混合为“所有技术缺陷0”。MAJOR技术部署/证据整改已由C关闭；008仍开放非阻断。范围/环境问题只关闭课程阶段，真实门禁仍阻断。'),h(1,'3 缺陷和处理记录')]
for i,d in enumerate(defects,1):
    id,sev,sym,ex,commit,retest,status,fr=d
    hashes=[]
    for short in re.findall(r'\b[0-9a-f]{7}\b',commit): hashes.append(git('rev-parse',short))
    defect += [h(2,f'3.{i} {id}'),para(f'严重度：{sev}。需求/AC：{fr}。现象：{sym}。期望/实际：{ex}。'),para(f'处理提交：{commit}；有效完整对象：'+('；'.join(hashes) or 'N/A，无产品修复提交，不补造。')),para(f'修复或延期及复测：{retest}。最终状态：{status}。原始证据回指：G5-01/defect-register.md、G5-01/functional-gate-raw.json、G5-02/manifest.json对应文件、control/issues.md该ID段及logs/reviews同任务Review；详细映射见defect-ledger.json。')]
defect += [h(1,'4 失败和弃用保留'),para('G5-01旧87PASS/28★PASS、24PASS、单事件检索断言不足、H5裁切旧图、CLI390图实际500CSS视口、CSS测试首轮1FAIL、DevTools清理EBUSY、PostgreSQL服务ECONNREFUSED和003迁移缺表均保留历史。G5-02Docker Hub超时、Node20/入口退出、初轮未取得共享数据库凭据、manifest校验失败和日志hash错误均保留。'),para('本任务文档工具首次系统Python无pypdf/fitz；bundled renderer首次因LibreOffice不在PATH及沙箱临时目录权限失败，随后使用已安装LibreOffice与授权渲染重试。不得将这些失败删掉或声称首次成功。'),h(1,'5 双结论'),para(boundary),para('课程候选没有已知未处理的严重产品缺陷；这不证明106条未通过AC没有缺陷。真实实现/资源门禁保留阻断，008后续运行复测尚未完成，C最终复核待执行。')]
save(EV/'defect-ledger.json', {'records':[dict(zip(['id','severity','symptom','expectedActual','commit','retest','finalStatus','frAc'],d)) for d in defects], 'productCountsG501':{'fatal':0,'major':0,'normal':1,'minor':0,'closedNormal':1},'openNonBlocking':['ISSUE-G5-02-008'],'realWorldBlockersPreserved':True})

rtm=common+[h(1,'2 状态判定和完整分母'),para('完整分母39FR/117AC/34★FR，34★对应102AC。PASS/NOT_RUN/BLOCKED为真实执行状态；COURSE_DEFERRED为课程处置；IMPLEMENTATION_DEFERRED、EVIDENCE_DEFERRED、REAL_WORLD_DEFERRED为后续义务，后两者可重叠，不相加冒充新分母。编号链齐全不等于执行证据齐全。'),para('设计源为G3已挂接rtm_v1.md设计段；每个FR有DLD-TR/DBD-TR/API-TR同序号。实现索引I1：backend/src/event-workflow.mjs、event-persistence.mjs及frontend/src/features/response/workflow.ts；I2：backend/src/sprint2-service.mjs、sprint2-persistence.mjs及frontend/src/features/formal/module-workbench.ts；I0：未实现backlog。标记代码存在不代表完整验收。'),para('E1：evidence/g5/G5-01/functional-gate-raw.json#acResults，逐AC定位；T1：frontend/tests/postgresql-system-e2e.test.ts专用系统包2/2；E2：evidence/g5/G5-02/原始专项包，只作局部技术支持，不提升AC PASS。J1：ISSUE-G5-00-001未实现；J2：ISSUE-G5-01-002证据/环境延期；具体条件、断言、实际和原始数组索引见rtm-v5.json。'),h(1,'3 逐需求结论')]
frrows=[]
for n in range(1,40):
    rows=[a for a in ac if a['frNumber']==n]; c=Counter(a['result'] for a in rows)
    frrows.append([f'G2-FR-{n:03}'+(' ★' if n in star else ''),f"{c['PASS']}/{c['NOT_RUN']}/{c['BLOCKED']}",'PASS' if c['PASS']==3 else 'BLOCKED','IMPLEMENTATION_DEFERRED' if n>=30 else ('NONE' if c['PASS']==3 else 'EVIDENCE_DEFERRED')])
rtm += [table('表 3-1 逐FR实际结论',['FR','P/N/B','FR结果','后续状态'],frrows),h(1,'4 逐AC链路附录'),para('下表每行连接需求编号、设计与实现、TC/测试与evidence、缺陷/延期和结论。TC为计划编号；T1仅用于具有真实PASS的AC，N/A表示缺完整执行断言，不表示不存在测试设计。详尽Given/When/Then、实际值、断言及补测解除条件在机器附件117行，不新增PASS。')]
arows=[]
for a in chains:
    n=a['frNumber']; sid=a['acId'].split('FR-')[1]
    impl='I0' if n>=30 else ('I1' if 13<=n<=16 or n==22 else 'I2')
    status=a['result']; defer='IMPL' if n>=30 else ('EVID' if status!='PASS' else 'NONE')
    arows.append([f'FR-{n:03} / AC-{sid}',f'DLD/DBD/API-{n:03}; {impl}',f'TC-{sid}; '+('T1; E1' if status=='PASS' else '执行N/A; E1缺项'),f'{status}; {defer}; '+('J1' if n>=30 else ('J2' if status!='PASS' else '无缺陷'))])
rtm += [table('表 4-1 117AC全链索引',['需求 AC','设计 实现','测试 evidence','真实结论 延期'],arows),h(1,'5 双结论和后续补测'),para(boundary),para('39FR中1个全部AC PASS，6个FR至少1个PASS（具体由机器核验计算）；34★FR为1/33。30AC未实现、59未执行及17其他阻断全部逐项保留。后续必须补功能完整AC、目标数据规模/性能、安全渗透、Chrome/Edge双版本、Android/iOS宿主、新用户计时、六真实系统和非乙方独立部署；C本轮只复核课程收口，不抹去这些义务。')]

docs=[('36-测试报告','测试报告',test,False),('37-性能压测报告','性能压测报告',perf,True),('38-缺陷看板与修复记录','缺陷看板与修复记录',defect,True),('39-RTM闭环核验记录','RTM闭环核验记录',rtm,True)]

# Physical copy first. Preserve actual cover/revision/layout styles, not approximations.
source=Document(REF)
body=source._element.body
cut=next(i for i,e in enumerate(body) if e.tag==qn('w:p') and ''.join(e.itertext()).startswith('1 测试概况'))
# itertext can duplicate lxml text; use paragraph identity instead.
cut=list(body).index(source.paragraphs[18]._p)
normal=next(x for x in source.paragraphs if x.text.startswith('覆盖率口径'))
heading=source.paragraphs[18]
heading2=next(x for x in source.paragraphs if x.text.startswith('3.1'))
caption=source.paragraphs[19]
template_table=source.tables[1]
manifest=[]
def replace_text(paragraph,text):
    runs=paragraph.runs
    if runs:
        runs[0].text=text
        for r in runs[1:]: r.text=''
    else: paragraph.add_run(text)
def addp(d,text,pattern):
    e=deepcopy(pattern._p)
    for child in list(e):
        if child.tag!=qn('w:pPr'): e.remove(child)
    from docx.text.paragraph import Paragraph
    p1=Paragraph(e,d._body); r=p1.add_run(text)
    if pattern.runs and pattern.runs[0]._r.rPr is not None: r._r.insert(0,deepcopy(pattern.runs[0]._r.rPr))
    d._element.body.insert(len(d._element.body)-1,e)
    return p1
for filename,title,blocks,fallback in docs:
    target=OUT/(filename+'.docx'); shutil.copy2(REF,target); d=Document(target)
    reps={0:f'文档编号：YJGL-2026-G5-{filename[:2]}　　版本号：V0.1',1:'密　　级：内部 · 课程交付',3:'某自然博物馆智能运营中心建设项目',4:title,5:'应急管理子系统 第五关复核候选',7:'编制单位：020202项目组',8:'编　　制：何思源（A）',9:'审　　核：任俊强（C，待复核）',10:'批　　准：待课程最终评审',11:'编制日期：2026 年 10 月 11 日'}
    for i,t in reps.items(): replace_text(d.paragraphs[i],t)
    revision=d.tables[0]
    while len(revision.rows)>2: revision._tbl.remove(revision.rows[-1]._tr)
    for cell,t in zip(revision.rows[1].cells,['V0.1','2026-10-11','全部','G5-03真实证据汇总候选；C最终复核待执行','何思源']): replace_text(cell.paragraphs[0],t)
    for e in list(d._element.body)[cut:]:
        if e.tag!=qn('w:sectPr'): d._element.body.remove(e)
    # No instructional case media participates in the new document.
    for block in blocks:
        if block[0]=='heading':
            p1=addp(d,block[2],heading if block[1]==1 else heading2); p1.paragraph_format.keep_with_next=True
        elif block[0]=='paragraph': addp(d,block[1],normal)
        else:
            _,cap,headers,rows=block
            cp=addp(d,cap,caption); cp.paragraph_format.keep_with_next=True
            t=d.add_table(rows=1,cols=len(headers))
            tp=deepcopy(template_table._tbl.tblPr); t._tbl.remove(t._tbl.tblPr); t._tbl.insert(0,tp)
            for cells,values in [(t.rows[0].cells,headers)]+[(t.add_row().cells,row) for row in rows]:
                for c,v in zip(cells,values):
                    # actual source cell and paragraph properties; business columns may differ.
                    orig=template_table.rows[0 if values is headers else 1].cells[0]
                    pr=c._tc.get_or_add_tcPr()
                    for prop in orig._tc.tcPr:
                        if prop.tag!=qn('w:tcW'): pr.append(deepcopy(prop))
                    p1=c.paragraphs[0]; p1._p.insert(0,deepcopy(orig.paragraphs[0]._p.pPr)) if orig.paragraphs[0]._p.pPr is not None else None
                    p1.paragraph_format.keep_with_next=False
                    r=p1.add_run(str(v)); r.font.color.rgb=RGBColor(0,0,0)
                    if orig.paragraphs[0].runs and orig.paragraphs[0].runs[0]._r.rPr is not None: r._r.insert(0,deepcopy(orig.paragraphs[0].runs[0]._r.rPr))
                trpr=cells[0]._tc.getparent().get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for part in d.part.package.parts:
        if part.partname.startswith('/word/header') or part.partname.startswith('/word/footer'):
            for node in part._element.iter(qn('w:t')):
                if node.text: node.text=node.text.replace('澜图','应急管理').replace('LTPT','YJGL')
    for node in d._element.iter(qn('w:color')):
        node.set(qn('w:val'),'000000')
        for key in list(node.attrib):
            if key!=qn('w:val'): del node.attrib[key]
    setting=OxmlElement('w:updateFields'); setting.set(qn('w:val'),'true'); d.settings._element.append(setting)
    d.save(target)
    md=[f'# {title}', '', '版本V0.1；主责A；C最终复核待执行。','']
    for b in blocks:
        if b[0]=='heading': md.extend(['#'*(b[1]+1)+' '+b[2],''])
        elif b[0]=='paragraph': md.extend([b[1],''])
        else:
            md += [b[1],'','| '+' | '.join(map(str,b[2]))+' |','| '+' | '.join(['---']*len(b[2]))+' |']
            md += ['| '+' | '.join(str(x).replace('|','/').replace('\n',' ') for x in row)+' |' for row in b[3]]
            md.append('')
    (WORK/(filename+'.md')).write_text('\n'.join(md),encoding='utf-8')
    expected=[b[2] if b[0]=='heading' else b[1] for b in blocks if b[0]!='table']
    verified=Document(target)
    actual='\n'.join(x.text for x in verified.paragraphs)
    actual_cells='\n'.join(c.text for t in verified.tables for r in t.rows for c in r.cells)
    assert all(x in actual for x in expected)
    for b in blocks:
        if b[0]=='table':
            for row in b[3]:
                for value in row: assert str(value) in actual_cells
    refzip=ZipFile(REF); outzip=ZipFile(target)
    preserve=['word/styles.xml','word/numbering.xml','word/theme/theme1.xml']
    preserved={name:refzip.read(name)==outzip.read(name) for name in preserve if name in refzip.namelist()}
    assert all(preserved.values())
    manifest.append({'file':target.relative_to(ROOT).as_posix(),'reference':REF.relative_to(ROOT).as_posix(),'referenceSha256':hashlib.sha256(REF.read_bytes()).hexdigest(),'physicalCopy':True,'mode':'NEW DOCUMENT','fallback':fallback,'fallbackAuthorization':'OVR-036 / PFC-G5-03-001' if fallback else None,'packagePartsPreservedBeforeFieldUpdate':preserved,'sectionGeometry':'MATCH inherited sectPr','coverAndRevision':'MATCH properties inherited; project fields replaced','body':'MATCH role pPr/rPr inherited; business chapter/row/page counts intentionally differ','tables':'source tblPr/tcPr/pPr mapped; business grids adjusted; repeat header/cantSplit added','numberingAndTheme':'MATCH unchanged before native field update','figures':'case figures intentionally removed; no new engineering figure required','toc':'MATCH inherited SDT/TOC field; must update all fields with native Word before final rendering','captionCount':sum(b[0]=='table' for b in blocks),'businessTableCount':sum(b[0]=='table' for b in blocks),'images':0,'contentConsistency':'PASS identical block source, every text/table value verified','renderStatus':'PENDING','structuralDeviation':'new project chapters replace case chapters; fallback document type differs' if fallback else 'new project chapter structure replaces teaching business content'})
save(EV/'document-manifests.json',manifest)
save(EV/'mechanical-summary.json',{'head':HEAD,'task':'G5-03','state':'PENDING_C_FINAL_REVIEW','acResults':dict(counts),'scope':f['scope'],'starAc':{'denominator':102,**dict(Counter(a['result'] for a in ac if a['frNumber'] in star))},'frAllAcPass':1,'frAtLeastOnePass':sum(any(a['result']=='PASS' for a in ac if a['frNumber']==n) for n in range(1,40)),'starResults':{'PASS':1,'BLOCKED':33},'performance':metrics,'rawSamples':1160,'failureCount':sum(m['failures'] for m in metrics),'unconditionalRealWorldPass':False,'courseDecision':'CANDIDATE_PENDING_C','realWorldSystemsTested':0,'realWorldSystemsDenominator':6,'realIndependentOperators':0})
print(json.dumps({'docs':len(docs),'AC':dict(counts),'starPass':star_pass,'rawSamples':len(csvrows),'frWithPass':sum(any(a['result']=='PASS' for a in ac if a['frNumber']==n) for n in range(1,40))},ensure_ascii=False))
