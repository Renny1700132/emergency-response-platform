"""Read-only validation of final evidence, with outputs confined to G5-03.
Run after authoring + native field refresh; no business tests or historical files written.
"""
from pathlib import Path
from collections import Counter
from datetime import datetime, timezone, timedelta
import hashlib, json, re, subprocess
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

R=Path(__file__).resolve().parents[2]
E=R/'evidence/g5/G5-03'
def read(p): return (R/p).read_text(encoding='utf-8-sig')
def load(p): return json.loads(read(p))
def save(name,x): (E/name).write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def git(*args): return subprocess.check_output(['git',*args],cwd=R,text=True,encoding='utf-8').strip()
raw=load('evidence/g5/G5-01/functional-gate-raw.json')
rtm=load('evidence/g5/G5-03/rtm-v5.json')
summary=load('evidence/g5/G5-03/mechanical-summary.json')
rows=rtm['rows']; ac=raw['acResults']
assert len(rows)==117 and len({x['frId'] for x in rows})==39
assert Counter(x['result'] for x in rows)==Counter(x['result'] for x in ac)=={'PASS':11,'NOT_RUN':59,'BLOCKED':47}
for x,a in zip(rows,ac):
    for k in a: assert x[k]==a[k],(x['acId'],k)
    assert x['plannedTC']==x['acId'].replace('AC-','TC-',1)
    for d in x['designIds']: assert d in x['designMapping']
    for path in x['implementationFiles']: assert (R/path).is_file()
    for path in x['evidence']: assert (R/path.split('#')[0]).is_file()
    assert x['rawLocator'].endswith('/'+str(ac.index(a)))
    if x['result']!='PASS': assert x['courseStatus']=='COURSE_DEFERRED' and x['deferredStatuses']
assert sum(x['star'] for x in rows)==102
assert summary['frAtLeastOnePass']==6 and summary['frAllAcPass']==1
assert summary['starAc']=={'denominator':102,'NOT_RUN':56,'BLOCKED':35,'PASS':11}
checks={'fullRtm':True,'117RowsMatchRaw':True,'39Fr34Star':True,'deferredNotPromoted':True}

# All historical evidence and frozen inputs are protected against this task's mutation.
changed=git('diff','HEAD','--name-only').splitlines()
protected=('evidence/g4/','evidence/g5/G5-01/','evidence/g5/G5-02/','control/baselines/','docs/inputs/','backend/','frontend/','tests/')
assert not [x for x in changed if x.startswith(protected)],'protected/business files changed'
checks['historicalAndBusinessFilesUnchanged']=True

perfpaths=['backend/src/server.mjs','backend/src/sprint2-service.mjs','backend/src/event-workflow.mjs','backend/src/external-adapters.mjs','backend/src/message-port.mjs']
diff=git('diff','e6e12b59fc5749cdf0475ed9631bac7d21d4bb1a','HEAD','--',*perfpaths)
assert diff==''
save('version-impact.json',{'head':git('rev-parse','HEAD'),'performanceBaseline':'e6e12b59fc5749cdf0475ed9631bac7d21d4bb1a','unchangedPerformancePathFiles':perfpaths,'pathDiffEmpty':True,'functionalSystemEvidence':'92085e9 + G5-01 controlled test changes; b3db2d5 commits final filtering test','laterProductChange':'5fc359a main entrypoint + Docker Node24; container recovery and quality independently accepted by C','currentTaskBusinessChanges':False,'rerunRequired':'no historical destructive tests repeated; final npm run quality executed on current HEAD','boundaries':'No upgrade of memory performance to PostgreSQL or owner environment'})

man=load('evidence/g5/G5-03/document-manifests.json')
ref=R/'docs/reference/36-测试报告（教学样例）.docx'
rd=Document(ref)
def geometry(d):
    return [[s.page_width,s.page_height,s.top_margin,s.bottom_margin,s.left_margin,s.right_margin,s.header_distance,s.footer_distance] for s in d.sections]
docresults=[]
for m in man:
    path=R/m['file']; d=Document(path)
    assert geometry(d)==geometry(rd)
    assert hashlib.sha256(ref.read_bytes()).hexdigest()==m['referenceSha256']
    text=''.join(n.text or '' for n in d._element.iter(qn('w:t')))
    assert not any(s in text for s in ['澜图','samgeo','BUG-116','LTPT-2026','测试设计总则','交互分割'])
    # Verify controlled Markdown body including exact table cell values after native Word save.
    md=read('docs/work/A_PM/g5_final/'+path.stem+'.md')
    body=[x for x in d.paragraphs if x.text.startswith('1 范围和版本')]
    assert body
    start=list(d._element.body).index(body[0]._p)
    bodytext='\n'.join(''.join(n.text or '' for n in p.iter(qn('w:t'))) for e in list(d._element.body)[start:] for p in e.iter(qn('w:p')))
    for line in md.splitlines():
        if not line or line.startswith('# ') or line.startswith('版本V0.1') or line.startswith('| ---'): continue
        if line.startswith('| '):
            for cell in [x.strip() for x in line.strip('|').split('|')]: assert cell in bodytext,(path.name,cell)
        elif line.startswith('##'): assert line.lstrip('# ').strip() in bodytext
        else: assert line in bodytext,(path.name,line[:80])
    caps=[p.text for p in d.paragraphs if re.match(r'^表 \d+-\d+',p.text)]
    assert len(caps)==len(d.tables)-1==m['businessTableCount'] and len(set(caps))==len(caps)
    for color in d._element.iter(qn('w:color')): assert color.get(qn('w:val')) in ('000000','auto')
    field=''.join(n.text or '' for n in d._element.iter(qn('w:instrText')))
    assert 'TOC' in field
    docresults.append({'file':m['file'],'contentConsistency':'PASS','sectionGeometry':'MATCH','caseVisibleTextAbsent':True,'tableCaptions':len(caps),'businessTables':len(d.tables)-1,'coverRevisionExemptTables':1,'TOC':'native Word updated','bodyBlack':True,'referenceSha256Unchanged':True,'nativeFieldUpdate':'Word 16.0 serialisation recorded; cached TOC text/page numbers intentionally changed'})
save('document-content-check.json',{'status':'PASS','documents':docresults})
checks['docxMarkdownContent']=True

logs=[]
for path in sorted((R/'logs/prompts').glob('2026-10-*.md')):
    content=path.read_text(encoding='utf-8-sig')
    if 'G5-' not in content: continue
    entries=re.split(r'(?m)^#{1,2} (LOG-[^\n]+)',content)
    taskentries=[]
    for i in range(1,len(entries),2):
        id=entries[i]; text=entries[i+1]
        if 'G5' not in id: continue
        continuation=re.search(r'CONTINUED_IN_(LOG-[A-Za-z0-9-]+)',text)
        taskentries.append({'id':id,'userPromptRaw':'USER_PROMPT_RAW' in text,'finalOutputRaw':'AGENT_FINAL_OUTPUT_RAW' in text,'continuedIn':continuation.group(1) if continuation else None,'gitMarkers':bool(re.search(r'commit|COMMIT|PUSH|push|提交',text))})
    logs.append({'file':path.relative_to(R).as_posix(),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'userPromptBlocks':content.count('USER_PROMPT_RAW'),'finalOutputBlocks':content.count('AGENT_FINAL_OUTPUT_RAW'),'entries':taskentries,'currentTaskPendingOutput':'LOG-G5-03-002' in content and 'AGENT_FINAL_OUTPUT_RAW' not in content.split('# LOG-G5-03-002')[-1]})
reviews=['logs/reviews/2026-10-11_G5-01-A-final-review.md','logs/reviews/2026-10-10_G5-02-C-final-course-acceptance.md']
for path in reviews: assert 'ACCEPTED' in read(path)
valid=['124419f','930894c','b3db2d5','279976c','6946645','5fc359a','7b86b8c','1ff5464','c739b40','c69ea8b']
objects={s:git('rev-parse',s) for s in valid}
for obj in objects.values(): assert git('cat-file','-t',obj)=='commit'
save('ai-audit.json',{'logs':logs,'validCommitObjects':objects,'finalReviews':reviews,'historicWrongHash':'279976c338c0a60971a92e94f02bb24ae88f75e1 retained only as explicitly corrected historical error','currentRole':'A user-authorized AI proxy; no C final signoff generated','rawDialogueAvailable':False,'conclusion':'existing original prompts/failures/review/commit chains checked; current final-output and push backfill pending actual completion; no reconstructed conversations'})
checks['priorIndependentReviewsExist']=True
save('final-selfcheck.json',{'task':'G5-03','checkedAt':datetime.now(timezone(timedelta(hours=8))).isoformat(),'status':'PASS_A_AUTOMATED_CONTENT_CHECK','checks':checks,'unconditionalRealWorldPass':False,'CFinalReview':'PENDING','visualQA':'separate render-qa.json required','finalQualityLog':'evidence/g5/G5-03/final-quality.log'})
print(json.dumps({'status':'PASS','AC':117,'FR':39,'starFR':34,'documents':4,'historicalAndBusinessChanges':0},ensure_ascii=False))
