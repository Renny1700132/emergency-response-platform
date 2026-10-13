"""A closeout verification and final archive view; no C source is mutated."""
from pathlib import Path
import subprocess,json,hashlib,re,zipfile,argparse
import sys
sys.dont_write_bytecode=True
from g6_04_archive import verify_archive
from collections import Counter
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'evidence/g6/G6-04'
SNAP='3b47e8f97caac1fa6dc989fd1ff5279bd1eb5450'
def git(*args): return subprocess.check_output(['git','-c','core.quotepath=false',*args],cwd=ROOT,timeout=30)
snapshot_files=git('ls-tree','-r','--name-only',SNAP).decode().splitlines()
cached_paths=[p for p in snapshot_files if p.startswith(('logs/reviews/','logs/prompts/')) and p.endswith('.md') or p in ['tasks.md','evidence/g6/G6-02/metrics-ledger.md','evidence/g6/G6-02/gaps-deferrals.md','evidence/g6/G6-02/delivery-index.json','evidence/g6/G6-02/rtm-acceptance-increment.json','evidence/g5/G5-01/functional-gate-raw.json','docs/work/G6/G6-01_A_rereview.md','docs/work/G6/G6-02_A_rereview.md']]
batch_input=''.join(f'{SNAP}:{p}\n' for p in cached_paths).encode('utf-8')
batch_output=subprocess.check_output(['git','cat-file','--batch'],input=batch_input,cwd=ROOT,timeout=30)
blob_cache={}
offset=0
for path in cached_paths:
    eol=batch_output.index(b'\n',offset)
    size=int(batch_output[offset:eol].split()[-1])
    blob_cache[path]=batch_output[eol+1:eol+1+size]
    offset=eol+size+2
original_git=git
def git(*args):
    if len(args)==2 and args[0]=='show' and args[1].startswith(SNAP+':'):
        path=args[1].split(':',1)[1]
        if path in blob_cache: return blob_cache[path]
    return original_git(*args)
def read(p): return json.loads((ROOT/p).read_text(encoding='utf-8-sig'))
def sha(b): return hashlib.sha256(b).hexdigest()
parser=argparse.ArgumentParser(description='Read-only G6-04 audit; index generation is a separate command')
parser.add_argument('--verify', action='store_true')
parser.add_argument('--working-tree', action='store_true', help='Draft only, not proof of a submitted snapshot')
parser.add_argument('--commit')
parser.add_argument('--index', default='evidence/g6/G6-04/final-archive-index.json')
args=parser.parse_args()
checks=[]
def check(name,condition,details=None):
    checks.append({'check':name,'pass':bool(condition),'details':details})
def same_content(actual,blob):
    return actual==blob or (b'\0' not in actual and actual.replace(b'\r\n',b'\n')==blob.replace(b'\r\n',b'\n'))

rtm=read('evidence/g6/G6-02/rtm-acceptance-increment.json')['rows']
rtm5={x['acId']:x for x in read('evidence/g5/G5-03/rtm-v5.json')['rows']}
check('Frozen denominators',len(rtm)==117 and len({x['frId'] for x in rtm})==39 and len({x['frId'] for x in rtm if x['star']})==34)
check('Actual results',Counter(x['actualResult'] for x in rtm)=={'PASS':11,'NOT_RUN':59,'BLOCKED':47})
for row in rtm: check(row['acId']+' history preserved',row['actualResult']==rtm5[row['acId']]['result'])
for path in ['evidence/g6/G6-02/gaps-deferrals.md','evidence/g6/G6-02/rtm-acceptance-increment.json','evidence/g5/G5-01/functional-gate-raw.json','logs/reviews/2026-10-12_G6-03-C-rereview.md','docs/work/G6/G6-01_A_rereview.md','docs/work/G6/G6-02_A_rereview.md']:
    check('Authoritative evidence unchanged '+path,same_content((ROOT/path).read_bytes(),git('show',f'{SNAP}:{path}')))
historic_delivery=json.loads(git('show',f'{SNAP}:evidence/g6/G6-02/delivery-index.json'))
current_delivery=read('evidence/g6/G6-02/delivery-index.json')
check('G6-02 unique index historical snapshots preserved',
      current_delivery['items']==historic_delivery['items'] and
      current_delivery['currentIncrement']==historic_delivery['currentIncrement'] and
      current_delivery['subsequentIncrements'][:2]==historic_delivery['subsequentIncrements'] and
      current_delivery['currentAcceptance']==historic_delivery['currentAcceptance'])
g604_increment=current_delivery['subsequentIncrements'][2] if len(current_delivery['subsequentIncrements'])==3 else {}
check('G6-04 C increment IDs and counts',
      [x.get('id') for x in g604_increment.get('items',[])]==[f'GI-{n:03d}' for n in range(51,66)] and
      g604_increment.get('counts',{}).get('items')==15 and
      g604_increment.get('counts',{}).get('realWorldAccepted')==0)
for item in g604_increment.get('items',[]):
    blob=git('show',f"{g604_increment['inputCommit']}:{item['file']}")
    check('G6-04 C increment canonical Git blob '+item['id'],len(blob)==item['bytes'] and hashlib.sha256(blob).hexdigest()==item['sha256'])
historic_metrics=git('show',f'{SNAP}:evidence/g6/G6-02/metrics-ledger.md').replace(b'\r\n',b'\n')
current_metrics=(ROOT/'evidence/g6/G6-02/metrics-ledger.md').read_bytes().replace(b'\r\n',b'\n')
allowed_metric_prefixes=(b'| G6-04 C\xe7\xac\xa6\xe5\x90\x88\xe6\x80\xa7\xe5\xa4\x8d\xe9\xaa\x8c\xe5\xa2\x9e\xe9\x87\x8f |',b'| 6da71f3\xe6\x97\xb6\xe7\x82\xb9\xe9\xa6\x96\xe7\x89\x88\xe8\xb7\xaf\xe5\xbe\x84\xe5\xad\x98\xe5\x9c\xa8 |')
projected_metrics=b'\n'.join(line for line in current_metrics.split(b'\n') if not line.startswith(allowed_metric_prefixes))
check('G6-02 metrics history preserved with two dated additions',projected_metrics==historic_metrics)
tasks=(ROOT/'tasks.md').read_text(encoding='utf-8-sig')
snapshot=git('show',f'{SNAP}:tasks.md').decode('utf-8-sig')
for task in ['G6-01','G6-02','G6-03']:
    old=next(x for x in snapshot.splitlines() if x.startswith('| '+task+' |'))
    new=next(x for x in tasks.splitlines() if x.startswith('| '+task+' |'))
    check(task+' task conclusion unchanged',old==new and '| DONE' in new)
check('G6-04 no premature DONE','| DONE' not in next(x for x in tasks.splitlines() if x.startswith('| G6-04 |')))
debts=read('evidence/g6/G6-04/technical-debt.json')
check('Debt IDs unique',len(debts)==12 and len({x['id'] for x in debts})==12)
for d in debts:
    check(d['id']+' transfer fields',all(d.get(k) for k in ['problem','source','impact','responsibility','priority','trigger','repaymentPlan','verification','plannedDate','realWorldStatus']))
transfer=(ROOT/'docs/work/G6/technical_debt_transfer.md').read_text(encoding='utf-8')
check('All C gaps covered',all(f'GAP{i:03}' in transfer for i in range(1,17)))
check('Real deferred preserved',all(x in transfer for x in ['REAL_WORLD_DEFERRED','KN-065','030—039','59 NOT_RUN','17个','试运行','签署','培训']))
report=(ROOT/'docs/work/G6/G6-04_project_development_summary.md').read_text(encoding='utf-8')
check('Required six report chapters',all('## '+x in report for x in ['1 项目概述','2 项目实施过程回顾','3 计划与实际偏差分析','4 AI辅助开发分析','5 技术债务','6 项目经验与流程改进']))
for d in read('evidence/g6/G6-04/deviations.json'):
    check(d['id']+' evidence/cause/impact/action',all(d.get(k) for k in ['plan','actual','evidence','cause','impact','improvement']))
census=read('evidence/g6/G6-04/ai-log-census.json')
check('Review corpus complete including Chinese paths',set(x['path'] for x in census['reviews'])==set(p for p in snapshot_files if p.startswith('logs/reviews/') and p.endswith('.md')))
check('Prompt corpus complete including Chinese paths',set(x['path'] for x in census['prompts'])==set(p for p in snapshot_files if p.startswith('logs/prompts/') and p.endswith('.md')))
check('Prompt block count matches details',census['counts']['archivedRawPromptBlocks']==sum(len(x['rawPromptBlocks']) for x in census['prompts']))
for row in census['prompts']+census['reviews']:
    check('AI source fingerprint '+row['path'],sha(git('show',f"{SNAP}:{row['path']}"))==row['sha256GitBlob'])
check('AI unknowns remain unknown',all(x+' | 未知' in report for x in ['AI代码占比','评审打回率','Prompt复用次数','返工成本']))
check('Reference source unchanged',sha((ROOT/'docs/reference/46-项目开发总结报告（复盘·教学样例）.docx').read_bytes())==read('evidence/g6/G6-04/document-pair-manifest.json')['referenceSha256'])
check('Final rendered OOXML inheritance',all(read('evidence/g6/G6-04/final-ooxml-inheritance.json')['canonicalXmlEquality'].values()))
check('Final visual QA',read('evidence/g6/G6-04/document-pair-manifest.json')['visualReview']=='PASS')
doc=Document(ROOT/'docs/deliverables/G6/46-项目开发总结报告.docx')
body_paras=doc.paragraphs
start=next(i for i,p in enumerate(body_paras) if p.text=='1 项目概述')
bodytext=[p.text for p in body_paras[start:] if p.text]
bodytable=[c.text for t in doc.tables[1:] for r in t.rows for c in r.cells]
source_lines=report.splitlines()
# Verify every source paragraph/heading and table cell, excluding Markdown
# delimiters/captions. Fields and front matter have their own audited contract.
source_text=[]
for line in source_lines[next(i for i,l in enumerate(source_lines) if l.startswith('## ')):]:
    if not line.strip() or line.startswith('| ---'): continue
    if line.startswith('|'):
        source_text.extend(c.strip() for c in line.strip().strip('|').split('|'))
    else: source_text.append(re.sub(r'^#{2,3}\s+','',line).replace('`',''))
available=Counter(bodytext+bodytable)
missing=[x for x in source_text if available[x]==0]
check('Work to DOCX content complete',not missing,missing)
check('Business captions unique',sum(p.text.startswith('表 1-1') for p in doc.paragraphs)==1 and sum(p.text.startswith('表 4-1') for p in doc.paragraphs)==1 and len(doc.tables)==3)
with zipfile.ZipFile(ROOT/'docs/deliverables/G6/46-项目开发总结报告.docx') as z:
    xml='\n'.join(z.read(x).decode('utf-8') for x in z.namelist() if x.endswith('.xml') and x.startswith('word/') and ('document' in x or 'header' in x or 'footer' in x))
check('No teaching case residue',not any(x in xml for x in ['澜图','segment-geospatial','TensorRT','GeoPackage','通关实训第 3 组']))
check('Current branch master',git('branch','--show-current').decode().strip()=='master')

# Independent verification of the already-existing archive. No file writes.
plan=(ROOT/'docs/work/A_PM/g6_delivery_acceptance_plan.md').read_text(encoding='utf-8-sig')
for task in ['G6-00','G6-01','G6-02','G6-03','G6-04']:
    taskrow=next(line for line in tasks.splitlines() if line.startswith('| '+task+' |'))
    planrow=next(line for line in plan.splitlines() if line.startswith('| '+task+' |'))
    check(task+' plan/task row exact consistency',taskrow==planrow)
    check(task+' unique controlled plan task row',sum(line.startswith('| '+task+' |') for line in plan.splitlines())==1)
check('Plan current unique C index entry','evidence/g6/G6-02/delivery-index.json' in plan and 'evidence/g6/G6-02/delivery-index.md' in plan)
check('Plan old unavailable claims removed','尚未生成' not in plan and 'G6-01—04均TODO' not in plan and '规划候选，未执行交付制作' not in plan)
for path in ['docs/work/G6/G6-01_A_rereview.md','docs/work/G6/G6-02_A_rereview.md','logs/reviews/2026-10-12_G6-03-C-rereview.md','logs/reviews/2026-10-13_G6-04-C-review.md']:
    check('Plan actual Review locator '+path,path in plan and (ROOT/path).is_file())
archive=verify_archive(index_file=args.index,commit=args.commit,working_tree=args.working_tree)
checks.extend(archive['checks'])
failures=[x for x in checks if not x['pass']]
result={'task':'G6-04','executor':'G6-04 mechanical audit; C review recorded separately; not B/A approval','factsBaselineMaster':SNAP,'status':'PASS' if not failures else 'FAIL','assertionCount':len(checks),'failures':failures,'checks':checks,'releaseDecision':'NOT_FROZEN_C_ACCEPTED_AND_INDEXED_PENDING_B_TECHNICAL_CONFIRMATION','realWorldStatus':'REAL_WORLD_DEFERRED','archiveRows':archive['archiveRows'],'archiveVerification':{k:v for k,v in archive.items() if k not in ['checks','failures']},'filesWritten':0}
print(json.dumps(result,ensure_ascii=False,indent=2))
raise SystemExit(bool(failures))
