from __future__ import annotations

import hashlib
import json
import re
import shutil
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "deliverables" / "G6"
WORK = ROOT / "docs" / "work" / "B_TECH"
EVIDENCE = ROOT / "evidence" / "g6" / "G6-01"
OUT.mkdir(parents=True, exist_ok=True)

PAIRS = [
    {
        "source": WORK / "g6_user_manual.md",
        "reference": ROOT / "docs" / "reference" / "34-用户手册（教学样例）.docx",
        "target": OUT / "34-用户手册.docx",
        "number": "YJGL-G6-34",
        "title": "用户手册",
        "body_start": 19,
        "cover": {3: "某自然博物馆智能运营中心建设项目", 4: "用户手册", 5: "应急管理子系统"},
    },
    {
        "source": WORK / "g6_deployment_operations_manual.md",
        "reference": ROOT / "docs" / "reference" / "45-部署与运维手册（教学样例）.docx",
        "target": OUT / "45-部署与运维手册.docx",
        "number": "YJGL-G6-45",
        "title": "部署与运维手册",
        "body_start": 18,
        "cover": {3: "某自然博物馆智能运营中心建设项目", 4: "部署与运维手册", 5: "应急管理子系统"},
    },
]


def set_text(paragraph: Paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def add_like(doc: Document, prototype: Paragraph, text: str) -> Paragraph:
    element = deepcopy(prototype._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    paragraph = Paragraph(element, doc._body)
    run = paragraph.add_run(text)
    if prototype.runs and prototype.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(prototype.runs[0]._r.rPr))
    doc._element.body.insert(len(doc._element.body) - 1, element)
    return paragraph


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], caption_proto: Paragraph, table_proto: Table) -> None:
    cap = add_like(doc, caption_proto, caption)
    cap.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=len(headers))
    old_pr = table._tbl.tblPr
    table._tbl.remove(old_pr)
    table._tbl.insert(0, deepcopy(table_proto._tbl.tblPr))
    ref_header = table_proto.rows[0].cells[0]
    ref_body = table_proto.rows[min(1, len(table_proto.rows) - 1)].cells[0]
    for row_index, values in enumerate([headers, *rows]):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for cell, value in zip(cells, values):
            ref = ref_header if row_index == 0 else ref_body
            tc_pr = cell._tc.get_or_add_tcPr()
            for prop in ref._tc.tcPr:
                if prop.tag != qn("w:tcW"):
                    tc_pr.append(deepcopy(prop))
            paragraph = cell.paragraphs[0]
            if ref.paragraphs[0]._p.pPr is not None:
                paragraph._p.insert(0, deepcopy(ref.paragraphs[0]._p.pPr))
            run = paragraph.add_run(value)
            if ref.paragraphs[0].runs and ref.paragraphs[0].runs[0]._r.rPr is not None:
                run._r.insert(0, deepcopy(ref.paragraphs[0].runs[0]._r.rPr))
            run.font.color.rgb = RGBColor(0, 0, 0)
        tr_pr = cells[0]._tc.getparent().get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))


def parse_markdown(path: Path) -> list[dict]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[dict] = []
    start = next(i for i, line in enumerate(lines) if line.startswith("## "))
    i = start
    pending_caption: str | None = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            blocks.append({"type": "code", "lines": code})
            i += 1
            continue
        if line.startswith("### "):
            blocks.append({"type": "heading", "level": 2, "text": line[4:]})
            i += 1
            continue
        if line.startswith("## "):
            blocks.append({"type": "heading", "level": 1, "text": line[3:]})
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                blocks.append({"type": "image", "alt": match.group(1), "path": (path.parent / match.group(2)).resolve(), "caption": pending_caption})
                pending_caption = None
            i += 1
            continue
        if line.startswith("| ") and i + 1 < len(lines) and lines[i + 1].startswith("| ---"):
            def split(row: str) -> list[str]:
                return [cell.strip().replace("`", "") for cell in row.strip().strip("|").split("|")]
            headers = split(line)
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(split(lines[i]))
                i += 1
            blocks.append({"type": "table", "caption": pending_caption or "表 受控信息", "headers": headers, "rows": rows})
            pending_caption = None
            continue
        if re.match(r"^[图表]\s*\d+-\d+", line):
            pending_caption = line
            i += 1
            continue
        if line.startswith("- ") or re.match(r"^\d+\. ", line):
            text = re.sub(r"^(?:- |\d+\. )", "", line)
            blocks.append({"type": "list", "text": text})
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{2,3} |```|\| |!\[|- |\d+\. |[图表]\s*\d+-\d+)", lines[i]):
            paragraph_lines.append(lines[i].strip())
            i += 1
        blocks.append({"type": "paragraph", "text": " ".join(paragraph_lines)})
    return blocks


def build(pair: dict) -> dict:
    shutil.copy2(pair["reference"], pair["target"])
    doc = Document(pair["target"])
    original = Document(pair["reference"])
    prototypes = {
        "normal": next(p for p in original.paragraphs if p.style.name == "Normal" and len(p.text) > 20),
        "heading1": next(p for p in original.paragraphs if p.style.name == "Heading 1"),
        "heading2": next(p for p in original.paragraphs if p.style.name == "Heading 2"),
        "list": next(p for p in original.paragraphs if p.style.name == "List Paragraph"),
        "caption": next(p for p in original.paragraphs if re.match(r"^[图表]\s*\d+-\d+", p.text.strip())),
        "code": next(p for p in original.paragraphs if p.style.name == "Normal" and (p.text.strip().startswith("#") or p.text.strip().startswith("docker") or p.text.strip().startswith("lantu"))),
    }
    first_table = original.tables[1]
    replacements = {
        0: f"文档编号：{pair['number']}　　版本号：V0.2",
        1: "密　　级：内部 · 课程交付",
        7: "编制单位：020202项目组",
        8: "编　　制：严宇（技术负责人）",
        9: "审　　核：何思源（项目负责人，待复核）",
        10: "批　　准：课程负责人（待验收评审）",
        11: "编制日期：2026 年 10 月 12 日",
        **pair["cover"],
    }
    for index, value in replacements.items():
        set_text(doc.paragraphs[index], value)
    revision = doc.tables[0]
    while len(revision.rows) > 3:
        revision._tbl.remove(revision.rows[-1]._tr)
    while len(revision.rows) < 3:
        revision._tbl.append(deepcopy(revision.rows[-1]._tr))
    revision_values = [
        ["V0.1", "2026-10-11", "全部", "G6-01 首版候选；保留为历史审计", "严宇"],
        ["V0.2", "2026-10-12", "全部", "按 A Review 补齐业务闭环、命令证据、完整部署回退和版本一致性；待 A 复验", "严宇"],
    ]
    for row, values in zip(revision.rows[1:3], revision_values):
        for cell, value in zip(row.cells, values):
            set_text(cell.paragraphs[0], value)
    cut_element = doc.paragraphs[pair["body_start"]]._p
    body_children = list(doc._element.body)
    cut_index = body_children.index(cut_element)
    for element in body_children[cut_index:]:
        if element.tag != qn("w:sectPr"):
            doc._element.body.remove(element)

    for block in parse_markdown(pair["source"]):
        kind = block["type"]
        if kind == "heading":
            paragraph = add_like(doc, prototypes["heading1" if block["level"] == 1 else "heading2"], block["text"])
            paragraph.paragraph_format.keep_with_next = True
        elif kind == "paragraph":
            add_like(doc, prototypes["normal"], block["text"].replace("`", ""))
        elif kind == "list":
            add_like(doc, prototypes["list"], f"· {block['text'].replace('`', '')}")
        elif kind == "code":
            for line in block["lines"]:
                add_like(doc, prototypes["code"], line)
        elif kind == "table":
            add_table(doc, block["caption"], block["headers"], block["rows"], prototypes["caption"], first_table)
        elif kind == "image":
            if block.get("caption"):
                cap = add_like(doc, prototypes["caption"], block["caption"])
                cap.paragraph_format.keep_with_next = True
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = False
            paragraph.add_run().add_picture(str(block["path"]), width=Inches(6.25))

    for part in doc.part.package.parts:
        if str(part.partname).startswith("/word/header") or str(part.partname).startswith("/word/footer"):
            for node in part._element.iter(qn("w:t")):
                if node.text:
                    node.text = node.text.replace("澜图", "应急管理").replace("LTPT", "YJGL")
    for node in doc._element.iter(qn("w:color")):
        node.set(qn("w:val"), "000000")
        for key in list(node.attrib):
            if key != qn("w:val"):
                del node.attrib[key]
    # The retained front matter already carries the body-start section break.
    # New chapter headings must not acquire additional section breaks, which
    # would isolate a heading on its own page and restart page numbering.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() and paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}:
            p_pr = paragraph._p.pPr
            if p_pr is not None:
                sect_pr = p_pr.find(qn("w:sectPr"))
                if sect_pr is not None:
                    p_pr.remove(sect_pr)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings._element.append(update)
    doc.save(pair["target"])

    output = Document(pair["target"])
    visible = "\n".join(p.text for p in output.paragraphs) + "\n" + "\n".join(cell.text for table in output.tables for row in table.rows for cell in row.cells)
    assert "澜图" not in visible and "遥感影像" not in visible
    assert pair["title"] in visible and "严宇" in visible
    source_sha = hashlib.sha256(pair["reference"].read_bytes()).hexdigest()
    return {
        "mode": "CONTENT EDIT",
        "target": pair["target"].relative_to(ROOT).as_posix(),
        "reference": pair["reference"].relative_to(ROOT).as_posix(),
        "referenceSha256": source_sha,
        "contentEditSource": pair["source"].relative_to(ROOT).as_posix(),
        "releaseCommit": "75f01c3bb33369d3785acf69a3a00eb983aa188b",
        "physicalCopy": True,
        "sectionGeometry": "INHERITED_FROM_PAIR",
        "cover": "REFERENCE_COVER_CLONE_PROJECT_FIELDS_REPLACED",
        "stylesNumberingTheme": "INHERITED_FROM_PAIR",
        "bodyMapping": "PAIR_ROLE_PARAGRAPH_AND_TABLE_PROPERTIES_CLONED",
        "businessStructureDeviation": "Project-specific manual chapters replace teaching-case business chapters by design.",
        "caseContaminationCheck": "PASS",
        "renderStatus": "PENDING_V0_2_QA",
    }


def main() -> None:
    manifests = [build(pair) for pair in PAIRS]
    (EVIDENCE / "document-pair-manifest.json").write_text(json.dumps(manifests, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": "PASS", "documents": [m["target"] for m in manifests]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
