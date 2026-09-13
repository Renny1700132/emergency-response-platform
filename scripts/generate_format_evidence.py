"""Extract DOCX structure/style/table/header/image evidence for all deliverable/reference pairs."""
from __future__ import annotations

import json
import re
import subprocess
import sys
import zipfile
from copy import deepcopy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = Path(r"C:\Users\15801\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")
sys.path.insert(0, str(RUNTIME / "Lib" / "site-packages"))
from docx import Document
from docx.oxml.ns import qn

PAIRS = {
    "00": ("docs/deliverables/00-投标文件技术标.docx", "docs/reference/00-澜图遥感影像智能解译平台-投标文件技术标（教学案例）.docx"),
    "01": ("docs/deliverables/01-项目建议书.docx", "docs/reference/01-项目建议书（教学样例）.docx"),
    "02": ("docs/deliverables/02-澄清与质询记录.docx", "docs/reference/02-澄清与质询记录（教学样例）.docx"),
    "03": ("docs/deliverables/03-项目计划v1（WBS与甘特图）.docx", "docs/reference/03-项目计划v1（WBS与甘特图·教学样例）.docx"),
    "04": ("docs/deliverables/04-风险登记册v1.docx", "docs/reference/04-风险登记册v1（教学样例）.docx"),
    "05": ("docs/deliverables/05-需求确认书.docx", "docs/reference/05-需求确认书（教学样例）.docx"),
    "06": ("docs/deliverables/06-述标答辩讲稿与策略.docx", "docs/reference/06-述标答辩讲稿与策略（教学样例）.docx"),
}


def pt(value):
    return None if value is None else round(value.pt, 2)


def style_info(style):
    pf = style.paragraph_format
    font = style.font
    rpr = style.element.rPr
    fonts = rpr.rFonts if rpr is not None else None
    return {
        "style_id": style.style_id,
        "name": style.name,
        "font": {
            "name_api": font.name,
            "ascii": fonts.get(qn("w:ascii")) if fonts is not None else None,
            "hAnsi": fonts.get(qn("w:hAnsi")) if fonts is not None else None,
            "eastAsia": fonts.get(qn("w:eastAsia")) if fonts is not None else None,
            "size_pt": pt(font.size), "bold": font.bold, "italic": font.italic,
        },
        "paragraph": {
            "alignment": str(pf.alignment), "left_indent_pt": pt(pf.left_indent),
            "right_indent_pt": pt(pf.right_indent), "first_line_indent_pt": pt(pf.first_line_indent),
            "line_spacing": str(pf.line_spacing), "line_spacing_rule": str(pf.line_spacing_rule),
            "space_before_pt": pt(pf.space_before), "space_after_pt": pt(pf.space_after),
            "keep_with_next": pf.keep_with_next, "keep_together": pf.keep_together,
            "widow_control": pf.widow_control, "page_break_before": pf.page_break_before,
        },
    }


def get_fill(cell):
    shd = cell._tc.tcPr.find(qn("w:shd"))
    return None if shd is None else shd.get(qn("w:fill"))


def table_info(table, idx):
    grid = table._tbl.tblGrid
    widths = [int(x.get(qn("w:w"))) for x in grid.gridCol_lst] if grid is not None else []
    first = table.rows[0] if table.rows else None
    first_text = [c.text.replace("\n", " / ")[:80] for c in first.cells] if first else []
    fills = [get_fill(c) for c in first.cells] if first else []
    repeat = False
    if first is not None:
        trpr = first._tr.trPr
        repeat = trpr is not None and trpr.find(qn("w:tblHeader")) is not None
    tblpr = table._tbl.tblPr
    borders = tblpr.find(qn("w:tblBorders"))
    border_summary = {}
    if borders is not None:
        for child in borders:
            border_summary[child.tag.rsplit("}", 1)[-1]] = {
                "val": child.get(qn("w:val")), "sz": child.get(qn("w:sz")), "color": child.get(qn("w:color"))
            }
    margins = tblpr.find(qn("w:tblCellMar"))
    margin_summary = {}
    if margins is not None:
        for child in margins:
            margin_summary[child.tag.rsplit("}", 1)[-1]] = child.get(qn("w:w"))
    return {
        "index": idx, "rows": len(table.rows), "cols": len(table.columns), "style": table.style.name if table.style else None,
        "grid_twips": widths, "header_text": first_text, "header_fills": fills, "repeat_header": repeat,
        "borders": border_summary, "cell_margins_twips": margin_summary,
    }


def pdf_pages(pdf):
    text = subprocess.check_output(["pdfinfo", str(pdf)], text=True, encoding="utf-8", errors="replace")
    m = re.search(r"^Pages:\s+(\d+)", text, re.M)
    return int(m.group(1)) if m else None


def inspect(path: Path, pdf: Path):
    doc = Document(path)
    sections = []
    for i, sec in enumerate(doc.sections, 1):
        sp = sec._sectPr
        title_pg = sp.find(qn("w:titlePg")) is not None
        sections.append({
            "index": i, "width_pt": pt(sec.page_width), "height_pt": pt(sec.page_height), "orientation": str(sec.orientation),
            "margin_top_pt": pt(sec.top_margin), "margin_bottom_pt": pt(sec.bottom_margin),
            "margin_left_pt": pt(sec.left_margin), "margin_right_pt": pt(sec.right_margin),
            "header_distance_pt": pt(sec.header_distance), "footer_distance_pt": pt(sec.footer_distance),
            "start_type": str(sec.start_type), "different_first_page": title_pg,
            "header_linked": sec.header.is_linked_to_previous, "footer_linked": sec.footer.is_linked_to_previous,
            "header_text": " ".join(p.text for p in sec.header.paragraphs).strip(),
            "footer_text": " ".join(p.text for p in sec.footer.paragraphs).strip(),
        })
    styles = {}
    for wanted in ["Title", "Heading 1", "Heading 2", "Heading 3", "Normal", "Caption", "List Paragraph"]:
        try: styles[wanted] = style_info(doc.styles[wanted])
        except KeyError: styles[wanted] = None
    front = []
    for idx, p in enumerate(doc.paragraphs[:80]):
        text = p.text.strip()
        if text:
            front.append({"index": idx, "style": p.style.name, "text": text[:240], "page_break_before": p.paragraph_format.page_break_before})
    captions = [p.text.strip() for p in doc.paragraphs if re.match(r"^[图表]\s*\d", p.text.strip())]
    with zipfile.ZipFile(path) as zf:
        media = sorted(n for n in zf.namelist() if n.startswith("word/media/"))
        parts = sorted((n, zf.getinfo(n).file_size) for n in zf.namelist())
    return {
        "path": str(path.relative_to(ROOT)), "bytes": path.stat().st_size,
        "pages": pdf_pages(pdf), "sections": sections, "styles": styles,
        "paragraph_count": len(doc.paragraphs), "front_paragraphs": front,
        "table_count": len(doc.tables), "tables": [table_info(t, i) for i, t in enumerate(doc.tables, 1)],
        "inline_shapes": len(doc.inline_shapes), "media_parts": media, "captions": captions,
        "package_part_count": len(parts),
    }


def main():
    pdf_dir = ROOT / ".tmp/final_rework/all_pairs"
    out = {}
    for key, (cur, ref) in PAIRS.items():
        out[key] = {
            "current": inspect(ROOT / cur, pdf_dir / f"{key}-current.pdf"),
            "reference": inspect(ROOT / ref, pdf_dir / f"{key}-reference.pdf"),
        }
    target = ROOT / ".tmp/final_rework/format_evidence.json"
    target.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(target)


if __name__ == "__main__":
    main()
