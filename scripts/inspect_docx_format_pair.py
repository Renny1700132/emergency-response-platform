"""Extract deterministic OOXML format evidence for one deliverable/reference pair."""

from __future__ import annotations

import argparse
import json
import zipfile
from pathlib import Path

from lxml import etree
from docx import Document
from docx.oxml.ns import qn


def pt(value):
    return None if value is None else round(value.pt, 3)


def xml_child(parent, name):
    child = parent.find(qn(name)) if parent is not None else None
    return None if child is None else etree.tostring(child, encoding="unicode")


def style_info(doc: Document, name: str):
    try:
        style = doc.styles[name]
    except KeyError:
        return None
    font = style.font
    pf = style.paragraph_format
    rpr = style.element.rPr
    fonts = None if rpr is None else rpr.rFonts
    return {
        "name": style.name,
        "style_id": style.style_id,
        "font": {
            "ascii": None if fonts is None else fonts.get(qn("w:ascii")),
            "eastAsia": None if fonts is None else fonts.get(qn("w:eastAsia")),
            "hAnsi": None if fonts is None else fonts.get(qn("w:hAnsi")),
            "size_pt": pt(font.size), "bold": font.bold, "italic": font.italic,
            "underline": str(font.underline), "color": None if font.color.rgb is None else str(font.color.rgb),
        },
        "paragraph": {
            "alignment": str(pf.alignment), "first_line_pt": pt(pf.first_line_indent),
            "left_pt": pt(pf.left_indent), "right_pt": pt(pf.right_indent),
            "before_pt": pt(pf.space_before), "after_pt": pt(pf.space_after),
            "line_spacing": str(pf.line_spacing), "line_rule": str(pf.line_spacing_rule),
            "keep_next": pf.keep_with_next, "keep_lines": pf.keep_together,
            "widow": pf.widow_control, "page_break_before": pf.page_break_before,
        },
        "rPr_xml": xml_child(style.element, "w:rPr"),
        "pPr_xml": xml_child(style.element, "w:pPr"),
    }


def section_info(section, index: int):
    sp = section._sectPr
    pg_mar = sp.find(qn("w:pgMar"))
    return {
        "index": index, "width_pt": pt(section.page_width), "height_pt": pt(section.page_height),
        "orientation": str(section.orientation), "top_pt": pt(section.top_margin),
        "bottom_pt": pt(section.bottom_margin), "left_pt": pt(section.left_margin),
        "right_pt": pt(section.right_margin),
        "gutter_twips": None if pg_mar is None else pg_mar.get(qn("w:gutter")),
        "header_pt": pt(section.header_distance), "footer_pt": pt(section.footer_distance),
        "start_type": str(section.start_type),
        "different_first": sp.find(qn("w:titlePg")) is not None,
        "vertical_align": None if sp.find(qn("w:vAlign")) is None else sp.find(qn("w:vAlign")).get(qn("w:val")),
        "header_linked": section.header.is_linked_to_previous,
        "footer_linked": section.footer.is_linked_to_previous,
        "header": [p.text for p in section.header.paragraphs],
        "footer": [p.text for p in section.footer.paragraphs],
    }


def table_info(table, index: int):
    grid = table._tbl.tblGrid
    tbl_pr = table._tbl.tblPr
    rows = []
    for row_index, row in enumerate(table.rows[:2], 1):
        rows.append({
            "index": row_index,
            "height_pt": pt(row.height),
            "height_rule": str(row.height_rule),
            "repeat_header": row._tr.get_or_add_trPr().find(qn("w:tblHeader")) is not None,
            "cant_split": row._tr.get_or_add_trPr().find(qn("w:cantSplit")) is not None,
            "cell_widths_twips": [cell._tc.tcPr.tcW.get(qn("w:w")) for cell in row.cells],
            "cell_v_align": [str(cell.vertical_alignment) for cell in row.cells],
            "cell_shading": [None if cell._tc.tcPr.find(qn("w:shd")) is None else cell._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")) for cell in row.cells],
        })
    return {
        "index": index, "rows": len(table.rows), "cols": len(table.columns),
        "style": None if table.style is None else table.style.name,
        "grid_twips": [] if grid is None else [x.get(qn("w:w")) for x in grid.gridCol_lst],
        "header_text": [c.text for c in table.rows[0].cells] if table.rows else [],
        "tblW": xml_child(tbl_pr, "w:tblW"), "tblLayout": xml_child(tbl_pr, "w:tblLayout"),
        "tblBorders": xml_child(tbl_pr, "w:tblBorders"), "tblCellMar": xml_child(tbl_pr, "w:tblCellMar"),
        "tblLook": xml_child(tbl_pr, "w:tblLook"), "rows_sample": rows,
    }


def inspect(path: Path):
    doc = Document(path)
    with zipfile.ZipFile(path) as zf:
        settings = zf.read("word/settings.xml") if "word/settings.xml" in zf.namelist() else b""
    return {
        "path": str(path), "sections": [section_info(s, i) for i, s in enumerate(doc.sections, 1)],
        "even_odd_headers": b"evenAndOddHeaders" in settings,
        "styles": {n: style_info(doc, n) for n in ["Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Normal", "List Paragraph", "Caption", "TOC 1", "TOC 2", "TOC 3"]},
        "tables": [table_info(t, i) for i, t in enumerate(doc.tables, 1)],
        "paragraphs": len(doc.paragraphs), "table_count": len(doc.tables),
        "inline_shapes": len(doc.inline_shapes),
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("current", type=Path)
    parser.add_argument("reference", type=Path)
    parser.add_argument("--out", required=True, type=Path)
    args = parser.parse_args()
    result = {"current": inspect(args.current), "reference": inspect(args.reference)}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(args.out)


if __name__ == "__main__":
    main()
