from __future__ import annotations

import sys
from pathlib import Path

from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION_START
from pptx import Presentation
import pypdfium2 as pdfium


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "deliverables" / "03-项目计划v1（WBS与甘特图）.docx"
PPTX = ROOT / "docs" / "work" / "A_PM" / "project_plan_v1_figures.pptx"
PDF = ROOT / "artifacts" / "g1-06-v12-render" / "project-plan-v12.pdf"


def inspect() -> None:
    doc = Document(DOCX)
    print(f"DOCX paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")
    for index, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() or "w:type=\"page\"" in paragraph._p.xml:
            print(f"P {index}: {paragraph.style.name!r} pagebreak={'w:type=\"page\"' in paragraph._p.xml} sect={'sectPr' in paragraph._p.xml} {paragraph.text!r}")
    for index, table in enumerate(doc.tables):
        print(f"T {index}: rows={len(table.rows)} cols={len(table.columns)} first={table.cell(0, 0).text!r}")
    for index, section in enumerate(doc.sections):
        print(f"SECTION {index}: start={section.start_type} size={section.page_width}x{section.page_height}")
    prs = Presentation(PPTX)
    for slide_index, slide in enumerate(prs.slides, start=1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text_frame") and shape.text.strip()]
        print(f"SLIDE {slide_index}: {' | '.join(texts)}")


def replace_paragraph_text(paragraph, old: str, new: str) -> bool:
    if paragraph.text != old:
        return False
    if not paragraph.runs:
        paragraph.add_run(new)
        return True
    paragraph.runs[0].text = new
    for run in paragraph.runs[1:]:
        run.text = ""
    return True


def replace_cell_text(cell, old: str, new: str) -> bool:
    full_text = "\n".join(paragraph.text for paragraph in cell.paragraphs)
    if full_text != old:
        return False
    replace_paragraph_text(cell.paragraphs[0], cell.paragraphs[0].text, new)
    for paragraph in cell.paragraphs[1:]:
        replace_paragraph_text(paragraph, paragraph.text, "")
    return True


def revise_docx() -> None:
    doc = Document(DOCX)
    paragraph_replacements = {
        "文档编号：【待人工确认】\u3000\u3000版本号：V1.1": "文档编号：【待人工确认】\u3000\u3000版本号：V1.2",
        "计划关键路径为：需求条款与 SRS → M1需求确认 → 总体及详细设计 → 视频/消息连通验证 → M2设计评审 → 公共基础与核心业务增量 → 指挥与H5集成 → 外部系统联动 → W14数据初始化候选与W15性能加固 → W15后段单元/集成测试 → W16前段初验部署 → W16性能兼容、安全可移植与功能追踪验证 → M3初验 → W17—W20逐周试运行与真实演练 → W20前段M4试运行准出 → W20后段M5竣工验收。": "计划关键路径为：需求条款与 SRS → M1需求确认 → 总体及详细设计 → 视频/消息连通验证 → M2设计评审 → 公共基础与核心业务增量 → 指挥与H5集成 → 外部系统联动 → W14数据初始化候选与W15性能加固 → W15后段单元/集成测试 → W16前段初验部署 → W16性能兼容、安全可移植与功能追踪验证 → M3初验 → 形成试运行启动记录并从该时点连续监测 → W17—W20逐周试运行与真实演练 → W20前段M4试运行准出 → W20后段M5竣工验收。",
        "数据初始化技术实施由 B 在 W14第5个工作日完成，C 仅提供规则、对账口径与符合性输入；W15—W16 的 B 主责工作按 WBS 所列工作日顺序执行，不以同周并列掩盖前置依赖。文档编制和测试设计与开发并行推进，但均受阶段评审门禁控制。任一关键路径工作包未满足完成判据，A 不得将后续里程碑写成已达成。": "数据初始化技术实施由 B 在 W14第5个工作日完成，C 仅提供规则、对账口径与符合性输入；W15—W16 的 B 主责工作按 WBS 所列工作日顺序执行，不以同周并列掩盖前置依赖。W17 中 6.1 的培训与 6.2 的首周运行监测并行，试运行连续监测不等待全部培训完成；6.2 的准入依据为 M3 初验包和试运行启动记录。文档编制和测试设计与开发并行推进，但均受阶段评审门禁控制。任一关键路径工作包未满足完成判据，A 不得将后续里程碑写成已达成。",
        "本计划 V1.0 进入 C 的计划完整性与一致性复核。执行期间 A 每周根据真实完成证据更新滚动计划，不追溯修改已发生记录。对不影响正式里程碑、范围和验收标准的内部任务调整，在周报记录原因和影响；涉及需求、范围、责任、关键数字、正式里程碑或验收条件的调整，必须取得书面确认并按变更流程升版。": "本计划 V1.2 进入 C 的计划完整性与一致性复验。执行期间 A 每周根据真实完成证据更新滚动计划，不追溯修改已发生记录。对不影响正式里程碑、范围和验收标准的内部任务调整，在周报记录原因和影响；涉及需求、范围、责任、关键数字、正式里程碑或验收条件的调整，必须取得书面确认并按变更流程升版。",
    }
    seen = {old: False for old in paragraph_replacements}
    for paragraph in doc.paragraphs:
        for old, new in paragraph_replacements.items():
            if replace_paragraph_text(paragraph, old, new):
                seen[old] = True

    revision_table = doc.tables[0]
    new_row = revision_table.add_row()
    for source_cell, target_cell in zip(revision_table.rows[-2].cells, new_row.cells):
        target_pr = target_cell._tc.get_or_add_tcPr()
        for child in list(target_pr):
            target_pr.remove(child)
        source_pr = source_cell._tc.tcPr
        if source_pr is not None:
            for child in source_pr:
                target_pr.append(deepcopy(child))
    revision_values = [
        "V1.2",
        "2026-09-11",
        "第2—3章",
        "按C复验意见明确W17培训与首周连续监测并行，修复异常空白页，待C再复验。",
        "成员A",
    ]
    for cell, value in zip(new_row.cells, revision_values):
        replace_paragraph_text(cell.paragraphs[0], cell.paragraphs[0].text, value)

    wbs = doc.tables[1]
    replacements = {
        "6.1": [
            "6.1", "试运行启动与培训包", "A", "M3初验包、培训材料、运行方案",
            "试运行启动记录、培训课件/录屏/签到/考核", "5.5", "W17，≤1周",
            "M3初验通过后先形成试运行启动记录并立即启用连续监测；指挥人员培训1天；值班/安保、处置、物资管理各半天；系统管理员培训1天；桌面推演式实操完成；培训不作为首周监测完成的前置条件",
        ],
        "6.2": [
            "6.2", "试运行第一周证据包", "A", "M3初验包、试运行启动记录、运行监测、服务工单、缺陷清单",
            "完整W17运行日志、工单和修复记录", "5.5；与6.1培训并行", "W17，≤1周",
            "自试运行启动记录所载时点起连续监测并覆盖完整W17；培训期间监测不中断；运行事件、可用性和缺陷均有日级记录",
        ],
    }
    found = set()
    for row in wbs.rows[1:]:
        key = row.cells[0].text.strip()
        if key in replacements:
            found.add(key)
            for cell, value in zip(row.cells, replacements[key]):
                replace_paragraph_text(cell.paragraphs[0], cell.paragraphs[0].text, value)

    # Blank page was caused by an odd-page section transition after the landscape WBS section.
    # A normal new-page transition preserves the intended section split without inserting a parity page.
    for section in doc.sections[1:]:
        section.start_type = WD_SECTION_START.NEW_PAGE

    missing = [old for old, value in seen.items() if not value]
    if missing or found != set(replacements):
        raise RuntimeError(f"DOCX replacement mismatch: missing_paragraphs={len(missing)}, wbs={found}")
    doc.save(DOCX)


def revise_pptx() -> None:
    prs = Presentation(PPTX)
    old = "W17—W20均形成运行监测、工单、缺陷和可用率证据。W20先完成M4试运行准出，再组织M5竣工验收；甲方验收窗口变化不降低准出条件。"
    new = "M3初验后形成试运行启动记录并立即连续监测；W17培训与监测并行。W17—W20均形成运行、工单、缺陷和可用率证据；W20先M4准出，再组织M5。"
    count = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text_frame") and shape.text == old:
                shape.text_frame.paragraphs[0].runs[0].text = new
                count += 1
    if count != 1:
        raise RuntimeError(f"PPTX replacement mismatch: {count}")
    prs.save(PPTX)


def revise() -> None:
    revise_docx()
    revise_pptx()
    print("Revised G1-06 DOCX and PPTX to V1.2")


def fix_blank_page() -> None:
    doc = Document(DOCX)
    target = None
    for index, paragraph in enumerate(doc.paragraphs):
        if index + 1 < len(doc.paragraphs) and doc.paragraphs[index + 1].text == "6 风险与计划联动":
            if not paragraph.text and "w:type=\"page\"" in paragraph._p.xml:
                target = paragraph
                break
    if target is None:
        raise RuntimeError("Expected page-break paragraph before chapter 6 was not found")
    target._element.getparent().remove(target._element)
    doc.save(DOCX)
    print("Removed overflowing page-break paragraph before chapter 6")


def render_pdf() -> None:
    output_dir = PDF.parent / "pages"
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(PDF)
    for index, page in enumerate(pdf):
        bitmap = page.render(scale=1.5)
        bitmap.to_pil().save(output_dir / f"page-{index + 1:02d}.png")
    print(f"Rendered {len(pdf)} pages to {output_dir}")


def sync_docx_timeline() -> None:
    doc = Document(DOCX)
    if len(doc.inline_shapes) < 2:
        raise RuntimeError("Expected at least two inline figures in DOCX")
    slide_image = PDF.parent / "ppt" / "幻灯片2.PNG"
    inline = doc.inline_shapes[1]._inline
    relationship_id = inline.graphic.graphicData.pic.blipFill.blip.embed
    image_part = doc.part.related_parts[relationship_id]
    image_part._blob = slide_image.read_bytes()
    doc.save(DOCX)
    print(f"Synchronized DOCX timeline figure from {slide_image.name}")


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "inspect":
        inspect()
    elif len(sys.argv) == 2 and sys.argv[1] == "revise":
        revise()
    elif len(sys.argv) == 2 and sys.argv[1] == "fix-blank-page":
        fix_blank_page()
    elif len(sys.argv) == 2 and sys.argv[1] == "render-pdf":
        render_pdf()
    elif len(sys.argv) == 2 and sys.argv[1] == "sync-docx-timeline":
        sync_docx_timeline()
    else:
        raise SystemExit("usage: revise_g1_06_v12.py inspect")
