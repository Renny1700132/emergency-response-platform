from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]


def set_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = value
        for extra in paragraph.runs[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        paragraph.add_run(value)
    for extra_p in cell.paragraphs[1:]:
        extra_p._element.getparent().remove(extra_p._element)


def set_paragraph_text(paragraph, value):
    if paragraph.runs:
        run = paragraph.runs[0]
        run.text = value
        for extra in paragraph.runs[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        paragraph.add_run(value)


def add_cell_text(cell, value, font_size=8):
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(value)
    run.font.size = Pt(font_size)
    run.font.name = '宋体'
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), '宋体')


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)


def set_table_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:type'), 'dxa')
            tc_w.set(qn('w:w'), str(int(width.cm * 567)))
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn('w:w'), str(int(width.cm * 567)))


def get_rtm_boundaries():
    result = {}
    for line in (ROOT / 'docs/work/C_REQ/rtm_v1.md').read_text(encoding='utf-8').splitlines():
        if not line.startswith('| G2-FR-'):
            continue
        cells = [x.strip() for x in line.strip('|').split('|')]
        if len(cells) == 9:
            result[cells[0]] = cells[5]
    return result


def append_revision(doc, version, sections, description):
    table = doc.tables[0]
    if any(row.cells[0].text.strip() == version for row in table.rows):
        return
    cells = table.add_row().cells
    for cell, value in zip(cells, [version, '2026-09-16', sections, description, '何思源']):
        set_text(cell, value)


def update_srs():
    path = ROOT / 'docs/deliverables/08-软件需求规格说明书SRS.docx'
    doc = Document(path)
    replacements = 0
    for p in doc.paragraphs:
        if 'spec/RTM待同步' in p.text:
            for run in p.runs:
                if 'spec/RTM待同步' in run.text:
                    run.text = run.text.replace('spec/RTM待同步', 'spec/RTM已同步')
                    replacements += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if 'spec/RTM待同步' in p.text:
                        for run in p.runs:
                            if 'spec/RTM待同步' in run.text:
                                run.text = run.text.replace('spec/RTM待同步', 'spec/RTM已同步')
                                replacements += 1
    append_revision(doc, 'V1.3', '第3章', 'G2-R05 同步 RCLR 双镜像状态，清除 spec/RTM 待同步旧表述。')
    doc.save(path)
    return replacements


def update_rtm():
    path = ROOT / 'docs/deliverables/10-需求追踪矩阵RTMv1.docx'
    doc = Document(path)
    mappings = get_rtm_boundaries()
    updated = 0
    for table in doc.tables:
        if not table.rows or not table.rows[0].cells:
            continue
        headers = [c.text.strip() for c in table.rows[0].cells]
        if '裁决边界' not in headers or 'G2-FR' not in headers:
            continue
        boundary_index = headers.index('裁决边界')
        fr_index = headers.index('G2-FR')
        for row in table.rows[1:]:
            fr = row.cells[fr_index].text.strip()
            if fr in mappings:
                set_text(row.cells[boundary_index], mappings[fr])
                updated += 1
    for p in doc.paragraphs:
        if p.text.startswith('3. `G2-CLR`'):
            set_paragraph_text(p, '3. `G2-CLR-001—012` 仅为 G1 上游继承索引；`G2-RCLR-001—010` 为本轮 G2 规格化人工业务裁决，并已挂接受影响 SRS、spec AC 和 RTM 追踪项。两类记录均不替代接口联调、性能测试、安全测试、部署验证或最终验收。')
    if updated != 29:
        raise RuntimeError(f'RTM 更新行数异常：{updated}')
    append_revision(doc, 'V1.1', '第1、2、4章', 'G2-R05 将 G2-RCLR-001—010 同步至裁决边界、追踪规则和覆盖统计。')
    doc.save(path)
    return updated


def update_clarifications():
    path = ROOT / 'docs/deliverables/09-AI反向澄清记录.docx'
    doc = Document(path)
    if any(p.text.strip().startswith('5 G2 规格化阶段人工业务裁决') for p in doc.paragraphs):
        substitutions = {
            'AC-G2-FR-017—020-01—03、024-01—03': 'AC-G2-FR-017-01—03、018-01—03、019-01—03、020-01—03、024-01—03',
            'AC-G2-FR-010—012-01—03、023-01—03': 'AC-G2-FR-010-01—03、011-01—03、012-01—03、023-01—03',
        }
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for old, new in substitutions.items():
                            if old in p.text:
                                set_paragraph_text(p, p.text.replace(old, new))
        set_table_widths(doc.tables[-1], [Cm(2.0), Cm(7.0), Cm(7.0)])
        append_revision(doc, 'V1.1', '第2、5章', 'G2-R05 明确 G1 继承索引并新增 10 条 G2-RCLR 双镜像追踪。')
        doc.save(path)
        return 0
    for p in doc.paragraphs:
        if p.text.strip() == '2 澄清记录':
            set_paragraph_text(p, '2 G1 上游继承裁决索引（G2-CLR-001—012）')
            break
    else:
        raise RuntimeError('未找到澄清记录章节标题')

    stats = doc.tables[1]
    set_text(stats.rows[1].cells[0], 'G1 上游继承索引')
    set_text(stats.rows[1].cells[1], '12 条（G2-CLR-001—012；不计入本轮新澄清）')
    added = stats.add_row().cells
    set_text(added[0], 'G2 规格化新澄清')
    set_text(added[1], '10 条（G2-RCLR-001—010；课程项目人工业务裁决）')

    heading_style = next(p.style for p in doc.paragraphs if p.text.strip().startswith('2 G1 上游'))
    normal_style = next(p.style for p in doc.paragraphs if p.text.strip().startswith('AI 反向澄清用于'))
    caption_style = next((p.style for p in doc.paragraphs if p.text.strip().startswith('图 1-1')), normal_style)
    h = doc.add_paragraph(style=heading_style)
    h.add_run('5 G2 规格化阶段人工业务裁决（G2-RCLR-001—010）')
    intro = doc.add_paragraph(style=normal_style)
    intro.add_run('本节同步 G2-R02 已记录的 10 条课程项目人工业务裁决。G2-CLR-001—012 仅保留为第 2 节所列 G1 上游继承索引；以下裁决不改变 39 条 FR、34 条★FR、冻结数字、责任边界或验收强度，也不表示接口、性能、部署或现场验收已通过。')
    cap = doc.add_paragraph(style=caption_style)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run('表 5-1  G2-RCLR 裁决与双镜像追踪')
    table = doc.add_table(rows=1, cols=3)
    if doc.tables:
        table.style = doc.tables[-1].style
    hdr = table.rows[0]
    for cell, text in zip(hdr.cells, ['编号', '人工裁决与受影响 SRS/FR', 'spec AC 与 RTM 追踪']):
        set_text(cell, text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    set_repeat_table_header(hdr)
    widths = [Cm(2.0), Cm(7.0), Cm(7.0)]
    for cell, width in zip(hdr.cells, widths):
        cell.width = width
    rows = [
        ('G2-RCLR-001', '可靠外部告警 ID 幂等；无可靠 ID 不自动合并。SRS：G2-FR-014、027。', 'AC-G2-FR-014-01—03、027-01—03；RTM：FR-014、027。'),
        ('G2-RCLR-002', '核实超时仅升级提醒，不自动认定有效或启动预案。SRS：G2-FR-014。', 'AC-G2-FR-014-01—03；RTM：FR-014。'),
        ('G2-RCLR-003', '重指派保留历史，默认不改截止时间；改期须显式审计。SRS：G2-FR-015、022。', 'AC-G2-FR-015-01—03、022-01—03；RTM：FR-015、022。'),
        ('G2-RCLR-004', '保存多时间；时间线优先业务时间，迟到数据不覆盖历史。SRS：G2-FR-004、013、014。', 'AC-G2-FR-004-01—03、013-01—03、014-01—03；RTM：FR-004、013、014。'),
        ('G2-RCLR-005', '二维码绑定任务/点位/时段；过期无效、重复扫码幂等。SRS：G2-FR-017—020、024。', 'AC-G2-FR-017—020-01—03、024-01—03；RTM：FR-017—020、024。'),
        ('G2-RCLR-006', '盘点以发布/开始快照比对；期间出入库独立记录并复核。SRS：G2-FR-009、025。', 'AC-G2-FR-009-01—03、025-01—03；RTM：FR-009、025。'),
        ('G2-RCLR-007', '取消/逾期/未完成不计完成；补演关联原计划且保留异常。SRS：G2-FR-010—012、023。', 'AC-G2-FR-010—012-01—03、023-01—03；RTM：FR-010—012、023。'),
        ('G2-RCLR-008', '迟到/重复回执幂等；保留超时、重试、降级与人工历史。SRS：G2-FR-020、022。', 'AC-G2-FR-020-01—03、022-01—03；RTM：FR-020、022。'),
        ('G2-RCLR-009', '位置过期仅显示最后有效位置，禁止自动调派。SRS：G2-FR-005。', 'AC-G2-FR-005-01—03；RTM：FR-005。'),
        ('G2-RCLR-010', '附件逻辑删除、审计保留，归档引用受保护。SRS：G2-FR-013、016、021、022。', 'AC-G2-FR-013-01—03、016-01—03、021-01—03、022-01—03；RTM：FR-013、016、021、022。'),
    ]
    for values in rows:
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, widths):
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            add_cell_text(cell, value)
    set_table_widths(table, widths)
    append_revision(doc, 'V1.1', '第2、5章', 'G2-R05 明确 G1 继承索引并新增 10 条 G2-RCLR 双镜像追踪。')
    doc.save(path)
    return len(rows)


if __name__ == '__main__':
    print({'srs_replacements': update_srs(), 'rtm_rows': update_rtm(), 'rclr_rows': update_clarifications()})
